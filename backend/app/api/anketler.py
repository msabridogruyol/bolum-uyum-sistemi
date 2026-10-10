# -*- coding: utf-8 -*-
"""
[2026-10-10] Anket ve envanterler (modül: anketler).

Yönetim:  GET /yonetim/anket-sablonlari · GET/POST /yonetim/okul/{okul_id}/anketler ·
          GET/PUT/DELETE /yonetim/anket/{anket_id} · POST /yonetim/anket/{anket_id}/durum ·
          GET /yonetim/anket/{anket_id}/sonuclar · GET /yonetim/anket/{anket_id}/excel
Öğrenci:  GET /ogrenci/anketler · GET /ogrenci/anket/{anket_id} · POST /ogrenci/anket/{anket_id}/yanit

Anonimlik: anonim ankette yanıt satırında ogrenci_id ve şube TUTULMAZ (yalnızca sınıf düzeyi); kimin katıldığı ayrı
anket_katilim tablosunda yalnızca "bir kez yanıtlama" kontrolü için durur ve yanıtla eşleştirilemez.
Anonim olmayan tarama formlarında öğrenci kendi seviyesini ve önerisini görür; rehber öğrenci bazında listeyi görür.
"""
import io
import json
import re
from collections import Counter, defaultdict
from datetime import date

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.api.deps import get_mevcut_ogrenci, get_mevcut_yonetim
from app.api.okul_yonetimi import _okul_kapsami
from app.core.anket_sablonlari import LIKERT, SABLONLAR, SORU_TURLERI, destek_gerekiyor_mu, puanla, seviye_bilgisi
from app.core.database import get_db
from app.core.hesap_yonetimi import denetim_yaz
from app.models import AdminKullanici, Ogrenci

yonetim_router = APIRouter(prefix="/yonetim", tags=["Anketler"])
ogrenci_router = APIRouter(prefix="/ogrenci", tags=["Anketler"])
DURUMLAR = {"taslak": "Taslak", "yayinda": "Yayında", "kapandi": "Kapandı"}
# [2026-10-10] Anonimlik eşiği: en az 5 kişiye gönderilir; anonim anket sonuçları en az 5 yanıtla ve en az 5 kişilik gruplarla gösterilir
from app.core.kucuk_grup import EN_AZ_GRUP as EN_AZ, siniflari_birlestir  # noqa: E402  (ortak küçük grup eşiği = 5)
SINIF_SIRA = {s: i for i, s in enumerate(["9. Sınıf", "10. Sınıf", "11. Sınıf", "12. Sınıf", "Mezun", "Aday"])}


def _j(v, bos):
    if v is None:
        return bos
    return v if isinstance(v, (list, dict)) else json.loads(v)


def _anket(db: Session, anket_id: int) -> dict:
    r = db.execute(text("SELECT * FROM anketler WHERE id = :i"), {"i": anket_id}).mappings().first()
    if r is None:
        raise HTTPException(404, "Anket bulunamadı.")
    d = dict(r)
    d["sorular"], d["hedef"], d["puanlama"] = _j(d["sorular"], []), _j(d["hedef"], {}), _j(d["puanlama"], None)
    return d


def _hedefte_mi(a: dict, o: Ogrenci) -> bool:
    h = a["hedef"] or {}
    siniflar, subeler = h.get("siniflar") or [], h.get("subeler") or []
    if not siniflar and not subeler:
        return True
    return (o.sinif in siniflar) or (f"{o.sinif}|{o.sube}" in subeler)


def _hedef_sayisi(db: Session, a: dict) -> int:
    ogr = db.query(Ogrenci).filter(Ogrenci.okul_id == a["okul_id"]).all()
    return sum(1 for o in ogr if _hedefte_mi(a, o) and not getattr(o, "test_hesabi", False))


def _ozet(db: Session, a: dict, katilim: int | None = None) -> dict:
    if katilim is None:
        katilim = db.execute(text("SELECT count(*) FROM anket_katilim WHERE anket_id = :i"), {"i": a["id"]}).scalar() or 0
    bugun = date.today()
    acik = a["durum"] == "yayinda" and (not a["baslangic"] or a["baslangic"] <= bugun) and (not a["bitis"] or a["bitis"] >= bugun)
    return {"id": a["id"], "baslik": a["baslik"], "aciklama": a["aciklama"], "sablon": a["sablon"], "anonim": a["anonim"],
            "durum": a["durum"], "durum_adi": DURUMLAR[a["durum"]], "acik": acik, "hedef": a["hedef"],
            "baslangic": a["baslangic"].isoformat() if a["baslangic"] else None, "bitis": a["bitis"].isoformat() if a["bitis"] else None,
            "soru_sayisi": len(a["sorular"]), "katilim": katilim, "envanter": bool(a["puanlama"]), "olusturan": a["olusturan"]}


# ============================================================================= yönetim
@yonetim_router.get("/anket-sablonlari")
def sablonlar(yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    return {"sablonlar": [{"kod": k, **v} for k, v in SABLONLAR.items()], "soru_turleri": SORU_TURLERI, "likert": LIKERT}


@yonetim_router.get("/okul/{okul_id}/anketler")
def okul_anketleri(okul_id: int, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    _okul_kapsami(db, yon, okul_id)
    ids = [r[0] for r in db.execute(text("SELECT id FROM anketler WHERE okul_id = :o ORDER BY olusturulma_zamani DESC"), {"o": okul_id}).all()]
    sayi = dict(db.execute(text("SELECT anket_id, count(*) FROM anket_katilim WHERE anket_id = ANY(:i) GROUP BY anket_id"), {"i": ids}).all()) if ids else {}
    liste = []
    for i in ids:
        a = _anket(db, i)
        liste.append({**_ozet(db, a, sayi.get(i, 0)), "hedef_sayisi": _hedef_sayisi(db, a)})
    return {"anketler": liste}


class SoruIstek(BaseModel):
    id: str | None = None
    metin: str = Field(min_length=2, max_length=300)
    tur: str
    secenekler: list[str] = Field(default_factory=list)
    zorunlu: bool = True
    ters: bool = False


class AnketIstek(BaseModel):
    baslik: str = Field(min_length=3, max_length=120)
    aciklama: str | None = Field(default=None, max_length=1000)
    sablon: str | None = None
    anonim: bool = True
    hedef: dict = Field(default_factory=dict)          # {"siniflar": ["12. Sınıf"], "subeler": ["11. Sınıf|A"]}
    baslangic: date | None = None
    bitis: date | None = None
    sorular: list[SoruIstek] = Field(default_factory=list)


def _sorulari_dogrula(istek: AnketIstek) -> tuple[list[dict], dict | None]:
    if not istek.sorular:
        raise HTTPException(400, "En az bir soru ekleyin.")
    if len(istek.sorular) > 60:
        raise HTTPException(400, "Bir ankette en fazla 60 soru olabilir.")
    if istek.baslangic and istek.bitis and istek.bitis < istek.baslangic:
        raise HTTPException(400, "Bitiş tarihi başlangıçtan önce olamaz.")
    sorular, gorulen = [], set()
    for n, s in enumerate(istek.sorular, 1):
        if s.tur not in SORU_TURLERI:
            raise HTTPException(400, f"{n}. soru: geçersiz tür.")
        sec = [x.strip()[:120] for x in s.secenekler if x and x.strip()][:12]
        if s.tur in ("tek", "coklu") and len(sec) < 2:
            raise HTTPException(400, f"{n}. soru: en az iki seçenek yazın.")
        sid = s.id if s.id and re.fullmatch(r"s\d{1,3}", s.id) and s.id not in gorulen else f"s{n}"
        while sid in gorulen:
            sid = f"s{int(sid[1:]) + 100}"
        gorulen.add(sid)
        sorular.append({"id": sid, "metin": s.metin.strip(), "tur": s.tur, "secenekler": sec if s.tur in ("tek", "coklu") else [],
                        "zorunlu": s.zorunlu if s.tur != "acik" else False, "ters": s.ters if s.tur == "likert" else False})
    puanlama = SABLONLAR[istek.sablon]["puanlama"] if istek.sablon in SABLONLAR else None
    return sorular, puanlama


def _hedef_temizle(h: dict) -> dict:
    return {"siniflar": [str(x) for x in (h.get("siniflar") or [])][:10], "subeler": [str(x) for x in (h.get("subeler") or [])][:60]}


@yonetim_router.post("/okul/{okul_id}/anketler", status_code=201)
def anket_olustur(okul_id: int, istek: AnketIstek, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    if _okul_kapsami(db, yon, okul_id) is None:
        raise HTTPException(400, "Okul harici öğrencilere anket gönderilemez.")
    sorular, puanlama = _sorulari_dogrula(istek)
    yeni = db.execute(text("""
        INSERT INTO anketler (okul_id, baslik, aciklama, sablon, anonim, hedef, baslangic, bitis, sorular, puanlama, olusturan)
        VALUES (:ok, :b, :a, :s, :an, CAST(:h AS JSONB), :bs, :bt, CAST(:q AS JSONB), CAST(:p AS JSONB), :y) RETURNING id
    """), {"ok": okul_id, "b": istek.baslik.strip(), "a": (istek.aciklama or "").strip() or None,
           "s": istek.sablon if istek.sablon in SABLONLAR else None, "an": istek.anonim, "h": json.dumps(_hedef_temizle(istek.hedef), ensure_ascii=False),
           "bs": istek.baslangic, "bt": istek.bitis, "q": json.dumps(sorular, ensure_ascii=False),
           "p": json.dumps(puanlama, ensure_ascii=False) if puanlama else None, "y": yon.ad_soyad}).scalar()
    denetim_yaz(db, yon, "anket_olustur", "anketler", yeni, istek.baslik, okul_id)
    db.commit()
    return {"id": yeni}


@yonetim_router.get("/anket/{anket_id}")
def anket_detay(anket_id: int, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    a = _anket(db, anket_id)
    _okul_kapsami(db, yon, a["okul_id"])
    return {**_ozet(db, a), "sorular": a["sorular"], "puanlama": a["puanlama"], "hedef_sayisi": _hedef_sayisi(db, a)}


@yonetim_router.put("/anket/{anket_id}", status_code=204)
def anket_duzenle(anket_id: int, istek: AnketIstek, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    a = _anket(db, anket_id)
    _okul_kapsami(db, yon, a["okul_id"])
    katilim = db.execute(text("SELECT count(*) FROM anket_katilim WHERE anket_id = :i"), {"i": anket_id}).scalar() or 0
    sorular, puanlama = _sorulari_dogrula(istek)
    if katilim and (json.dumps(sorular, sort_keys=True) != json.dumps(a["sorular"], sort_keys=True) or istek.anonim != a["anonim"]):
        raise HTTPException(400, "Yanıt gelmiş bir anketin soruları ve anonimliği değiştirilemez; yalnızca başlık, açıklama, hedef ve tarihler değişebilir.")
    db.execute(text("""
        UPDATE anketler SET baslik = :b, aciklama = :a, anonim = :an, hedef = CAST(:h AS JSONB), baslangic = :bs, bitis = :bt,
               sorular = CAST(:q AS JSONB) WHERE id = :i
    """), {"b": istek.baslik.strip(), "a": (istek.aciklama or "").strip() or None, "an": istek.anonim,
           "h": json.dumps(_hedef_temizle(istek.hedef), ensure_ascii=False), "bs": istek.baslangic, "bt": istek.bitis,
           "q": json.dumps(sorular, ensure_ascii=False), "i": anket_id})
    _ = puanlama
    denetim_yaz(db, yon, "anket_duzenle", "anketler", anket_id, istek.baslik, a["okul_id"])
    db.commit()


class DurumIstek(BaseModel):
    durum: str = Field(pattern="^(taslak|yayinda|kapandi)$")


@yonetim_router.post("/anket/{anket_id}/durum", status_code=204)
def anket_durum(anket_id: int, istek: DurumIstek, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    a = _anket(db, anket_id)
    _okul_kapsami(db, yon, a["okul_id"])
    if istek.durum == "yayinda" and _hedef_sayisi(db, a) < EN_AZ:
        raise HTTPException(400, f"Anket en az {EN_AZ} öğrenciye gönderilmelidir; hedef kitleyi genişletin.")
    db.execute(text("UPDATE anketler SET durum = :d WHERE id = :i"), {"d": istek.durum, "i": anket_id})
    if istek.durum == "yayinda" and a["durum"] == "taslak":   # ilk yayında hedef öğrencilere bildirim
        from app.core.bildirim import bildir
        hedef = [o.id for o in db.query(Ogrenci).filter(Ogrenci.okul_id == a["okul_id"]).all() if _hedefte_mi(a, o)]
        bildir(db, "ogrenci", hedef, "anket", f"Yeni anket: {a['baslik']}", "Birkaç dakikanı ayırıp yanıtlar mısın?", f"/anketler?anket={anket_id}", a["okul_id"])
    denetim_yaz(db, yon, f"anket_{istek.durum}", "anketler", anket_id, a["baslik"], a["okul_id"])
    db.commit()


@yonetim_router.delete("/anket/{anket_id}", status_code=204)
def anket_sil(anket_id: int, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    a = _anket(db, anket_id)
    _okul_kapsami(db, yon, a["okul_id"])
    db.execute(text("DELETE FROM anketler WHERE id = :i"), {"i": anket_id})
    denetim_yaz(db, yon, "anket_sil", "anketler", anket_id, a["baslik"], a["okul_id"])
    db.commit()


def _kucuk_siniflari_birlestir(siniflar: list) -> dict:
    """Anonim ankette EN_AZ'dan az yanıtlı sınıflar 'Diğer sınıflar' olarak birleşir; birleşik grup da azsa sınıf bilgisi gösterilmez."""
    return siniflari_birlestir(siniflar, "Diğer sınıflar")


def _sonuclar(db: Session, a: dict) -> dict:
    rows = db.execute(text("""
        SELECT y.cevaplar, y.puan, y.seviye, y.sinif, y.sube, y.olusturulma_zamani, o.ad_soyad, o.id AS oid
          FROM anket_yanitlari y LEFT JOIN ogrenciler o ON o.id = y.ogrenci_id WHERE y.anket_id = :i ORDER BY y.id
    """), {"i": a["id"]}).all()
    if a["anonim"] and len(rows) < EN_AZ:
        return {"sorular": [], "toplam": len(rows), "gizli": True, "esik": EN_AZ}
    cev = [_j(r.cevaplar, {}) for r in rows]
    sorular = []
    for s in a["sorular"]:
        vals = [c.get(s["id"]) for c in cev if c.get(s["id"]) not in (None, "", [])]
        x = {**s, "yanit": len(vals)}
        if s["tur"] in ("likert", "puan"):
            nums = [v for v in vals if isinstance(v, (int, float))]
            ust = 5 if s["tur"] == "likert" else 10
            x["dagilim"] = [sum(1 for v in nums if v == k) for k in range(1, ust + 1)]
            x["ortalama"] = round(sum(nums) / len(nums), 2) if nums else None
            x["etiketler"] = LIKERT if s["tur"] == "likert" else [str(k) for k in range(1, 11)]
        elif s["tur"] in ("tek", "coklu"):
            say = Counter()
            for v in vals:
                for t in (v if isinstance(v, list) else [v]):
                    say[t] += 1
            x["dagilim"] = [say.get(t, 0) for t in s["secenekler"]]
            x["etiketler"] = s["secenekler"]
        else:
            x["metinler"] = [str(v)[:1000] for v in vals][-200:]
        sorular.append(x)
    sonuc = {"sorular": sorular, "toplam": len(rows), "esik": EN_AZ}
    if a["puanlama"]:
        puanlar = [float(r.puan) for r in rows if r.puan is not None]
        seviye = Counter(r.seviye for r in rows if r.seviye)
        sonuc["envanter"] = {
            "ortalama": round(sum(puanlar) / len(puanlar), 2) if puanlar else None, "yon": a["puanlama"]["yon"],
            "seviyeler": [{"kod": s["kod"], "ad": s["ad"], "sayi": seviye.get(s["kod"], 0)} for s in a["puanlama"]["seviyeler"]],
        }
        sinif = defaultdict(list)
        etiket = _kucuk_siniflari_birlestir([r.sinif or "—" for r in rows]) if a["anonim"] else None
        for r in rows:
            if r.puan is not None:
                k = etiket.get(r.sinif or "—") if etiket is not None else (r.sinif or "—")
                if k:
                    sinif[k].append(float(r.puan))
        sonuc["envanter"]["siniflar"] = sorted([{"sinif": k, "sayi": len(v), "ortalama": round(sum(v) / len(v), 2)} for k, v in sinif.items()],
                                               key=lambda x: (SINIF_SIRA.get(x["sinif"], 99), x["sinif"]))
        if not a["anonim"]:
            sonuc["ogrenciler"] = sorted([{"ogrenci_id": str(r.oid) if r.oid else None, "ad_soyad": r.ad_soyad or "(silinmiş hesap)",
                                           "sinif": f"{(r.sinif or '').replace('. Sınıf', '')}-{r.sube}" if r.sube else (r.sinif or ""),
                                           "puan": float(r.puan) if r.puan is not None else None, "seviye": r.seviye,
                                           "destek": destek_gerekiyor_mu(a["puanlama"], r.seviye)} for r in rows],
                                         key=lambda x: (not x["destek"], -(x["puan"] or 0) if a["puanlama"]["yon"] == "yuksek_kotu" else (x["puan"] or 0)))
    return sonuc


@yonetim_router.get("/anket/{anket_id}/sonuclar")
def anket_sonuclari(anket_id: int, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    a = _anket(db, anket_id)
    _okul_kapsami(db, yon, a["okul_id"])
    return {"anket": {**_ozet(db, a), "hedef_sayisi": _hedef_sayisi(db, a)}, **_sonuclar(db, a)}


@yonetim_router.get("/anket/{anket_id}/excel")
def anket_excel(anket_id: int, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from app.api.raporlar import _dosya_adi
    a = _anket(db, anket_id)
    _okul_kapsami(db, yon, a["okul_id"])
    rows = db.execute(text("""
        SELECT y.cevaplar, y.puan, y.seviye, y.sinif, y.sube, y.olusturulma_zamani, o.ad_soyad
          FROM anket_yanitlari y LEFT JOIN ogrenciler o ON o.id = y.ogrenci_id WHERE y.anket_id = :i ORDER BY y.id
    """), {"i": anket_id}).all()
    etiket = None
    if a["anonim"]:   # anonimlik: eşik altı dosya yok, tarih yok, küçük sınıflar birleşik, satır sırası karışık
        if len(rows) < EN_AZ:
            raise HTTPException(400, f"Anonim anketin sonuçları en az {EN_AZ} yanıt gelince indirilebilir.")
        import random
        rows = list(rows)
        random.shuffle(rows)
        etiket = _kucuk_siniflari_birlestir([r.sinif or "—" for r in rows])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Yanıtlar"
    kol = (["Sınıf"] if a["anonim"] else ["Öğrenci", "Şube", "Sınıf", "Tarih"]) + [f"{i}. {s['metin'][:60]}" for i, s in enumerate(a["sorular"], 1)]
    if a["puanlama"]:
        kol += ["Puan (1–5)", "Seviye"]
    for j, k in enumerate(kol, 1):
        c = ws.cell(row=1, column=j, value=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5D8A")
    sev = {s["kod"]: s["ad"] for s in (a["puanlama"] or {}).get("seviyeler", [])}
    for i, r in enumerate(rows, 2):
        c = _j(r.cevaplar, {})
        deger = []
        for s in a["sorular"]:
            v = c.get(s["id"])
            deger.append(", ".join(v) if isinstance(v, list) else v)
        satir = ([etiket.get(r.sinif or "—") or "—"] if a["anonim"] else [r.ad_soyad, r.sube, r.sinif, r.olusturulma_zamani.date().isoformat()]) + deger
        if a["puanlama"]:
            satir += [float(r.puan) if r.puan is not None else None, sev.get(r.seviye, r.seviye)]
        for j, v in enumerate(satir, 1):
            ws.cell(row=i, column=j, value=v)
    ws.freeze_panes = "A2"
    b = io.BytesIO()
    wb.save(b)
    denetim_yaz(db, yon, "anket_excel", "anketler", anket_id, a["baslik"], a["okul_id"])
    db.commit()
    return Response(b.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{_dosya_adi("Anket", a["baslik"])}.xlsx"', "Cache-Control": "no-store"})


# ============================================================================= öğrenci
def _ogrenci_anketleri(db: Session, o: Ogrenci) -> list[dict]:
    if not o.okul_id:
        return []
    ids = [r[0] for r in db.execute(text("SELECT id FROM anketler WHERE okul_id = :o AND durum IN ('yayinda','kapandi') ORDER BY olusturulma_zamani DESC"),
                                    {"o": o.okul_id}).all()]
    katildi = {r[0] for r in db.execute(text("SELECT anket_id FROM anket_katilim WHERE ogrenci_id = :o"), {"o": o.id}).all()}
    sonuc = []
    for i in ids:
        a = _anket(db, i)
        if not _hedefte_mi(a, o):
            continue
        oz = _ozet(db, a)
        if not oz["acik"] and i not in katildi:
            continue
        oz["yanitladim"] = i in katildi
        if i in katildi and not a["anonim"] and a["puanlama"]:
            r = db.execute(text("SELECT seviye FROM anket_yanitlari WHERE anket_id = :a AND ogrenci_id = :o"), {"a": i, "o": o.id}).first()
            oz["sonucum"] = seviye_bilgisi(a["puanlama"], r.seviye if r else None)
        sonuc.append(oz)
    return sonuc


@ogrenci_router.get("/anketler")
def anketlerim(db: Session = Depends(get_db), o: Ogrenci = Depends(get_mevcut_ogrenci)):
    liste = _ogrenci_anketleri(db, o)
    return {"anketler": liste, "bekleyen": sum(1 for x in liste if x["acik"] and not x["yanitladim"])}


@ogrenci_router.get("/anket/{anket_id}")
def anket_ac(anket_id: int, db: Session = Depends(get_db), o: Ogrenci = Depends(get_mevcut_ogrenci)):
    a = _anket(db, anket_id)
    if a["okul_id"] != o.okul_id or not _hedefte_mi(a, o) or a["durum"] == "taslak":
        raise HTTPException(404, "Anket bulunamadı.")
    yanitladim = db.execute(text("SELECT 1 FROM anket_katilim WHERE anket_id = :a AND ogrenci_id = :o"), {"a": anket_id, "o": o.id}).first() is not None
    return {**_ozet(db, a), "sorular": a["sorular"], "likert": LIKERT, "yanitladim": yanitladim}


class YanitIstek(BaseModel):
    cevaplar: dict


@ogrenci_router.post("/anket/{anket_id}/yanit")
def anket_yanitla(anket_id: int, istek: YanitIstek, db: Session = Depends(get_db), o: Ogrenci = Depends(get_mevcut_ogrenci)):
    a = _anket(db, anket_id)
    if a["okul_id"] != o.okul_id or not _hedefte_mi(a, o) or not _ozet(db, a)["acik"]:
        raise HTTPException(400, "Bu anket şu an yanıtlanamaz.")
    if db.execute(text("SELECT 1 FROM anket_katilim WHERE anket_id = :a AND ogrenci_id = :o"), {"a": anket_id, "o": o.id}).first():
        raise HTTPException(400, "Bu anketi zaten yanıtladın.")
    temiz = {}
    for s in a["sorular"]:
        v = istek.cevaplar.get(s["id"])
        if s["tur"] == "likert":
            v = int(v) if isinstance(v, (int, float)) and 1 <= int(v) <= 5 else None
        elif s["tur"] == "puan":
            v = int(v) if isinstance(v, (int, float)) and 1 <= int(v) <= 10 else None
        elif s["tur"] == "tek":
            v = v if v in s["secenekler"] else None
        elif s["tur"] == "coklu":
            v = [x for x in (v if isinstance(v, list) else []) if x in s["secenekler"]] or None
        else:
            v = str(v).strip()[:1000] if v else None
        if v is None and s.get("zorunlu"):
            raise HTTPException(400, f"Lütfen şu soruyu yanıtla: {s['metin'][:80]}")
        if v is not None:
            temiz[s["id"]] = v
    try:
        # anonim ankette zaman damgaları gün düzeyinde tutulur (katılım ile yanıt saatten eşleştirilemesin)
        db.execute(text("INSERT INTO anket_katilim (anket_id, ogrenci_id, zaman) VALUES (:a, :o, CASE WHEN :an THEN date_trunc('day', now()) ELSE now() END)"),
                   {"a": anket_id, "o": o.id, "an": a["anonim"]})
    except Exception:
        db.rollback()
        raise HTTPException(400, "Bu anketi zaten yanıtladın.")
    puan, seviye = puanla(a["sorular"], temiz, a["puanlama"])
    db.execute(text("""
        INSERT INTO anket_yanitlari (anket_id, ogrenci_id, sinif, sube, cevaplar, puan, seviye, olusturulma_zamani)
        VALUES (:a, :o, :s, :sb, CAST(:c AS JSONB), :p, :sv, CASE WHEN :an THEN date_trunc('day', now()) ELSE now() END)
    """), {"an": a["anonim"], "a": anket_id, "o": None if a["anonim"] else o.id, "s": o.sinif, "sb": None if a["anonim"] else o.sube,
           "c": json.dumps(temiz, ensure_ascii=False), "p": puan, "sv": seviye})
    db.commit()
    return {"tamam": True, "sonucum": None if a["anonim"] else seviye_bilgisi(a["puanlama"], seviye)}


# ============================================================================= erken uyarı için
def tarama_uyarilari(db: Session, ogrenci_idler: list) -> dict:
    """{ogrenci_id: [(anket başlığı, seviye adı)]} — son 120 gündeki anonim olmayan taramalarda destek gerektiren sonuçlar."""
    if not ogrenci_idler:
        return {}
    rows = db.execute(text("""
        SELECT DISTINCT ON (y.ogrenci_id, y.anket_id) y.ogrenci_id, y.seviye, a.baslik, a.puanlama
          FROM anket_yanitlari y JOIN anketler a ON a.id = y.anket_id
         WHERE y.ogrenci_id = ANY(:i) AND NOT a.anonim AND a.puanlama IS NOT NULL AND y.olusturulma_zamani >= now() - interval '120 days'
         ORDER BY y.ogrenci_id, y.anket_id, y.id DESC
    """), {"i": ogrenci_idler}).all()
    sonuc = defaultdict(list)
    for r in rows:
        p = _j(r.puanlama, None)
        if destek_gerekiyor_mu(p, r.seviye):
            sonuc[r.ogrenci_id].append((r.baslik, (seviye_bilgisi(p, r.seviye) or {}).get("ad", r.seviye)))
    return sonuc
