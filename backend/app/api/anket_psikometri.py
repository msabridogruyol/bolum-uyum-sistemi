# -*- coding: utf-8 -*-
"""
[2026-10-10] Anket şablonlarının psikometrisi (yalnızca süper admin).

GET /yonetim/anket-psikometri        — her şablon için tüm okulların yanıtları birleşik: n, okul sayısı, Cronbach alfa (+ Feldt %95 GA),
                                       madde istatistikleri, uyarı bayrakları (hesap: app/core/psikometri.py)
GET /yonetim/anket-psikometri/excel  — aynı tablo, Excel

Dahil edilen yanıtlar: anketler.sablon = şablon kodu VE anketin likert maddeleri (id, metin, ters, sıra) şablonla birebir aynı.
Okulun maddeleri değiştirdiği / eklediği / çıkardığı anketler dışarıda kalır (sayıları raporlanır). Anonim yanıtlar da madde
düzeyinde (anket_yanitlari.cevaplar) saklandığı için dahildir; burada yalnızca toplu istatistik döner, tek yanıt gösterilmez.
Test hesaplarının (ogrenciler.test_hesabi) anonim olmayan yanıtları dışlanır. Yalnızca tüm likert maddeleri yanıtlanmış satırlar sayılır.
"""
import io
import json

from fastapi import APIRouter, Depends
from fastapi.responses import Response
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.api.deps import get_mevcut_super_admin
from app.core.anket_sablonlari import SABLONLAR
from app.core.database import get_db
from app.core.hesap_yonetimi import denetim_yaz
from app.core.psikometri import EN_AZ_ISTATISTIK, EN_AZ_YORUM, ESIK_ALFA, ESIK_RIT, sablon_analizi
from app.models import AdminKullanici

router = APIRouter(prefix="/yonetim", tags=["Anket psikometrisi"])

DURUM_ADI = {"veri_bekleniyor": "Veri bekleniyor", "yetersiz_orneklem": "Yetersiz örneklem", "dikkat": "Dikkat", "iyi": "İyi"}


def _j(v, bos):
    if v is None:
        return bos
    return v if isinstance(v, (list, dict)) else json.loads(v)


def _imza(sorular: list[dict]) -> list[tuple]:
    return [(s.get("id"), (s.get("metin") or "").strip(), bool(s.get("ters"))) for s in sorular if s.get("tur") == "likert"]


def _hesapla(db: Session) -> dict:
    anketler = db.execute(text("SELECT id, okul_id, sablon, sorular FROM anketler WHERE sablon IS NOT NULL")).all()
    sonuc = []
    for kod, sab in SABLONLAR.items():
        imza = _imza(sab["sorular"])
        uygun, degismis = {}, 0
        for a in anketler:
            if a.sablon != kod:
                continue
            if _imza(_j(a.sorular, [])) == imza:
                uygun[a.id] = a.okul_id
            else:
                degismis += 1
        rows = db.execute(text("""
            SELECT y.anket_id, y.cevaplar FROM anket_yanitlari y LEFT JOIN ogrenciler o ON o.id = y.ogrenci_id
             WHERE y.anket_id = ANY(:i) AND NOT COALESCE(o.test_hesabi, false)
        """), {"i": list(uygun)}).all() if uygun else []
        cevaplar = [_j(r.cevaplar, {}) for r in rows]
        analiz = sablon_analizi(sab["sorular"], cevaplar, sab.get("alt_boyutlar"))
        # okul sayısı yalnızca analize giren (tamamlanmış) yanıtların okulları
        likert = [s["id"] for s in sab["sorular"] if s.get("tur") == "likert"]
        tam = [r for r, c in zip(rows, cevaplar) if all(isinstance(c.get(i), int) and 1 <= c.get(i) <= 5 for i in likert)]
        sonuc.append({
            "kod": kod, "baslik": sab["baslik"], "anonim": sab["anonim"], "toplam_yanit": len(rows),
            "okul_sayisi": len({uygun[r.anket_id] for r in tam}), "anket_sayisi": len(uygun), "degismis_anket": degismis,
            **analiz, "durum_adi": DURUM_ADI[analiz["durum"]],
        })
    return {"sablonlar": sonuc, "esikler": {"en_az_istatistik": EN_AZ_ISTATISTIK, "en_az_yorum": EN_AZ_YORUM, "alfa": ESIK_ALFA, "r_it": ESIK_RIT}}


@router.get("/anket-psikometri")
def anket_psikometri(db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_super_admin)):
    return _hesapla(db)


@router.get("/anket-psikometri/excel")
def anket_psikometri_excel(db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_super_admin)):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from app.api.raporlar import _dosya_adi
    from datetime import date

    veri = _hesapla(db)
    wb = openpyxl.Workbook()

    def baslik_yaz(ws, kol):
        for j, k in enumerate(kol, 1):
            c = ws.cell(row=1, column=j, value=k)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2F5D8A")
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Özet"
    baslik_yaz(ws, ["Şablon", "Tamamlanmış yanıt (n)", "Okul", "Uygun anket", "Değiştirilmiş anket (dışarıda)", "Madde",
                    "Cronbach alfa", "GA alt (%95)", "GA üst (%95)", "Durum", "Uyarılar"])
    for i, s in enumerate(veri["sablonlar"], 2):
        ga = s.get("alfa_ga") or [None, None]
        for j, v in enumerate([s["baslik"], s["n"], s["okul_sayisi"], s["anket_sayisi"], s["degismis_anket"], s["madde_sayisi"],
                               s.get("alfa"), ga[0], ga[1], s["durum_adi"], "; ".join(s.get("uyarilar") or [])], 1):
            ws.cell(row=i, column=j, value=v)
    satir = len(veri["sablonlar"]) + 3
    ws.cell(row=satir, column=1, value=f"n < {EN_AZ_ISTATISTIK}: istatistik verilmez. n < {EN_AZ_YORUM}: yorum için yetersiz örneklem. "
                                       f"Eşikler: alfa ≥ {ESIK_ALFA}, düzeltilmiş madde-toplam r ≥ {ESIK_RIT}. GA: Feldt yaklaşımı.")
    for s in veri["sablonlar"]:
        alt = s.get("alt_boyutlar") or []
        for a in alt:
            satir += 1
            ga = a.get("alfa_ga") or [None, None]
            for j, v in enumerate([f"{s['baslik']} — {a['ad']}", a["n"], None, None, None, a["madde_sayisi"], a.get("alfa"), ga[0], ga[1],
                                   DURUM_ADI.get(a["durum"]), "; ".join(a.get("uyarilar") or [])], 1):
                ws.cell(row=satir, column=j, value=v)

    for s in veri["sablonlar"]:
        ws = wb.create_sheet(s["kod"][:31])
        baslik_yaz(ws, ["Madde", "Metin", "Ters", "Ortalama", "SS", "Düzeltilmiş r_it", "Silinirse alfa", "Taban %", "Tavan %", "Uyarı"])
        if not s.get("maddeler"):
            ws.cell(row=2, column=1, value=f"Veri bekleniyor (n = {s['n']}, en az {EN_AZ_ISTATISTIK} tamamlanmış yanıt gerekir).")
            continue
        for i, m in enumerate(s["maddeler"], 2):
            for j, v in enumerate([m["id"], m["metin"], "evet" if m["ters"] else "", m["ortalama"], m["ss"], m["r_it"], m["silinirse_alfa"],
                                   m["taban_yuzde"], m["tavan_yuzde"], "; ".join(m["uyarilar"])], 1):
                ws.cell(row=i, column=j, value=v)
        ws.column_dimensions["B"].width = 60

    b = io.BytesIO()
    wb.save(b)
    denetim_yaz(db, yon, "anket_psikometri_excel", "anketler", "-", "Anket psikometrisi Excel")
    db.commit()
    return Response(b.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{_dosya_adi("Anket_psikometrisi", date.today().isoformat())}.xlsx"',
                             "Cache-Control": "no-store"})
