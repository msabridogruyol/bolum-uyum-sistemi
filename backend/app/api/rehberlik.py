# -*- coding: utf-8 -*-
"""
[2026-10-10] Rehberlik: görüşme kayıtları (randevu, not, takip) ve erken uyarı listesi.

Okul yetkilisi (rehber öğretmen) + süper admin — modül: rehberlik
  GET    /yonetim/okul/{okul_id}/erken-uyari            — dikkat gerektiren öğrenciler (kural bazlı, canlı hesaplanır)
  POST   /yonetim/ogrenci/{ogrenci_id}/risk-ertele      — {kural, gun}  "görüştüm, N gün gösterme"
  GET    /yonetim/okul/{okul_id}/gorusmeler             — görüşmeler + yaklaşan randevular + bekleyen takipler + özet
  GET    /yonetim/okul/{okul_id}/gorusmeler/excel       — yıl sonu rehberlik raporu için tablo
  GET    /yonetim/ogrenci/{ogrenci_id}/gorusmeler       — öğrencinin görüşmeleri + şu anki uyarıları
  POST   /yonetim/ogrenci/{ogrenci_id}/gorusmeler       — yeni görüşme / randevu
  PUT    /yonetim/gorusme/{gorusme_id}                  — düzenle (randevuyu "yapıldı" yap, takibi kapat …)
  DELETE /yonetim/gorusme/{gorusme_id}

Öğrenci: planlanan görüşme Takvim'inde "Rehberlik görüşmesi" olarak görünür (notlar asla gösterilmez).
"""
import io
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.api.deps import get_mevcut_yonetim
from app.api.okul_yonetimi import _durum, _ilerleme, _ogrenci_kapsami, _okul_kapsami, _sinif_metni
from app.core.database import get_db
from app.core.hesap_yonetimi import denetim_yaz
from app.core.paketler import okul_modulleri
from app.models import AdminKullanici, Ogrenci

router = APIRouter(prefix="/yonetim", tags=["Rehberlik"])
TR = ZoneInfo("Europe/Istanbul")

TURLER = {"bireysel": "Bireysel görüşme", "veli": "Veli görüşmesi", "grup": "Grup görüşmesi", "ogretmen": "Öğretmenle görüşme"}
KONULAR = {"akademik": "Akademik / ders", "kariyer": "Bölüm ve meslek", "sinav": "Sınav hazırlığı", "motivasyon": "Motivasyon",
           "kisisel": "Kişisel / sosyal", "aile": "Aile", "devamsizlik": "Devamsızlık", "diger": "Diğer"}
DURUMLAR = {"planlandi": "Planlandı", "yapildi": "Yapıldı", "iptal": "İptal"}

SEVIYE_PUAN = {"yuksek": 3, "orta": 2, "dusuk": 1}
KURALLAR = {
    "giris_yok": "Hiç giriş yapmadı",
    "uzak_kaldi": "Uzun süredir girmiyor",
    "yarim_test": "Testi yarıda kaldı",
    "test_baslamadi": "Teste başlamadı",
    "net_dususu": "Netlerinde düşüş",
    "gorev_birakti": "Görevleri bıraktı",
    "hedef_yok": "Hedef bölüm seçmedi",
    "tarama": "Tarama formunda destek ihtiyacı",
}


def _gun(z) -> int | None:
    if z is None:
        return None
    if z.tzinfo is None:
        z = z.replace(tzinfo=timezone.utc)
    return (datetime.now(timezone.utc) - z).days


def _net_dususleri(db: Session, idler: list) -> dict:
    """{ogrenci_id: [(oturum, onceki_ort, son, yuzde)]} — son deneme, önceki en çok 3 denemenin ortalamasından %15+ düşükse."""
    if not idler:
        return {}
    satirlar = db.execute(text("""
        SELECT ogrenci_id, oturum, toplam_net FROM (
            SELECT ogrenci_id, oturum, toplam_net,
                   row_number() OVER (PARTITION BY ogrenci_id, oturum ORDER BY tarih DESC, id DESC) AS sira
              FROM ogrenci_denemeleri WHERE ogrenci_id = ANY(:idler)
        ) x WHERE sira <= 4 ORDER BY ogrenci_id, oturum, sira
    """), {"idler": idler}).all()
    gruplar = defaultdict(list)
    for r in satirlar:
        gruplar[(r.ogrenci_id, r.oturum)].append(float(r.toplam_net))
    sonuc = defaultdict(list)
    for (oid, oturum), netler in gruplar.items():
        if len(netler) < 3:
            continue
        son, onceki = netler[0], netler[1:]
        ort = sum(onceki) / len(onceki)
        if ort >= 5 and son < ort * 0.85:
            sonuc[oid].append((oturum, round(ort, 1), round(son, 1), round((ort - son) / ort * 100)))
    return sonuc


def _gorev_birakanlar(db: Session, idler: list) -> set:
    """Son 3 haftada hiç görev tamamlamamış ama önceki 5 haftada en az 2 görev tamamlamış öğrenciler."""
    if not idler:
        return set()
    bugun = datetime.now(timezone.utc)
    satirlar = db.execute(text("""
        SELECT ogrenci_id,
               count(*) FILTER (WHERE tamamlanma_zamani >= :s3) AS yeni,
               count(*) FILTER (WHERE tamamlanma_zamani < :s3 AND tamamlanma_zamani >= :s8) AS eski
          FROM ogrenci_haftalik_gorev
         WHERE ogrenci_id = ANY(:idler) AND durum = 'tamamlandi'
         GROUP BY ogrenci_id
    """), {"idler": idler, "s3": bugun - timedelta(days=21), "s8": bugun - timedelta(days=56)}).all()
    return {r.ogrenci_id for r in satirlar if r.yeni == 0 and r.eski >= 2}


def uyarilari_hesapla(db: Session, okul_id: int, ogrenci_id=None) -> list[dict]:
    """Okuldaki (ya da tek öğrencinin) uyarıları. Ertelenen kurallar çıkarılır."""
    q = db.query(Ogrenci).filter(Ogrenci.okul_id.is_(None) if okul_id == 0 else Ogrenci.okul_id == okul_id)
    if ogrenci_id is not None:
        q = q.filter(Ogrenci.id == ogrenci_id)
    ogrenciler = [o for o in q.all() if not getattr(o, "test_hesabi", False)]
    if not ogrenciler:
        return []
    idler = [o.id for o in ogrenciler]
    moduller = set(okul_modulleri(db, okul_id or None))
    il = _ilerleme(db, okul_id, ogrenci_id) if ogrenci_id is not None else _ilerleme(db, okul_id)
    netler = _net_dususleri(db, idler) if "net_takibi" in moduller else {}
    gorev = _gorev_birakanlar(db, idler) if "kocluk" in moduller else set()
    tarama = {}
    if "anketler" in moduller:   # [2026-10-10] anonim olmayan tarama formlarında destek gerektiren sonuç
        try:
            from app.api.anketler import tarama_uyarilari
            tarama = tarama_uyarilari(db, idler)
        except Exception:
            db.rollback()
    ertelenen = {(r.ogrenci_id, r.kural) for r in db.execute(text(
        "SELECT ogrenci_id, kural FROM risk_ertelemeleri WHERE ogrenci_id = ANY(:idler) AND bitis >= CURRENT_DATE"), {"idler": idler}).all()}
    son_gorusme = dict(db.execute(text("""
        SELECT ogrenci_id, max(zaman) FROM rehberlik_gorusmeleri WHERE ogrenci_id = ANY(:idler) AND durum = 'yapildi' GROUP BY ogrenci_id
    """), {"idler": idler}).all())
    yaklasan = dict(db.execute(text("""
        SELECT ogrenci_id, min(zaman) FROM rehberlik_gorusmeleri
         WHERE ogrenci_id = ANY(:idler) AND durum = 'planlandi' AND zaman >= now() - interval '1 day' GROUP BY ogrenci_id
    """), {"idler": idler}).all())

    sonuc = []
    for o in ogrenciler:
        x = il.get(o.id) or {}
        kod, _ = _durum(o, x)
        hesap_gun = _gun(o.olusturulma_zamani) or 0
        giris_gun = _gun(o.son_giris_zamani)
        r = []
        if o.son_giris_zamani is None:
            if hesap_gun >= 7:
                r.append(("giris_yok", "orta" if hesap_gun < 21 else "yuksek", f"Hesabı {hesap_gun} gün önce açıldı, hiç giriş yapmadı."))
        elif giris_gun >= 14:
            r.append(("uzak_kaldi", "yuksek" if giris_gun >= 30 else "orta", f"Son girişi {giris_gun} gün önce."))
        if kod == "devam" and (giris_gun is None or giris_gun >= 7):
            r.append(("yarim_test", "orta", f"Değerlendirmede {x.get('biten', 0)}/{x.get('toplam', 4)} bölümde kaldı."))
        if kod == "baslamadi" and o.son_giris_zamani is not None and hesap_gun >= 14:
            r.append(("test_baslamadi", "dusuk", "Giriş yaptı ama değerlendirmeye başlamadı."))
        for oturum, ort, son, yuzde in netler.get(o.id, []):
            r.append(("net_dususu", "yuksek" if yuzde >= 25 else "orta", f"{oturum} son deneme {son} net; önceki ortalama {ort} (−%{yuzde})."))
        if o.id in gorev:
            r.append(("gorev_birakti", "orta", "Önceki haftalarda görev yapıyordu, son 3 haftada hiç tamamlamadı."))
        for baslik, seviye in tarama.get(o.id, []):
            r.append(("tarama", "orta", f"{baslik}: {seviye}."))
        if x.get("durum") == "tamamlandi" and not x.get("hedef") and o.sinif in ("12. Sınıf", "Mezun"):
            r.append(("hedef_yok", "dusuk", "Testi bitirdi ama hedef bölüm seçmedi (son sınıf)."))
        r = [t for t in r if (o.id, t[0]) not in ertelenen]
        if not r:
            continue
        puan = sum(SEVIYE_PUAN[s] for _, s, _ in r)
        en = max(r, key=lambda t: SEVIYE_PUAN[t[1]])[1]
        sg = son_gorusme.get(o.id)
        yk = yaklasan.get(o.id)
        sonuc.append({
            "ogrenci_id": str(o.id), "ad_soyad": o.ad_soyad, "sinif_metni": _sinif_metni(o), "sinif": o.sinif,
            "seviye": en, "puan": puan,
            "uyarilar": [{"kural": k, "ad": KURALLAR[k], "seviye": s, "metin": m} for k, s, m in r],
            "son_gorusme": sg.astimezone(TR).date().isoformat() if sg else None,
            "yaklasan_gorusme": yk.astimezone(TR).isoformat() if yk else None,
        })
    sonuc.sort(key=lambda s: (-SEVIYE_PUAN[s["seviye"]], -s["puan"], s["ad_soyad"].lower()))
    return sonuc


@router.get("/okul/{okul_id}/erken-uyari")
def erken_uyari(okul_id: int, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    _okul_kapsami(db, yon, okul_id)
    liste = uyarilari_hesapla(db, okul_id)
    kurallar = Counter(u["kural"] for s in liste for u in s["uyarilar"])
    return {
        "ogrenciler": liste,
        "ozet": {"toplam": len(liste), "yuksek": sum(1 for s in liste if s["seviye"] == "yuksek"),
                 "orta": sum(1 for s in liste if s["seviye"] == "orta"), "dusuk": sum(1 for s in liste if s["seviye"] == "dusuk"),
                 "gorusulmemis_yuksek": sum(1 for s in liste if s["seviye"] == "yuksek" and not s["son_gorusme"] and not s["yaklasan_gorusme"])},
        "kurallar": [{"kod": k, "ad": a, "sayi": kurallar.get(k, 0)} for k, a in KURALLAR.items()],
    }


class ErteleIstek(BaseModel):
    kural: str
    gun: int = Field(default=14, ge=1, le=120)


@router.post("/ogrenci/{ogrenci_id}/risk-ertele", status_code=204)
def risk_ertele(ogrenci_id: str, istek: ErteleIstek, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    o = _ogrenci_kapsami(db, yon, ogrenci_id)
    if istek.kural not in KURALLAR:
        raise HTTPException(400, "Geçersiz uyarı türü.")
    db.execute(text("""
        INSERT INTO risk_ertelemeleri (ogrenci_id, kural, bitis, olusturan) VALUES (:o, :k, :b, :y)
        ON CONFLICT (ogrenci_id, kural) DO UPDATE SET bitis = EXCLUDED.bitis, olusturan = EXCLUDED.olusturan, olusturulma_zamani = now()
    """), {"o": o.id, "k": istek.kural, "b": date.today() + timedelta(days=istek.gun), "y": yon.ad_soyad})
    denetim_yaz(db, yon, "risk_ertele", "ogrenciler", o.id, f"{KURALLAR[istek.kural]} · {istek.gun} gün", o.okul_id)
    db.commit()


# ----------------------------------------------------------------------------- görüşmeler
def _satir(r) -> dict:
    d = dict(r._mapping) if hasattr(r, "_mapping") else dict(r)
    z = d["zaman"].astimezone(TR) if d["zaman"].tzinfo else d["zaman"]
    return {
        "id": d["id"], "ogrenci_id": str(d["ogrenci_id"]), "ad_soyad": d.get("ad_soyad"), "sinif_metni": d.get("sinif_metni"),
        "zaman": z.isoformat(), "tarih": z.date().isoformat(), "saat": z.strftime("%H:%M"),
        "durum": d["durum"], "durum_adi": DURUMLAR.get(d["durum"], d["durum"]),
        "tur": d["tur"], "tur_adi": TURLER.get(d["tur"], d["tur"]), "konu": d["konu"], "konu_adi": KONULAR.get(d["konu"], d["konu"]),
        "baslik": d["baslik"], "notlar": d["notlar"], "takip_tarihi": d["takip_tarihi"].isoformat() if d["takip_tarihi"] else None,
        "takip_tamam": d["takip_tamam"], "ogrenciye_goster": d["ogrenciye_goster"], "olusturan": d["olusturan"],
    }


SECIM = """
    SELECT g.*, o.ad_soyad, o.sinif, o.sube FROM rehberlik_gorusmeleri g JOIN ogrenciler o ON o.id = g.ogrenci_id
"""


def _sinif_ekle(rows) -> list[dict]:
    sonuc = []
    for r in rows:
        m = dict(r._mapping)
        sf, sb = m.get("sinif"), m.get("sube")
        m["sinif_metni"] = (f"{sf.replace('. Sınıf', '')}-{sb}" if sb and sf != "Mezun" else sf) if sf else ""
        sonuc.append(_satir(m))
    return sonuc


@router.get("/okul/{okul_id}/gorusmeler")
def okul_gorusmeleri(okul_id: int, gun: int = Query(90, ge=7, le=400), db: Session = Depends(get_db),
                     yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    _okul_kapsami(db, yon, okul_id)
    kosul = "o.okul_id IS NULL" if okul_id == 0 else "o.okul_id = :ok"
    p = {"ok": okul_id, "bas": datetime.now(timezone.utc) - timedelta(days=gun)}
    gecmis = _sinif_ekle(db.execute(text(SECIM + f" WHERE {kosul} AND g.durum <> 'planlandi' AND g.zaman >= :bas ORDER BY g.zaman DESC"), p).all())
    yaklasan = _sinif_ekle(db.execute(text(SECIM + f" WHERE {kosul} AND g.durum = 'planlandi' ORDER BY g.zaman"), p).all())
    takip = _sinif_ekle(db.execute(text(SECIM + f" WHERE {kosul} AND g.takip_tarihi IS NOT NULL AND NOT g.takip_tamam "
                                         "AND g.durum = 'yapildi' ORDER BY g.takip_tarihi"), p).all())
    yapilan = [g for g in gecmis if g["durum"] == "yapildi"]
    bu_ay = date.today().replace(day=1).isoformat()
    return {
        "gecmis": gecmis, "yaklasan": yaklasan, "takipler": takip,
        "ozet": {"bu_ay": sum(1 for g in yapilan if g["tarih"] >= bu_ay), "donem": len(yapilan),
                 "ogrenci": len({g["ogrenci_id"] for g in yapilan}), "geciken_takip": sum(1 for g in takip if g["takip_tarihi"] < date.today().isoformat()),
                 "konular": [{"kod": k, "ad": KONULAR.get(k, k), "sayi": n} for k, n in Counter(g["konu"] for g in yapilan).most_common()]},
        "secenekler": {"turler": TURLER, "konular": KONULAR, "durumlar": DURUMLAR},
    }


@router.get("/okul/{okul_id}/gorusmeler/excel")
def gorusmeler_excel(okul_id: int, gun: int = Query(365, ge=7, le=800), db: Session = Depends(get_db),
                     yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    okul = _okul_kapsami(db, yon, okul_id)
    kosul = "o.okul_id IS NULL" if okul_id == 0 else "o.okul_id = :ok"
    satirlar = _sinif_ekle(db.execute(text(SECIM + f" WHERE {kosul} AND g.zaman >= :bas ORDER BY g.zaman"),
                                      {"ok": okul_id, "bas": datetime.now(timezone.utc) - timedelta(days=gun)}).all())
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Görüşmeler"
    kol = ["Tarih", "Saat", "Öğrenci", "Sınıf", "Görüşme türü", "Konu", "Başlık", "Durum", "Notlar", "Takip tarihi", "Takip", "Kaydeden"]
    for j, k in enumerate(kol, 1):
        c = ws.cell(row=1, column=j, value=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5D8A")
    for i, g in enumerate(satirlar, 2):
        for j, v in enumerate([g["tarih"], g["saat"], g["ad_soyad"], g["sinif_metni"], g["tur_adi"], g["konu_adi"], g["baslik"],
                               g["durum_adi"], g["notlar"], g["takip_tarihi"], ("Tamam" if g["takip_tamam"] else "Bekliyor") if g["takip_tarihi"] else "",
                               g["olusturan"]], 1):
            c = ws.cell(row=i, column=j, value=v)
            if j == 9:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate([12, 7, 24, 8, 18, 18, 26, 10, 60, 12, 9, 18], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    # konu × ay özeti (rehberlik faaliyet raporu için)
    oz = wb.create_sheet("Özet")
    yapilan = [g for g in satirlar if g["durum"] == "yapildi"]
    aylar = sorted({g["tarih"][:7] for g in yapilan})
    oz.cell(row=1, column=1, value="Konu").font = Font(bold=True)
    for j, a in enumerate(aylar, 2):
        oz.cell(row=1, column=j, value=a).font = Font(bold=True)
    oz.cell(row=1, column=len(aylar) + 2, value="Toplam").font = Font(bold=True)
    for i, (k, ad) in enumerate(KONULAR.items(), 2):
        oz.cell(row=i, column=1, value=ad)
        for j, a in enumerate(aylar, 2):
            oz.cell(row=i, column=j, value=sum(1 for g in yapilan if g["konu"] == k and g["tarih"].startswith(a)))
        oz.cell(row=i, column=len(aylar) + 2, value=sum(1 for g in yapilan if g["konu"] == k))
    oz.column_dimensions["A"].width = 22
    b = io.BytesIO()
    wb.save(b)
    denetim_yaz(db, yon, "gorusme_excel", "rehberlik_gorusmeleri", okul_id, f"{len(satirlar)} kayıt", okul_id or None)
    db.commit()
    ad = (okul.ad if okul else "okul_harici").replace(" ", "_")
    return Response(b.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="rehberlik_gorusmeleri_{date.today().isoformat()}.xlsx"',
                             "X-Dosya-Adi": f"{ad}_rehberlik_gorusmeleri.xlsx".encode("ascii", "ignore").decode()})


@router.get("/ogrenci/{ogrenci_id}/gorusmeler")
def ogrenci_gorusmeleri(ogrenci_id: str, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    o = _ogrenci_kapsami(db, yon, ogrenci_id)
    liste = _sinif_ekle(db.execute(text(SECIM + " WHERE g.ogrenci_id = :o ORDER BY g.zaman DESC"), {"o": o.id}).all())
    uyari = uyarilari_hesapla(db, o.okul_id or 0, o.id)
    return {"gorusmeler": liste, "uyarilar": uyari[0]["uyarilar"] if uyari else [],
            "secenekler": {"turler": TURLER, "konular": KONULAR, "durumlar": DURUMLAR}}


class GorusmeIstek(BaseModel):
    zaman: datetime
    durum: str = "yapildi"
    tur: str = "bireysel"
    konu: str = "akademik"
    baslik: str | None = Field(default=None, max_length=120)
    notlar: str | None = Field(default=None, max_length=5000)
    takip_tarihi: date | None = None
    takip_tamam: bool = False
    ogrenciye_goster: bool = True


def _dogrula(istek: GorusmeIstek) -> dict:
    if istek.durum not in DURUMLAR or istek.tur not in TURLER or istek.konu not in KONULAR:
        raise HTTPException(400, "Geçersiz görüşme türü, konusu ya da durumu.")
    z = istek.zaman if istek.zaman.tzinfo else istek.zaman.replace(tzinfo=TR)   # tarayıcıdan gelen yerel saat
    t = lambda v: (v or "").strip() or None  # noqa: E731
    return {"z": z, "du": istek.durum, "tu": istek.tur, "ko": istek.konu, "ba": t(istek.baslik), "no": t(istek.notlar),
            "tt": istek.takip_tarihi, "ta": istek.takip_tamam, "og": istek.ogrenciye_goster}


@router.post("/ogrenci/{ogrenci_id}/gorusmeler", status_code=201)
def gorusme_ekle(ogrenci_id: str, istek: GorusmeIstek, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    o = _ogrenci_kapsami(db, yon, ogrenci_id)
    v = _dogrula(istek)
    yeni = db.execute(text("""
        INSERT INTO rehberlik_gorusmeleri (okul_id, ogrenci_id, zaman, durum, tur, konu, baslik, notlar, takip_tarihi, takip_tamam,
                                           ogrenciye_goster, olusturan_id, olusturan)
        VALUES (:ok, :o, :z, :du, :tu, :ko, :ba, :no, :tt, :ta, :og, :yi, :y) RETURNING id
    """), {**v, "ok": o.okul_id, "o": o.id, "yi": yon.id, "y": yon.ad_soyad}).scalar()
    denetim_yaz(db, yon, "gorusme_ekle", "rehberlik_gorusmeleri", yeni, f"{o.ad_soyad} · {TURLER[v['tu']]} · {DURUMLAR[v['du']]}", o.okul_id)
    if v["du"] == "planlandi" and v["og"]:   # [2026-10-10] bildirim (+ e-posta): yalnızca zaman, konu ve not yok
        from app.core.bildirim import bildir
        z = v["z"].astimezone(TR)
        bildir(db, "ogrenci", [o.id], "rehberlik_randevu", "Rehber öğretmeninle görüşmen planlandı",
               f"{z.strftime('%d.%m.%Y %H:%M')} — rehberlik servisinde görüşelim.", "/takvim", o.okul_id, eposta=True)
    db.commit()
    return {"id": yeni}


def _gorusme(db: Session, yon: AdminKullanici, gorusme_id: int):
    g = db.execute(text("SELECT id, ogrenci_id FROM rehberlik_gorusmeleri WHERE id = :i"), {"i": gorusme_id}).first()
    if g is None:
        raise HTTPException(404, "Görüşme bulunamadı.")
    return g, _ogrenci_kapsami(db, yon, str(g.ogrenci_id))


@router.put("/gorusme/{gorusme_id}", status_code=204)
def gorusme_duzenle(gorusme_id: int, istek: GorusmeIstek, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    g, o = _gorusme(db, yon, gorusme_id)
    v = _dogrula(istek)
    db.execute(text("""
        UPDATE rehberlik_gorusmeleri SET zaman = :z, durum = :du, tur = :tu, konu = :ko, baslik = :ba, notlar = :no,
               takip_tarihi = :tt, takip_tamam = :ta, ogrenciye_goster = :og, guncelleme_zamani = now() WHERE id = :i
    """), {**v, "i": g.id})
    denetim_yaz(db, yon, "gorusme_duzenle", "rehberlik_gorusmeleri", g.id, f"{o.ad_soyad} · {DURUMLAR[v['du']]}", o.okul_id)
    db.commit()


@router.delete("/gorusme/{gorusme_id}", status_code=204)
def gorusme_sil(gorusme_id: int, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    g, o = _gorusme(db, yon, gorusme_id)
    db.execute(text("DELETE FROM rehberlik_gorusmeleri WHERE id = :i"), {"i": g.id})
    denetim_yaz(db, yon, "gorusme_sil", "rehberlik_gorusmeleri", g.id, o.ad_soyad, o.okul_id)
    db.commit()


def ogrenci_randevulari(db: Session, o: Ogrenci, bas: date) -> list[dict]:
    """Öğrencinin takvimi için planlanmış görüşmeler — yalnızca zaman ve tür; notlar hiçbir zaman."""
    sonuc = []
    for r in db.execute(text("""
        SELECT id, zaman, tur FROM rehberlik_gorusmeleri
         WHERE ogrenci_id = :o AND durum = 'planlandi' AND ogrenciye_goster AND zaman >= :bas ORDER BY zaman
    """), {"o": o.id, "bas": bas}).all():
        z = r.zaman.astimezone(TR)
        sonuc.append({"id": f"rehber-{r.id}", "baslik": "Rehberlik görüşmesi" if r.tur != "veli" else "Rehberlik: veli görüşmesi",
                      "aciklama": "Rehber öğretmeninle görüşme", "tur": "toplanti", "tur_adi": "Rehberlik", "ikon": "🧭",
                      "baslangic": z.date(), "bitis": None, "saat": z.strftime("%H:%M"), "link": None, "hedef_sinif": None,
                      "kaynak": "rehberlik", "duzenlenebilir": False})
    return sonuc
