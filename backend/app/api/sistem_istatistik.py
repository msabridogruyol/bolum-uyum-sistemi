# -*- coding: utf-8 -*-
"""
[2026-10-10] Süper admin — sistem geneli istatistikler, rapor merkezi dışa aktarımları ve Kontrol Paneli "dikkat gerektirenler".

GET /yonetim/istatistik/genel?donem=30|90|yil|tum&okul_id=&test=false   — İstatistikler sayfasının tüm verisi (tek istek)
GET /yonetim/istatistik/genel/excel?...                                 — aynı verinin Excel'i
GET /yonetim/istatistik/dikkat                                          — Kontrol Paneli: kısa özet + dikkat gerektirenler
GET /yonetim/audit-log/excel?gun=90                                     — Audit Log dışa aktarımı

Kapsam: öğrenciler (okul filtresi; test hesapları varsayılan hariç). "Dönem" olay tarihlerine uygulanır (tamamlanan tur,
giriş, modül kullanımı, zaman serileri); öğrenci sayıları, huni ve okul tablosu güncel durumdur.
Performans: her blok tek toplu SQL (öğrenci başına sorgu yok); modül tabloları yoksa (eski şema) blok boş döner.
KVKK küçük grup (app/core/kucuk_grup.py): okul satırlarında 5'ten az öğrencili okulun oran / ortalamaları gizlenir (sayı görünür);
tek bir küçük okul seçilirse sayfanın tüm oran / ortalama / dağılımları gizlenir.
"""
import io
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.api.deps import get_mevcut_super_admin
from app.core.database import get_db
from app.core.hesap_yonetimi import denetim_yaz
from app.core.kucuk_grup import DIPNOT, EN_AZ_GRUP, GIZLI_METIN, gizle, yeterli
from app.models import AdminKullanici

router = APIRouter(prefix="/yonetim", tags=["Sistem istatistikleri"])

TZ = ZoneInfo("Europe/Istanbul")
DONEMLER = {"30": "Son 30 gün", "90": "Son 90 gün", "yil": "Bu eğitim yılı", "tum": "Tümü"}
SINIF_SIRASI = ["Aday", "9. Sınıf", "10. Sınıf", "11. Sınıf", "12. Sınıf", "Mezun"]
DURUM_SERILERI = ["Tamamladı", "Devam ediyor", "Başlamadı", "Giriş yapmadı"]
UYUM_BANTLARI = [("Yetiyor (70+)", 70), ("Sınırda (40–69)", 40), ("Yetmiyor (<40)", -1)]

# Modül kullanımı: (anahtar, ad, ikon, tablo (ogrenci_id sütunu x'te), ek koşul, zaman sütunu | None, olay birimi)
MODULLER = [
    ("kocluk", "Koçluk adımı tamamlayan", "🎯", "ogrenci_gelisim_adim_durumu x", "x.durum = 'tamamlandi'", "x.guncelleme_zamani", "adım"),
    ("haftalik", "Haftalık görev tamamlayan", "✅", "ogrenci_haftalik_gorev x", "x.durum = 'tamamlandi'", "x.tamamlanma_zamani", "görev"),
    ("filiz", "Filiz'e yazan", "💬", "ogrenci_koclugu_mesajlari m JOIN ogrenci_koclugu_oturumlari x ON x.id = m.oturum_id",
     "m.rol = 'ogrenci'", "m.olusturulma_zamani", "mesaj"),
    ("simulasyon", "Meslek simülasyonu yapan", "🎬", "simulasyon_sonuclari x", "TRUE", "x.olusturulma_zamani", "simülasyon"),
    ("net", "Deneme neti giren", "📈", "ogrenci_denemeleri x", "TRUE", "x.olusturulma_zamani", "deneme"),
    ("calisma", "Çalışma kaydı giren", "⏱️", "calisma_kayitlari x", "TRUE", "x.olusturulma_zamani", "kayıt"),
    ("portfolyo", "Portfolyo kaydı ekleyen", "🗂️", "portfolyo_kayitlari x", "TRUE", "x.olusturulma_zamani", "kayıt"),
    ("anket", "Anket yanıtlayan", "📋", "anket_katilim x", "TRUE", "x.zaman", "yanıt"),
    ("rehberlik", "Rehberlik görüşmesi yapılan", "🧭", "rehberlik_gorusmeleri x", "x.durum = 'yapildi'", "x.zaman", "görüşme"),
    ("kulup", "Kulübe katılan", "🎭", "kulup_uyelikleri x", "x.durum = 'onaylandi'", "x.talep_zamani", "üyelik"),
    ("kutuphane", "Kütüphanesine ekleyen", "📚", "ogrenci_kutuphane x", "TRUE", "x.olusturulma_zamani", "eser"),
    ("tercih", "Tercih listesi hazırlayan", "🎓", "tercihler x", "TRUE", None, "tercih"),
    ("koc_talebi", "Eğitim koçu talebi açan", "👩‍🏫", "koc_gorusme_talepleri x", "TRUE", "x.olusturulma_zamani", "talep"),
]


# ----------------------------------------------------------------------------- yardımcılar
def _guvenli(db: Session, fn, varsayilan):
    """Modül tablosu yoksa (eski şema) bloğu boş geç; ana işlemi bozma (SAVEPOINT)."""
    try:
        with db.begin_nested():
            return fn()
    except Exception:
        return varsayilan


def _oran(x, n):
    return round(100 * x / n, 1) if n else None


def _ort(liste):
    liste = [float(v) for v in liste if v is not None]
    return round(sum(liste) / len(liste), 1) if liste else None


def _bugun() -> date:
    return datetime.now(TZ).date()


def _donem_baslangic(donem: str) -> date | None:
    b = _bugun()
    if donem == "30":
        return b - timedelta(days=29)
    if donem == "90":
        return b - timedelta(days=89)
    if donem == "yil":
        return date(b.year if b.month >= 9 else b.year - 1, 9, 1)
    return None


def _kapsam(okul_id: int | None, test: bool) -> tuple[str, dict]:
    """Öğrenci kapsamı (o takma adlı ogrenciler)."""
    parca, p = ["TRUE"], {}
    if okul_id == 0:
        parca.append("o.okul_id IS NULL")
    elif okul_id is not None:
        parca.append("o.okul_id = :okul")
        p["okul"] = okul_id
    if not test:
        parca.append("NOT coalesce(o.test_hesabi, false)")
    return " AND ".join(parca), p


def _sinif(ham) -> str:
    from app.api.okul_yonetimi import sinif_ayir
    return sinif_ayir(ham)[0] or "Belirtilmedi"


# ----------------------------------------------------------------------------- ana hesap
def istatistik_verisi(db: Session, donem: str = "30", okul_id: int | None = None, test: bool = False) -> dict:
    if donem not in DONEMLER:
        donem = "30"
    kosul, p = _kapsam(okul_id, test)
    bas_gun = _donem_baslangic(donem)
    bas = datetime.combine(bas_gun, datetime.min.time(), TZ) if bas_gun else None
    if bas is not None:
        p["bas"] = bas

    def zf(sutun):
        return f" AND {sutun} >= :bas" if bas is not None and sutun else ""

    simdi = datetime.now(timezone.utc)

    # --- 1) Öğrenci satırları: güncel durum bayrakları (tek sorgu, EXISTS alt sorguları) ---
    ogr = db.execute(text(f"""
        SELECT o.id, o.okul_id, o.sinif, o.son_giris_zamani AS son_giris, o.olusturulma_zamani AS acilis,
               (o.son_giris_zamani IS NOT NULL OR EXISTS (SELECT 1 FROM ogrenci_hesap_olaylari e
                                                          WHERE e.ogrenci_id = o.id AND e.olay = 'giris')) AS giris,
               EXISTS (SELECT 1 FROM ogrenci_katman_oturumlari k JOIN katmanlar km ON km.id = k.katman_id
                        WHERE k.ogrenci_id = o.id AND k.durum <> 'baslamadi' AND NOT km.kosullu_mu) AS basladi,
               EXISTS (SELECT 1 FROM ogrenci_degerlendirme_turu t WHERE t.ogrenci_id = o.id AND t.durum = 'tamamlandi') AS tamamladi,
               EXISTS (SELECT 1 FROM ogrenci_dal_oturumlari d WHERE d.ogrenci_id = o.id AND d.durum = 'tamamlandi') AS k5,
               EXISTS (SELECT 1 FROM ogrenci_hedef_bolum h WHERE h.ogrenci_id = o.id AND h.aktif_mi) AS hedef,
               EXISTS (SELECT 1 FROM ogrenci_gelisim_adim_durumu a WHERE a.ogrenci_id = o.id AND a.durum = 'tamamlandi') AS kocluk
          FROM ogrenciler o WHERE {kosul}
    """), p).mappings().all()
    # Teste başlayan / tamamlayan öğrenci giriş yapmıştır (eski kayıtlarda son_giris_zamani boş kalabiliyor)
    ogr = [{**r, "giris": r["giris"] or r["basladi"] or r["tamamladi"]} for r in ogr]
    n = len(ogr)

    def aktif(r, gun):
        z = r["son_giris"]
        if z is None:
            return False
        z = z if z.tzinfo else z.replace(tzinfo=timezone.utc)
        return z >= simdi - timedelta(days=gun)

    # --- 2) Tamamlanan turlar (dönem içinde): güven, geçerlilik ---
    turlar = db.execute(text(f"""
        SELECT t.guven_skoru, t.sonuc_gecerli_mi, o.okul_id
          FROM ogrenci_degerlendirme_turu t JOIN ogrenciler o ON o.id = t.ogrenci_id
         WHERE t.durum = 'tamamlandi' AND {kosul} {zf('t.tamamlanma_zamani')}
    """), p).all()
    guvenler = [round(float(t.guven_skoru), 1) for t in turlar if t.guven_skoru is not None]
    gecersiz = sum(1 for t in turlar if not t.sonuc_gecerli_mi)

    # Öğrencinin dönem içindeki son tamamlanan turu (puan dağılımları ve sonuçlar için)
    tt = f"""tt AS (
        SELECT DISTINCT ON (t.ogrenci_id) t.id, t.ogrenci_id, o.okul_id
          FROM ogrenci_degerlendirme_turu t JOIN ogrenciler o ON o.id = t.ogrenci_id
         WHERE t.durum = 'tamamlandi' AND {kosul} {zf('t.tamamlanma_zamani')}
         ORDER BY t.ogrenci_id, t.tur_no DESC)"""

    # --- 3) Katman (K1–K4) puanları: öğrenci başına katman ortalaması ---
    katman_sat = db.execute(text(f"""
        WITH {tt}
        SELECT tt.ogrenci_id, tt.okul_id, km.kod, km.sira, avg(s.puan) AS puan
          FROM tt JOIN ogrenci_degisken_skorlari s ON s.ogrenci_id = tt.ogrenci_id AND s.tur_id = tt.id
          JOIN degiskenler d ON d.id = s.degisken_id JOIN katmanlar km ON km.id = d.katman_id
         WHERE d.dal_id IS NULL AND NOT km.kosullu_mu
         GROUP BY tt.ogrenci_id, tt.okul_id, km.kod, km.sira
    """), p).all()
    katman_kodlari = [k for k, _ in sorted({(r.kod, r.sira) for r in katman_sat}, key=lambda x: x[1])] or \
        [r[0] for r in db.execute(text("SELECT kod FROM katmanlar WHERE NOT kosullu_mu ORDER BY sira")).all()]
    katman_degerleri = defaultdict(list)
    okul_katman = defaultdict(lambda: defaultdict(list))
    for r in katman_sat:
        v = round(float(r.puan), 1)
        katman_degerleri[r.kod].append(v)
        okul_katman[r.okul_id or 0][r.kod].append(v)

    # --- 4) Sonuçlar: 1. sıra önerilen bölüm, alan; hedef bölümler ve hedefte uyum ---
    ilk = db.execute(text(f"""
        WITH {tt}
        SELECT ilk.bolum_id, count(*) AS n FROM (
            SELECT DISTINCT ON (s.tur_id) s.tur_id, s.bolum_id FROM ogrenci_bolum_uyum_skorlari s JOIN tt ON s.ogrenci_id = tt.ogrenci_id AND s.tur_id = tt.id
             ORDER BY s.tur_id, s.toplam_uyum DESC) ilk
         GROUP BY ilk.bolum_id
    """), p).all()
    hedefler = db.execute(text(f"""
        SELECT h.bolum_id, count(*) AS n,
               count(*) FILTER (WHERE u.toplam_uyum >= 70) AS yetiyor,
               count(*) FILTER (WHERE u.toplam_uyum >= 40 AND u.toplam_uyum < 70) AS sinirda,
               count(*) FILTER (WHERE u.toplam_uyum < 40) AS yetmiyor,
               count(*) FILTER (WHERE u.toplam_uyum IS NULL) AS yok
          FROM ogrenci_hedef_bolum h JOIN ogrenciler o ON o.id = h.ogrenci_id
          LEFT JOIN LATERAL (
              SELECT t.id FROM ogrenci_degerlendirme_turu t WHERE t.ogrenci_id = o.id AND t.durum = 'tamamlandi'
               ORDER BY t.tur_no DESC LIMIT 1) st ON TRUE
          LEFT JOIN ogrenci_bolum_uyum_skorlari u ON u.ogrenci_id = o.id AND u.tur_id = st.id AND u.bolum_id = h.bolum_id
         WHERE h.aktif_mi AND {kosul}
         GROUP BY h.bolum_id
    """), p).all()
    bolum_idler = {r.bolum_id for r in ilk} | {r.bolum_id for r in hedefler}
    bolum_ad = dict(db.execute(text("SELECT id, ad FROM bolumler WHERE id = ANY(:i)"), {"i": list(bolum_idler) or [-1]}).all())
    from app.core.skor_motoru import bolum_alan_haritasi
    alan_of = bolum_alan_haritasi(db)
    onerilen = sorted(({"ad": bolum_ad.get(r.bolum_id, "?"), "deger": r.n} for r in ilk), key=lambda x: -x["deger"])
    alanlar = Counter()
    for r in ilk:
        alanlar[alan_of.get(r.bolum_id) or "Alanı tanımsız"] += r.n
    hedeflenen = sorted(({"ad": bolum_ad.get(r.bolum_id, "?"), "deger": r.n, "yetiyor": r.yetiyor, "sinirda": r.sinirda,
                          "yetmiyor": r.yetmiyor, "yok": r.yok} for r in hedefler), key=lambda x: -x["deger"])

    # --- 5) Zaman serileri ---
    if bas_gun is None:
        ilk_gun = min((r["acilis"] for r in ogr if r["acilis"]), default=None)
        bas_seri = ilk_gun.astimezone(TZ).date() if ilk_gun else _bugun() - timedelta(days=29)
    else:
        bas_seri = bas_gun
    aralik = (_bugun() - bas_seri).days
    birim = "day" if aralik <= 45 else ("week" if aralik <= 400 else "month")

    def kova(g: date) -> date:
        if birim == "week":
            return g - timedelta(days=g.weekday())
        if birim == "month":
            return g.replace(day=1)
        return g

    kovalar, g = [], kova(bas_seri)
    while g <= _bugun():
        kovalar.append(g)
        g = g + timedelta(days=1) if birim == "day" else (g + timedelta(days=7) if birim == "week" else
                                                            (g.replace(year=g.year + 1, month=1) if g.month == 12 else g.replace(month=g.month + 1)))
    ps = {**p, "sbas": datetime.combine(kovalar[0], datetime.min.time(), TZ), "birim": birim}
    yerel = "AT TIME ZONE 'Europe/Istanbul'"
    seri_sql = {
        "yeni_hesap": f"SELECT date_trunc(:birim, o.olusturulma_zamani {yerel})::date, count(*) FROM ogrenciler o "
                      f"WHERE {kosul} AND o.olusturulma_zamani >= :sbas GROUP BY 1",
        "giris": f"SELECT date_trunc(:birim, e.zaman {yerel})::date, count(DISTINCT e.ogrenci_id) FROM ogrenci_hesap_olaylari e "
                 f"JOIN ogrenciler o ON o.id = e.ogrenci_id WHERE e.olay = 'giris' AND {kosul} AND e.zaman >= :sbas GROUP BY 1",
        "tur": f"SELECT date_trunc(:birim, t.tamamlanma_zamani {yerel})::date, count(*) FROM ogrenci_degerlendirme_turu t "
               f"JOIN ogrenciler o ON o.id = t.ogrenci_id WHERE t.durum = 'tamamlandi' AND {kosul} AND t.tamamlanma_zamani >= :sbas GROUP BY 1",
    }
    seriler = {}
    for k, sql in seri_sql.items():
        d = dict(db.execute(text(sql), ps).all())
        seriler[k] = [int(d.get(x, 0)) for x in kovalar]
    zaman = {"birim": {"day": "gün", "week": "hafta", "month": "ay"}[birim], "x": [x.isoformat() for x in kovalar], **seriler}

    # --- 6) Modül kullanımı (dönem içinde; tercih listesi güncel) ---
    moduller = []
    for k, ad, ikon, tablo, ek, zs, olay in MODULLER:
        r = _guvenli(db, lambda: db.execute(text(
            f"SELECT count(DISTINCT x.ogrenci_id), count(*) FROM {tablo} JOIN ogrenciler o ON o.id = x.ogrenci_id "
            f"WHERE {kosul} AND {ek} {zf(zs)}"), p).one(), None)
        if r is not None:
            moduller.append({"k": k, "ad": ad, "ikon": ikon, "ogrenci": r[0], "sayi": r[1], "olay": olay, "donemli": zs is not None})
    ayrinti = {}
    ayrinti["haftalik"] = _guvenli(db, lambda: dict(db.execute(text(
        f"""SELECT count(*) AS toplam, count(*) FILTER (WHERE x.durum = 'tamamlandi') AS tamam FROM ogrenci_haftalik_gorev x
            JOIN ogrenciler o ON o.id = x.ogrenci_id WHERE {kosul} {zf('x.olusturulma_zamani')}"""), p).mappings().one()), None)
    ayrinti["filiz"] = _guvenli(db, lambda: dict(db.execute(text(
        f"""SELECT count(*) AS oturum, count(*) FILTER (WHERE x.durum = 'tamamlandi') AS biten FROM ogrenci_koclugu_oturumlari x
            JOIN ogrenciler o ON o.id = x.ogrenci_id WHERE {kosul} {zf('x.baslama_zamani')}"""), p).mappings().one()), None)
    ayrinti["simulasyon"] = _guvenli(db, lambda: dict(db.execute(text(
        f"""SELECT count(*) AS sayi, round(avg(x.keyif)::numeric, 1)::float AS keyif FROM simulasyon_sonuclari x
            JOIN ogrenciler o ON o.id = x.ogrenci_id WHERE {kosul} {zf('x.olusturulma_zamani')}"""), p).mappings().one()), None)
    ayrinti["simulasyon_meslekler"] = _guvenli(db, lambda: [
        {"ad": r[0], "deger": r[1], "keyif": float(r[2]) if r[2] is not None else None} for r in db.execute(text(
            f"""SELECT coalesce(x.meslek_ad, '?'), count(*), round(avg(x.keyif)::numeric, 1) FROM simulasyon_sonuclari x
                JOIN ogrenciler o ON o.id = x.ogrenci_id WHERE {kosul} {zf('x.olusturulma_zamani')}
                GROUP BY 1 ORDER BY 2 DESC, 1 LIMIT 15"""), p).all()], [])
    ayrinti["rehberlik"] = _guvenli(db, lambda: dict(db.execute(text(
        f"""SELECT count(*) FILTER (WHERE x.durum = 'yapildi') AS yapildi, count(*) FILTER (WHERE x.durum = 'planlandi') AS planlandi,
                   count(*) AS toplam FROM rehberlik_gorusmeleri x JOIN ogrenciler o ON o.id = x.ogrenci_id
             WHERE {kosul} {zf('x.zaman')}"""), p).mappings().one()), None)
    ayrinti["kulup"] = _guvenli(db, lambda: dict(db.execute(text(
        f"""SELECT x.durum, count(*) FROM kulup_uyelikleri x JOIN ogrenciler o ON o.id = x.ogrenci_id
             WHERE {kosul} {zf('x.talep_zamani')} GROUP BY 1"""), p).all()), {})
    ayrinti["portfolyo"] = _guvenli(db, lambda: dict(db.execute(text(
        f"""SELECT count(*) AS toplam, count(*) FILTER (WHERE x.dogrulandi) AS dogrulanan FROM portfolyo_kayitlari x
            JOIN ogrenciler o ON o.id = x.ogrenci_id WHERE {kosul} {zf('x.olusturulma_zamani')}"""), p).mappings().one()), None)
    ayrinti["tercih_listesi"] = _guvenli(db, lambda: dict(db.execute(text(
        f"""SELECT x.durum, count(*) FROM tercih_listesi x JOIN ogrenciler o ON o.id = x.ogrenci_id WHERE {kosul} GROUP BY 1"""), p).all()), {})

    # Erken uyarı (güncel): rehberlik modülü açık okullarda kural motoru (okul başına bir çağrı, öğrenci başına değil)
    def _erken_uyari():
        from app.api.rehberlik import KURALLAR, uyarilari_hesapla
        from app.core.paketler import okul_modulleri
        okul_q = "SELECT id FROM okullar WHERE aktif_mi" + (" AND id = :okul" if okul_id else "")
        if okul_id == 0:
            return None
        seviye, kural = Counter(), Counter()
        for (oid,) in db.execute(text(okul_q), {"okul": okul_id} if okul_id else {}).all():
            if "rehberlik" not in okul_modulleri(db, oid):
                continue
            for s in uyarilari_hesapla(db, oid):
                seviye[s["seviye"]] += 1
                for u in s["uyarilar"]:
                    kural[KURALLAR.get(u["kural"], u["kural"])] += 1
        return {"seviye": [{"ad": a, "deger": seviye[k]} for k, a in (("yuksek", "Yüksek"), ("orta", "Orta"), ("dusuk", "Düşük"))],
                "kural": [{"ad": a, "deger": v} for a, v in kural.most_common()]}
    erken_uyari = _guvenli(db, _erken_uyari, None)

    # --- 7) Okullar ---
    okul_sat = db.execute(text(
        "SELECT o.id, o.ad, o.aktif_mi, o.paket, p.ad AS paket_ad FROM okullar o LEFT JOIN paketler p ON p.kod = o.paket ORDER BY o.ad")).mappings().all()
    yetkili = dict(db.execute(text(
        "SELECT okul_id, count(*) FROM admin_kullanicilar WHERE rol = 'okul_yetkilisi' AND aktif_mi"
        + ("" if test else " AND NOT coalesce(test_hesabi, false)") + " GROUP BY okul_id")).all())
    okul_turlari = defaultdict(lambda: {"guven": [], "tur": 0, "gecersiz": 0})
    for t in turlar:
        x = okul_turlari[t.okul_id or 0]
        x["tur"] += 1
        x["gecersiz"] += not t.sonuc_gecerli_mi
        if t.guven_skoru is not None:
            x["guven"].append(t.guven_skoru)
    grup = defaultdict(list)
    for r in ogr:
        grup[r["okul_id"] or 0].append(r)
    okul_ad = {o["id"]: o for o in okul_sat}
    okullar = []
    kimlikler = list(grup) if okul_id is None else [okul_id]
    if okul_id is None:
        kimlikler += [o["id"] for o in okul_sat if o["id"] not in grup and o["aktif_mi"]]
    for oid in kimlikler:
        l = grup.get(oid, [])
        k = len(l)
        o = okul_ad.get(oid)
        ok_tamam = sum(1 for r in l if r["tamamladi"])
        kat = {kk: _ort(okul_katman[oid].get(kk, [])) for kk in katman_kodlari}
        n_kat = max((len(okul_katman[oid].get(kk, [])) for kk in katman_kodlari), default=0)
        ot = okul_turlari[oid]
        satir = {
            "id": oid, "ad": o["ad"] if o else "Okul harici öğrenciler", "paket": (o["paket_ad"] or "Paket yok (tümü)") if o else "—",
            "aktif": bool(o["aktif_mi"]) if o else True,
            "ogrenci": k, "tamamlayan": ok_tamam, "yetkili": yetkili.get(oid, 0) if oid else 0,
            "giris": gizle(_oran(sum(r["giris"] for r in l), k), k),
            "aktif7": gizle(_oran(sum(aktif(r, 7) for r in l), k), k),
            "aktif30": gizle(_oran(sum(aktif(r, 30) for r in l), k), k),
            "tamamlama": gizle(_oran(ok_tamam, k), k),
            "hedef": gizle(_oran(sum(r["hedef"] for r in l), k), k),
            "guven": gizle(_ort(ot["guven"]), k) if yeterli(len(ot["guven"])) else None,
            "gecersiz": gizle(_oran(ot["gecersiz"], ot["tur"]), k) if yeterli(ot["tur"]) else None,
            "katman": {kk: (v if yeterli(n_kat) and yeterli(k) else None) for kk, v in kat.items()},
            "gizli": not yeterli(k),
        }
        okullar.append(satir)
    okullar.sort(key=lambda x: (-x["ogrenci"], x["ad"]))

    paket = Counter()
    paket_ogr = Counter()
    for o in okul_sat:
        if not o["aktif_mi"] or (okul_id is not None and o["id"] != okul_id):
            continue
        a = o["paket_ad"] or "Paket yok (tümü)"
        paket[a] += 1
        paket_ogr[a] += len(grup.get(o["id"], []))

    # --- 8) Sınıf düzeyine göre durum ---
    sinif = defaultdict(Counter)
    for r in ogr:
        d = "Tamamladı" if r["tamamladi"] else "Devam ediyor" if r["basladi"] else "Başlamadı" if r["giris"] else "Giriş yapmadı"
        sinif[_sinif(r["sinif"])][d] += 1
    sinif_sirali = sorted(sinif, key=lambda s: (SINIF_SIRASI.index(s) if s in SINIF_SIRASI else 99, s))

    tamamlayan = sum(1 for r in ogr if r["tamamladi"])
    okul_aktif = [o for o in okul_sat if o["aktif_mi"] and (okul_id is None or o["id"] == okul_id)]
    kpi = {
        "okul": len(okul_aktif) if okul_id != 0 else 0,
        "okul_paketli": sum(1 for o in okul_aktif if o["paket"]),
        "ogrenci": n,
        "aktif7": sum(aktif(r, 7) for r in ogr),
        "aktif30": sum(aktif(r, 30) for r in ogr),
        "tamamlayan": tamamlayan,
        "tamamlama_orani": _oran(tamamlayan, n),
        "guven_ort": _ort(guvenler),
        "tur": len(turlar),
        "gecersiz": gecersiz,
        "gecersiz_orani": _oran(gecersiz, len(turlar)),
        "hedef": sum(1 for r in ogr if r["hedef"]),
        "yetkili": sum(v for kk, v in yetkili.items() if okul_id is None or kk == okul_id) if okul_id != 0 else 0,
    }
    huni = [
        {"ad": "Hesap açıldı", "deger": n},
        {"ad": "Giriş yaptı", "deger": sum(1 for r in ogr if r["giris"])},
        {"ad": "Teste başladı", "deger": sum(1 for r in ogr if r["basladi"] or r["tamamladi"])},
        {"ad": "K1–K4 tamamladı", "deger": tamamlayan},
        {"ad": "K5 tamamladı", "deger": sum(1 for r in ogr if r["k5"])},
        {"ad": "Hedef seçti", "deger": kpi["hedef"]},
        {"ad": "Koçlukta ≥1 adım", "deger": sum(1 for r in ogr if r["kocluk"])},
    ]

    sonuc = {
        "filtre": {"donem": donem, "donem_ad": DONEMLER[donem], "baslangic": bas_gun.isoformat() if bas_gun else None,
                   "okul_id": okul_id, "test": test},
        "okul_secenekleri": [{"id": o["id"], "ad": o["ad"], "aktif": o["aktif_mi"]} for o in okul_sat],
        "kpi": kpi,
        "zaman": zaman,
        "huni": huni,
        "katmanlar": [{"kod": kk, "ortalama": _ort(katman_degerleri.get(kk, [])), "n": len(katman_degerleri.get(kk, [])),
                       "degerler": katman_degerleri.get(kk, [])} for kk in katman_kodlari],
        "guven_degerleri": guvenler,
        "onerilen": onerilen[:15],
        "hedeflenen": hedeflenen[:15],
        "alanlar": [{"ad": a, "deger": v} for a, v in alanlar.most_common()],
        "siniflar": {"kategoriler": sinif_sirali, "seriler": [{"ad": d, "degerler": [sinif[s][d] for s in sinif_sirali]} for d in DURUM_SERILERI]},
        "moduller": moduller,
        "ayrinti": ayrinti,
        "erken_uyari": erken_uyari,
        "okullar": okullar,
        "paketler": [{"ad": a, "deger": v, "ogrenci": paket_ogr[a]} for a, v in paket.most_common()],
        "kucuk_grup": {"esik": EN_AZ_GRUP, "metin": GIZLI_METIN, "dipnot": DIPNOT, "sayfa_gizli": False},
    }
    # Tek bir küçük okul seçildiyse: sayılar kalır, oran / ortalama / dağılımlar gizlenir
    if okul_id is not None and not yeterli(n):
        sonuc["kucuk_grup"]["sayfa_gizli"] = True
        for k in ("tamamlama_orani", "guven_ort", "gecersiz_orani"):
            sonuc["kpi"][k] = None
        for kat in sonuc["katmanlar"]:
            kat["ortalama"], kat["degerler"] = None, []
        sonuc["guven_degerleri"] = []
        for k in ("onerilen", "hedeflenen", "alanlar"):
            sonuc[k] = []
        sonuc["siniflar"] = {"kategoriler": [], "seriler": []}
        sonuc["ayrinti"]["simulasyon_meslekler"] = []
        if sonuc["ayrinti"].get("simulasyon"):
            sonuc["ayrinti"]["simulasyon"]["keyif"] = None
        sonuc["erken_uyari"] = None
    return sonuc


def _okul_param(okul_id: str | None) -> int | None:
    if okul_id in (None, "", "tum"):
        return None
    try:
        return int(okul_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="Geçersiz okul.")


@router.get("/istatistik/genel")
def genel_istatistik(donem: str = Query("30"), okul_id: str | None = Query(None), test: bool = Query(False),
                     db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_super_admin)):
    return istatistik_verisi(db, donem, _okul_param(okul_id), test)


# ----------------------------------------------------------------------------- Excel
def _xlsx_sayfa(wb, ad: str, basliklar: list[str], satirlar: list[list], not_metni: str | None = None):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet(ad[:31])
    for j, b in enumerate(basliklar, 1):
        c = ws.cell(row=1, column=j, value=b)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5D8A")
    for i, s in enumerate(satirlar, 2):
        for j, v in enumerate(s, 1):
            ws.cell(row=i, column=j, value=v)
    if not_metni:
        ws.cell(row=len(satirlar) + 3, column=1, value=not_metni).font = Font(italic=True, color="7A5C00")
    for j in range(1, len(basliklar) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 34 if j == 1 else 16
    ws.freeze_panes = "A2"
    return ws


def istatistik_xlsx(v: dict) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    f, k = v["filtre"], v["kpi"]
    okul_adi = next((o["ad"] for o in v["okul_secenekleri"] if o["id"] == f["okul_id"]), None) if f["okul_id"] else None
    gz = lambda x: GIZLI_METIN if x is None and v["kucuk_grup"]["sayfa_gizli"] else x  # noqa: E731
    kapsam = f"Dönem: {f['donem_ad']}{' (' + f['baslangic'] + ' itibarıyla)' if f['baslangic'] else ''} · Okul: " \
             f"{okul_adi or ('Okul harici' if f['okul_id'] == 0 else 'Tümü')} · Test hesapları: {'dahil' if f['test'] else 'hariç'}"
    _xlsx_sayfa(wb, "Özet", ["Gösterge", "Değer"], [
        ["Aktif okul", k["okul"]], ["Paket atanmış okul", k["okul_paketli"]], ["Öğrenci", k["ogrenci"]],
        ["Son 7 günde aktif öğrenci", k["aktif7"]], ["Son 30 günde aktif öğrenci", k["aktif30"]],
        ["Testi tamamlayan (en az bir tur)", k["tamamlayan"]], ["Tamamlama oranı (%)", gz(k["tamamlama_orani"])],
        ["Dönemde tamamlanan tur", k["tur"]], ["Ortalama güven puanı", gz(k["guven_ort"])],
        ["Geçersiz tur", k["gecersiz"]], ["Geçersiz tur oranı (%)", gz(k["gecersiz_orani"])],
        ["Hedef bölüm seçen", k["hedef"]], ["Okul yetkilisi", k["yetkili"]],
    ], f"Filizyol · {date.today().strftime('%d.%m.%Y')} · {kapsam}. {DIPNOT}")
    z = v["zaman"]
    _xlsx_sayfa(wb, "Zaman serisi", [f"Dönem başı ({z['birim']})", "Yeni hesap", "Giriş yapan öğrenci", "Tamamlanan tur"],
                [[x, z["yeni_hesap"][i], z["giris"][i], z["tur"][i]] for i, x in enumerate(z["x"])])
    _xlsx_sayfa(wb, "Huni", ["Adım", "Öğrenci", "İlk adıma göre (%)"],
                [[h["ad"], h["deger"], _oran(h["deger"], v["huni"][0]["deger"])] for h in v["huni"]], "Güncel durum (dönemden bağımsız).")
    _xlsx_sayfa(wb, "Katman puanları", ["Katman", "Ortalama", "Öğrenci"],
                [[x["kod"], gz(x["ortalama"]), x["n"]] for x in v["katmanlar"]],
                "Öğrencinin dönem içindeki son tamamlanmış turu; katman puanı = katmandaki özellik puanlarının ortalaması (0–100).")
    _xlsx_sayfa(wb, "Önerilen bölümler", ["Bölüm (1. sıra)", "Öğrenci"], [[x["ad"], x["deger"]] for x in v["onerilen"]])
    _xlsx_sayfa(wb, "Hedeflenen bölümler", ["Bölüm", "Hedefleyen", "Yetiyor (70+)", "Sınırda (40–69)", "Yetmiyor (<40)", "Hesaplanmadı"],
                [[x["ad"], x["deger"], x["yetiyor"], x["sinirda"], x["yetmiyor"], x["yok"]] for x in v["hedeflenen"]])
    _xlsx_sayfa(wb, "Alanlar", ["Üst alan (1. öneri)", "Öğrenci"], [[x["ad"], x["deger"]] for x in v["alanlar"]])
    s = v["siniflar"]
    _xlsx_sayfa(wb, "Sınıf düzeyi", ["Sınıf düzeyi"] + [x["ad"] for x in s["seriler"]],
                [[kat] + [x["degerler"][i] for x in s["seriler"]] for i, kat in enumerate(s["kategoriler"])])
    _xlsx_sayfa(wb, "Modül kullanımı", ["Modül", "Öğrenci", "Olay sayısı", "Birim"],
                [[m["ad"], m["ogrenci"], m["sayi"], m["olay"] + ("" if m["donemli"] else " (güncel)")] for m in v["moduller"]])
    if v["ayrinti"].get("simulasyon_meslekler"):
        _xlsx_sayfa(wb, "Simüle edilen meslekler", ["Meslek", "Simülasyon", "Ortalama keyif"],
                    [[x["ad"], x["deger"], x["keyif"]] for x in v["ayrinti"]["simulasyon_meslekler"]])
    if v.get("erken_uyari"):
        _xlsx_sayfa(wb, "Erken uyarı", ["Seviye / kural", "Öğrenci"],
                    [[x["ad"] + " seviye", x["deger"]] for x in v["erken_uyari"]["seviye"]] + [[x["ad"], x["deger"]] for x in v["erken_uyari"]["kural"]])
    kk = [x["kod"] for x in v["katmanlar"]]
    _xlsx_sayfa(wb, "Okullar", ["Okul", "Paket", "Öğrenci", "Tamamlayan", "Okul yetkilisi", "Giriş yapan (%)", "Son 7 gün aktif (%)",
                                "Son 30 gün aktif (%)", "Tamamlama (%)", "Hedef seçen (%)", "Ort. güven", "Geçersiz tur (%)"] + [f"{x} ort." for x in kk],
                [[o["ad"], o["paket"], o["ogrenci"], o["tamamlayan"], o["yetkili"]]
                 + [GIZLI_METIN if o["gizli"] and o[a] is None else o[a] for a in ("giris", "aktif7", "aktif30", "tamamlama", "hedef", "guven", "gecersiz")]
                 + [GIZLI_METIN if o["gizli"] else o["katman"].get(x) for x in kk] for o in v["okullar"]], DIPNOT)
    _xlsx_sayfa(wb, "Paketler", ["Paket", "Aktif okul", "Öğrenci"], [[x["ad"], x["deger"], x["ogrenci"]] for x in v["paketler"]])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@router.get("/istatistik/genel/excel")
def genel_istatistik_excel(donem: str = Query("30"), okul_id: str | None = Query(None), test: bool = Query(False),
                           db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_super_admin)):
    oid = _okul_param(okul_id)
    v = istatistik_verisi(db, donem, oid, test)
    icerik = istatistik_xlsx(v)
    denetim_yaz(db, yon, "istatistik_excel", "istatistik", oid if oid is not None else "tum",
                f"{v['filtre']['donem_ad']}, test {'dahil' if test else 'hariç'}")
    db.commit()
    return Response(icerik, media_type=XLSX, headers={
        "Content-Disposition": f'attachment; filename="filizyol_istatistik_{date.today().isoformat()}.xlsx"', "Cache-Control": "no-store"})


@router.get("/audit-log/excel")
def audit_log_excel(gun: int = Query(90, ge=0, le=3650), db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_super_admin)):
    """Audit Log dışa aktarımı (gun=0 → tümü, en çok 50.000 kayıt)."""
    import openpyxl
    satir = db.execute(text(f"""
        SELECT a.zaman, coalesce(k.ad_soyad, a.yapan_ad, 'silinmiş yönetici') AS yapan, k.rol, a.islem, a.hedef_tablo, a.hedef_id,
               a.gerekce, ok.ad AS okul
          FROM audit_log a LEFT JOIN admin_kullanicilar k ON k.id = a.admin_id LEFT JOIN okullar ok ON ok.id = a.okul_id
         {"WHERE a.zaman >= now() - make_interval(days => :g)" if gun else ""}
         ORDER BY a.zaman DESC LIMIT 50000"""), {"g": gun}).all()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    rol = {"super_admin": "Süper admin", "okul_yetkilisi": "Okul yetkilisi"}
    _xlsx_sayfa(wb, "Audit Log", ["Zaman", "Yapan", "Rol", "İşlem", "Hedef tablo", "Hedef", "Gerekçe / ayrıntı", "Okul"],
                [[r.zaman.astimezone(TZ).strftime("%d.%m.%Y %H:%M") if r.zaman else None, r.yapan, rol.get(r.rol, r.rol), r.islem,
                  r.hedef_tablo, r.hedef_id, r.gerekce, r.okul] for r in satir],
                f"Filizyol · {date.today().strftime('%d.%m.%Y')} · {'Son ' + str(gun) + ' gün' if gun else 'Tüm kayıtlar'} · {len(satir)} kayıt")
    b = io.BytesIO()
    wb.save(b)
    denetim_yaz(db, yon, "audit_log_excel", "audit_log", gun or "tum", f"{len(satir)} kayıt")
    db.commit()
    return Response(b.getvalue(), media_type=XLSX, headers={
        "Content-Disposition": f'attachment; filename="audit_log_{date.today().isoformat()}.xlsx"', "Cache-Control": "no-store"})


# ----------------------------------------------------------------------------- Kontrol Paneli
@router.get("/istatistik/dikkat")
def dikkat_gerektirenler(db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_super_admin)):
    ozet = db.execute(text("""
        SELECT count(*) AS ogrenci,
               count(*) FILTER (WHERE o.son_giris_zamani >= now() - interval '7 days') AS aktif7,
               count(*) FILTER (WHERE EXISTS (SELECT 1 FROM ogrenci_degerlendirme_turu t WHERE t.ogrenci_id = o.id AND t.durum = 'tamamlandi')) AS tamamlayan
          FROM ogrenciler o WHERE NOT coalesce(o.test_hesabi, false)""")).mappings().one()
    okul = db.execute(text("SELECT count(*) FILTER (WHERE aktif_mi) AS aktif, count(*) AS toplam FROM okullar")).mappings().one()
    s = lambda sql: _guvenli(db, lambda: db.execute(text(sql)).scalar() or 0, 0)  # noqa: E731
    maddeler = [
        {"k": "pipeline", "seviye": "uyari", "ikon": "🧪", "baslik": "Onay bekleyen pipeline taslağı", "link": "/admin/pipeline",
         "sayi": s("SELECT count(DISTINCT yukleme_grubu) FROM bolum_agirliklari_taslak WHERE durum = 'bekliyor'"),
         "aciklama": "Yüklenen bölüm ağırlıkları onaylanana kadar skor motorunda kullanılmaz."},
        {"k": "gecersiz", "seviye": "uyari", "ikon": "🛡️", "baslik": "Son 30 günde geçersiz sayılan tur", "link": "/admin/guvenlik",
         "sayi": s("""SELECT count(*) FROM ogrenci_degerlendirme_turu t JOIN ogrenciler o ON o.id = t.ogrenci_id
                       WHERE t.durum = 'tamamlandi' AND NOT t.sonuc_gecerli_mi AND t.tamamlanma_zamani >= now() - interval '30 days'
                         AND NOT coalesce(o.test_hesabi, false)"""),
         "aciklama": "Güven puanı eşiğin altında kalan değerlendirmeler; öğrenciye yeniden değerlendirme önerilir."},
        {"k": "yarim", "seviye": "bilgi", "ikon": "⏸️", "baslik": "14 günden uzun süredir yarım kalan değerlendirme", "link": "/admin/karsilastirma",
         "sayi": s("""SELECT count(*) FROM ogrenci_degerlendirme_turu t JOIN ogrenciler o ON o.id = t.ogrenci_id
                       WHERE t.durum = 'devam_ediyor' AND t.baslama_zamani < now() - interval '14 days' AND NOT coalesce(o.test_hesabi, false)"""),
         "aciklama": "Okulların rehberlik ekibi erken uyarı listesinde görür."},
        {"k": "giris_yok", "seviye": "bilgi", "ikon": "🔑", "baslik": "Hesabı 14+ gün önce açılıp hiç giriş yapmayan öğrenci", "link": "/admin/karsilastirma",
         "sayi": s("""SELECT count(*) FROM ogrenciler o WHERE o.son_giris_zamani IS NULL AND o.olusturulma_zamani < now() - interval '14 days'
                       AND NOT coalesce(o.test_hesabi, false)"""),
         "aciklama": "Giriş listesinin öğrencilere ulaştırılıp ulaştırılmadığını okulla kontrol edin."},
        {"k": "yetkilisiz", "seviye": "uyari", "ikon": "🏫", "baslik": "Okul yetkilisi tanımlanmamış aktif okul", "link": "/admin/okullar",
         "sayi": s("""SELECT count(*) FROM okullar ok WHERE ok.aktif_mi AND NOT EXISTS (
                       SELECT 1 FROM admin_kullanicilar a WHERE a.okul_id = ok.id AND a.rol = 'okul_yetkilisi' AND a.aktif_mi)"""),
         "aciklama": "Okulun paneline girecek kimse yok."},
        {"k": "paketsiz", "seviye": "bilgi", "ikon": "📦", "baslik": "Paket atanmamış aktif okul (tüm modüller açık)", "link": "/admin/okullar",
         "sayi": s("SELECT count(*) FROM okullar WHERE aktif_mi AND paket IS NULL"),
         "aciklama": "Paket atanmayan okulda bütün modüller açıktır."},
        {"k": "yetkili_sifre", "seviye": "bilgi", "ikon": "🔐", "baslik": "Henüz ilk girişini yapmamış okul yetkilisi", "link": "/admin/okullar",
         "sayi": s("""SELECT count(*) FROM admin_kullanicilar WHERE rol = 'okul_yetkilisi' AND aktif_mi AND son_giris_zamani IS NULL
                       AND NOT coalesce(test_hesabi, false)"""),
         "aciklama": "Geçici şifre iletilmemiş olabilir."},
        {"k": "koc", "seviye": "bilgi", "ikon": "👩‍🏫", "baslik": "Yanıt bekleyen eğitim koçu talebi", "link": "/admin/koclar",
         "sayi": s("SELECT count(*) FROM koc_gorusme_talepleri WHERE durum = 'beklemede'"),
         "aciklama": "Öğrenci ya da velinin açtığı görüşme talepleri."},
        {"k": "test_suresi", "seviye": "bilgi", "ikon": "🧪", "baslik": "Giriş bağlantısının süresi dolmuş test hesabı", "link": "/admin/test-hesaplari",
         "sayi": s("""SELECT (SELECT count(*) FROM ogrenciler WHERE test_hesabi AND test_giris_bitis < now())
                           + (SELECT count(*) FROM admin_kullanicilar WHERE test_hesabi AND test_giris_bitis < now())"""),
         "aciklama": "Kullanılmayacaksa silin ya da yeni bağlantı üretin."},
        {"k": "bolum_taslak", "seviye": "bilgi", "ikon": "📚", "baslik": "Yayında olmayan bölüm (taslak / test ediliyor)", "link": "/admin/bolumler",
         "sayi": s("SELECT count(*) FROM bolumler WHERE durum <> 'yayinda'"),
         "aciklama": "Öğrencilere önerilmez."},
    ]
    return {
        "ozet": {"okul_aktif": okul["aktif"], "okul_toplam": okul["toplam"], "ogrenci": ozet["ogrenci"], "aktif7": ozet["aktif7"],
                 "tamamlayan": ozet["tamamlayan"], "tamamlama_orani": _oran(ozet["tamamlayan"], ozet["ogrenci"])},
        "maddeler": [m for m in maddeler if m["sayi"]],
    }
