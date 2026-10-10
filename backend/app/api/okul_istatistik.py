# -*- coding: utf-8 -*-
"""
[2026-10-10] Okul paneli → İstatistikler (okul yetkilisi + süper admin).

GET /yonetim/okul/{okul_id}/istatistik?sinif=&sube=&bas=&bit=        — tüm okul istatistikleri tek istekte
GET /yonetim/okul/{okul_id}/istatistik/excel?sinif=&sube=&bas=&bit=  — aynı veriler Excel (sayfa sayfa)

Kapsam: okul yetkilisi yalnızca kendi okulunu, süper admin her okulu görür (okul_yonetimi._okul_kapsami).
Filtreler: sınıf düzeyi / şube öğrenci kümesini, tarih aralığı zaman içindeki etkinlikleri (girişler, tamamlamalar,
görevler, görüşmeler, denemeler, simülasyonlar…) süzer. Durum sayıları (tamamlayan, hedef seçen…) güncel durumdur.
Test hesapları sayılmaz. Okul kendi öğrencilerini isimle gördüğü için bu ekranda küçük grup gizleme UYGULANMAZ
(PDF/Excel toplu raporlarındaki gizleme app/core/kucuk_grup.py ile ayrıdır). Öğrenci tarafına hiçbir veri açılmaz.
Paket modülü kapalı bölümler hesaplanmaz ve dönmez.
Performans: sabit sayıda toplu SQL; öğrenci başına sorgu yok.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Query
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.api.deps import get_mevcut_yonetim
from app.api.okul_yonetimi import SINIFLAR, _durum, _ilerleme, _okul_kapsami, sube_etiketi
from app.core.database import get_db
from app.core.hesap_yonetimi import denetim_yaz
from app.core.paketler import okul_modulleri
from app.models import AdminKullanici, Ogrenci

router = APIRouter(prefix="/yonetim", tags=["Okul istatistikleri"])
TR = ZoneInfo("Europe/Istanbul")
VARSAYILAN_GUN = 90
EN_FAZLA_GUN = 730
DURUM_ADLARI = [("tamamlandi", "Tamamladı"), ("devam", "Devam ediyor"), ("baslamadi", "Başlamadı"), ("giris_yok", "Giriş yok")]


def _baslik(metin: str | None) -> str:
    try:
        from app.core.koclugu_servisi import turkce_baslik
        return turkce_baslik(metin or "")
    except Exception:
        return metin or ""


def _ort(liste) -> float | None:
    l = [float(x) for x in liste if x is not None]
    return round(sum(l) / len(l), 1) if l else None


def _yuzde(a, b) -> float | None:
    return round(100 * a / b, 1) if b else None


def _aralik(bas: date | None, bit: date | None) -> tuple[date, date]:
    bugun = datetime.now(TR).date()
    bit = min(bit or bugun, bugun)
    bas = bas or (bit - timedelta(days=VARSAYILAN_GUN - 1))
    if bas > bit:
        raise HTTPException(400, "Başlangıç tarihi bitişten sonra olamaz.")
    if (bit - bas).days > EN_FAZLA_GUN:
        raise HTTPException(400, f"Tarih aralığı en fazla {EN_FAZLA_GUN} gün olabilir.")
    return bas, bit


def _haftalar(bas: date, bit: date) -> list[date]:
    h = bas - timedelta(days=bas.weekday())
    son = bit - timedelta(days=bit.weekday())
    l = []
    while h <= son:
        l.append(h)
        h += timedelta(days=7)
    return l


def _sinif_sira(s: str | None) -> int:
    return SINIFLAR.index(s) if s in SINIFLAR else 98 if s else 99


# ============================================================================= ortak yardımcılar
# (İstatistikler ve Rapor Merkezi → Tek Bakışta [okul_tek_bakista.py] birlikte kullanır)
def ogrenci_kumesi(db: Session, okul_id: int, sinif: str | None = None, sube: str | None = None) -> dict:
    """Okulun öğrencileri (test hesapları hariç) ve sınıf / şube filtresi.
    → {tum, ogr (filtreli), sinif, sube, etiket, secenekler (filtreden bağımsız)}"""
    sinif = (sinif or "").strip() or None
    sube = ((sube or "").strip().upper() or None) if sinif else None
    tum = [o for o in db.query(Ogrenci).filter(Ogrenci.okul_id.is_(None) if okul_id == 0 else Ogrenci.okul_id == okul_id).all()
           if not getattr(o, "test_hesabi", False)]
    secenek = defaultdict(set)
    for o in tum:
        if o.sinif:
            secenek[o.sinif]
            if o.sube:
                secenek[o.sinif].add(o.sube)
    ogr = [o for o in tum if not sinif or (o.sinif == sinif and (not sube or (o.sube or "") == sube))]
    return {"tum": tum, "ogr": ogr, "sinif": sinif, "sube": sube,
            "etiket": sube_etiketi(sinif, sube) if sinif else "Tüm okul",
            "secenekler": [{"sinif": s, "subeler": sorted(b)} for s, b in sorted(secenek.items(), key=lambda kv: _sinif_sira(kv[0]))]}


def kocluk_etkinligi(db: Session, ids: list, bas: date, bit: date) -> tuple[set, set]:
    """(hiç gelişim adımı tamamlayan, aralıkta adım ya da haftalık görev tamamlayan) öğrenci kümeleri."""
    if not ids:
        return set(), set()
    p = {"ids": ids, "bas": bas, "bit": bit}
    TRG = "(%s AT TIME ZONE 'Europe/Istanbul')::date"
    adim_ogr, aralikta = set(), set()
    for r in db.execute(text(f"""
        SELECT ogrenci_id, count(*) FILTER (WHERE {TRG % 'guncelleme_zamani'} BETWEEN :bas AND :bit) AS ar
          FROM ogrenci_gelisim_adim_durumu WHERE ogrenci_id = ANY(:ids) AND durum = 'tamamlandi' GROUP BY 1"""), p).all():
        adim_ogr.add(r.ogrenci_id)
        if r.ar:
            aralikta.add(r.ogrenci_id)
    aralikta |= {r[0] for r in db.execute(text(f"""
        SELECT DISTINCT ogrenci_id FROM ogrenci_haftalik_gorev WHERE ogrenci_id = ANY(:ids) AND durum = 'tamamlandi'
           AND {TRG % 'tamamlanma_zamani'} BETWEEN :bas AND :bit"""), p).all()}
    return adim_ogr, aralikta


# ============================================================================= veri
def okul_istatistik_verisi(db: Session, okul_id: int, sinif: str | None = None, sube: str | None = None,
                           bas: date | None = None, bit: date | None = None) -> dict:
    bas, bit = _aralik(bas, bit)
    moduller = set(okul_modulleri(db, okul_id or None))
    kume = ogrenci_kumesi(db, okul_id, sinif, sube)
    sinif, sube, ogr = kume["sinif"], kume["sube"], kume["ogr"]
    ids = [o.id for o in ogr]
    p = {"ids": ids, "bas": bas, "bit": bit}
    TRG = "(%s AT TIME ZONE 'Europe/Istanbul')::date"   # zaman damgası → İstanbul günü
    simdi = datetime.now(timezone.utc)

    il_tum = _ilerleme(db, okul_id)
    il = {i: il_tum.get(i) or {} for i in ids}
    durum = {o.id: _durum(o, il[o.id])[0] for o in ogr}
    tamam_tur = {i: x["tur_id"] for i, x in il.items() if x.get("durum") == "tamamlandi" and x.get("tur_id")}
    son_tur = {i: x["tur_id"] for i, x in il.items() if x.get("tur_id")}
    tur_ids = list(tamam_tur.values())
    adlar = {r.id: _baslik(r.ad) for r in db.execute(text("SELECT id, ad FROM bolumler")).all()}

    # ------------------------------------------------------------------ değerlendirme turları (güven, geçersiz)
    turlar = db.execute(text("""
        SELECT t.id, t.ogrenci_id, t.durum, t.guven_skoru, t.sonuc_gecerli_mi
          FROM ogrenci_degerlendirme_turu t WHERE t.ogrenci_id = ANY(:ids)"""), p).all() if ids else []
    guvenler = [float(t.guven_skoru) for t in turlar if t.id in set(tur_ids) and t.guven_skoru is not None]
    gecersiz = sum(1 for t in turlar if t.durum == "tamamlandi" and t.sonuc_gecerli_mi is False)

    # ------------------------------------------------------------------ K5 (alan derinleşme)
    k5_ogr: dict = {}
    k5_alan = Counter()
    if tur_ids:
        for r in db.execute(text("""
            SELECT d.ogrenci_id, bool_and(d.durum = 'tamamlandi') AS hepsi,
                   array_agg(dl.ad) FILTER (WHERE d.durum = 'tamamlandi') AS alanlar
              FROM ogrenci_dal_oturumlari d JOIN dallar dl ON dl.id = d.dal_id
             WHERE d.tur_id = ANY(:t) GROUP BY d.ogrenci_id"""), {"t": tur_ids}).all():
            k5_ogr[r.ogrenci_id] = bool(r.hepsi)
            for a in r.alanlar or []:
                k5_alan[a] += 1
    k5_tamam = sum(1 for v in k5_ogr.values() if v)

    # ------------------------------------------------------------------ katman (K1–K4) adımları — huni
    katman_tamam = Counter()
    if son_tur:
        for kod, n in db.execute(text("""
            SELECT k.kod, count(DISTINCT ko.ogrenci_id) FROM ogrenci_katman_oturumlari ko JOIN katmanlar k ON k.id = ko.katman_id
             WHERE ko.tur_id = ANY(:t) AND ko.durum = 'tamamlandi' AND NOT k.kosullu_mu GROUP BY k.kod"""),
                {"t": list(son_tur.values())}).all():
            katman_tamam[kod] = n
    katmanlar = [(r.kod, r.ad) for r in db.execute(text("SELECT kod, ad FROM katmanlar WHERE NOT kosullu_mu ORDER BY sira")).all()]

    # ------------------------------------------------------------------ katılım trendi (haftalık)
    haftalar = _haftalar(bas, bit)
    giris_h, tamam_h = Counter(), Counter()
    if ids:
        for h, n in db.execute(text(f"""
            SELECT date_trunc('week', e.zaman AT TIME ZONE 'Europe/Istanbul')::date, count(DISTINCT e.ogrenci_id)
              FROM ogrenci_hesap_olaylari e WHERE e.ogrenci_id = ANY(:ids) AND e.olay = 'giris'
               AND {TRG % 'e.zaman'} BETWEEN :bas AND :bit GROUP BY 1"""), p).all():
            giris_h[h] = n
        for h, n in db.execute(text(f"""
            SELECT date_trunc('week', t.tamamlanma_zamani AT TIME ZONE 'Europe/Istanbul')::date, count(DISTINCT t.ogrenci_id)
              FROM ogrenci_degerlendirme_turu t WHERE t.ogrenci_id = ANY(:ids) AND t.durum = 'tamamlandi'
               AND {TRG % 't.tamamlanma_zamani'} BETWEEN :bas AND :bit GROUP BY 1"""), p).all():
            tamam_h[h] = n
    aralikta_giris = db.execute(text(f"""
        SELECT count(DISTINCT ogrenci_id) FROM ogrenci_hesap_olaylari WHERE ogrenci_id = ANY(:ids) AND olay = 'giris'
           AND {TRG % 'zaman'} BETWEEN :bas AND :bit"""), p).scalar() if ids else 0

    # ------------------------------------------------------------------ koçluk adımı (huni + KPI)
    kocluk_acik = "kocluk" in moduller
    adim_ogr, kocluk_aralik = kocluk_etkinligi(db, ids, bas, bit) if kocluk_acik else (set(), set())
    kocluk_aktif = len(kocluk_aralik)

    toplam = len(ogr)
    giris_yapan = sum(1 for o in ogr if o.son_giris_zamani is not None)
    aktif7 = sum(1 for o in ogr if o.son_giris_zamani and o.son_giris_zamani >= simdi - timedelta(days=7))
    tamamlayan = sum(1 for k in durum.values() if k == "tamamlandi")
    baslayan = sum(1 for k in durum.values() if k in ("tamamlandi", "devam"))
    hedef_secen = sum(1 for x in il.values() if x.get("hedef"))

    v: dict = {
        "moduller": sorted(moduller),
        "filtre": {"sinif": sinif, "sube": sube, "etiket": kume["etiket"],
                   "bas": bas.isoformat(), "bit": bit.isoformat(), "secenekler": kume["secenekler"]},
        "kpi": {
            "ogrenci": toplam, "giris_yapan": giris_yapan, "aktif_7": aktif7, "aralikta_giris": aralikta_giris,
            "teste_baslayan": baslayan, "tamamlayan": tamamlayan, "tamamlama_orani": _yuzde(tamamlayan, toplam),
            "k5_tamamlayan": k5_tamam, "hedef_secen": hedef_secen,
            "ort_guven": _ort(guvenler), "gecersiz_tur": gecersiz,
            "kocluk_aktif": kocluk_aktif if kocluk_acik else None,
        },
    }

    # ------------------------------------------------------------------ Katılım
    # teste başlayan öğrenci giriş yapmıştır (eski / içe aktarılmış hesaplarda son giriş zamanı boş olabilir)
    giris_huni = sum(1 for o in ogr if o.son_giris_zamani is not None or durum[o.id] in ("tamamlandi", "devam"))
    huni = [{"ad": "Hesabı açılan", "deger": toplam}, {"ad": "Giriş yaptı", "deger": giris_huni},
            {"ad": "Teste başladı", "deger": baslayan}]
    huni += [{"ad": f"{kod} tamamlandı", "deger": katman_tamam.get(kod, 0)} for kod, _ in katmanlar]
    huni += [{"ad": "K5 tamamlandı", "deger": k5_tamam}, {"ad": "Hedef bölüm seçti", "deger": hedef_secen}]
    if kocluk_acik:
        huni.append({"ad": "Koçluk adımı tamamladı", "deger": len(adim_ogr)})
    gruplar = defaultdict(Counter)
    for o in ogr:
        anahtar = (o.sinif or "Belirtilmedi", o.sube or "")
        gruplar[anahtar][durum[o.id]] += 1
    gruplar_sirali = sorted(gruplar.items(), key=lambda kv: (_sinif_sira(kv[0][0]), kv[0][1]))
    v["katilim"] = {
        "haftalik": [{"hafta": h.isoformat(), "giris": giris_h.get(h, 0), "tamamlama": tamam_h.get(h, 0)} for h in haftalar],
        "huni": huni,
        "subeler": [{"etiket": sube_etiketi(k[0], k[1]) if k[1] else k[0], "sinif": k[0], "sube": k[1], "toplam": sum(c.values()),
                     **{kod: c.get(kod, 0) for kod, _ in DURUM_ADLARI}} for k, c in gruplar_sirali],
        "durum_adlari": dict(DURUM_ADLARI),
    }

    # ------------------------------------------------------------------ Profil
    katman_ort, degisken_ort, sube_katman, ogr_katman = [], [], [], {}
    if tur_ids:
        sinif_of = {o.id: o.sinif or "Belirtilmedi" for o in ogr}
        sube_of = {o.id: (o.sinif or "Belirtilmedi", o.sube or "") for o in ogr}
        rows = db.execute(text("""
            SELECT s.ogrenci_id, k.kod, avg(s.puan) AS ort FROM ogrenci_degisken_skorlari s
              JOIN degiskenler d ON d.id = s.degisken_id JOIN katmanlar k ON k.id = d.katman_id
             WHERE s.tur_id = ANY(:t) AND d.dal_id IS NULL AND NOT k.kosullu_mu GROUP BY s.ogrenci_id, k.kod"""), {"t": tur_ids}).all()
        okul_k, sinif_k, sube_k = defaultdict(list), defaultdict(lambda: defaultdict(list)), defaultdict(lambda: defaultdict(list))
        for r in rows:
            puan = float(r.ort)
            okul_k[r.kod].append(puan)
            sinif_k[sinif_of.get(r.ogrenci_id)][r.kod].append(puan)
            sube_k[sube_of.get(r.ogrenci_id)][r.kod].append(puan)
            ogr_katman.setdefault(r.kod, []).append(round(puan, 1))
        kodlar = [k for k, _ in katmanlar]
        katman_ort = [{"ad": "Okul geneli" if not sinif else sube_etiketi(sinif, sube), "okul": True,
                       "degerler": {k: _ort(okul_k[k]) for k in kodlar}}]
        if not sinif or not sube:
            katman_ort += [{"ad": s, "okul": False, "ogrenci": len({i for i in tamam_tur if sinif_of.get(i) == s}),
                            "degerler": {k: _ort(d[k]) for k in kodlar}}
                           for s, d in sorted(sinif_k.items(), key=lambda kv: _sinif_sira(kv[0]))]
        sube_katman = [{"ad": sube_etiketi(s[0], s[1]) if s[1] else s[0], "degerler": {k: _ort(d[k]) for k in kodlar}}
                       for s, d in sorted(sube_k.items(), key=lambda kv: (_sinif_sira(kv[0][0]), kv[0][1]))]
        degisken_ort = [{"ad": r.ad, "katman": r.kod, "ortalama": round(float(r.ort), 1), "n": r.n} for r in db.execute(text("""
            SELECT d.ad, k.kod, avg(s.puan) AS ort, count(DISTINCT s.ogrenci_id) AS n FROM ogrenci_degisken_skorlari s
              JOIN degiskenler d ON d.id = s.degisken_id JOIN katmanlar k ON k.id = d.katman_id
             WHERE s.tur_id = ANY(:t) AND d.dal_id IS NULL AND NOT k.kosullu_mu
             GROUP BY d.ad, k.kod ORDER BY ort DESC"""), {"t": tur_ids}).all()]
    v["profil"] = {
        "katmanlar": [{"kod": k, "ad": a} for k, a in katmanlar],
        "tamamlayan": len(tur_ids),
        "katman_ort": katman_ort,
        "guclu": degisken_ort[:8],
        "zayif": list(reversed(degisken_ort))[:8],
        "sube_katman": sube_katman,
        "histogram": ogr_katman,   # {K1: [öğrenci başına katman ortalaması, …]}
    }

    # ------------------------------------------------------------------ Bölüm ve meslek
    ilk, hedef = Counter(), Counter()
    for i, x in il.items():
        if x.get("durum") == "tamamlandi" and x.get("ilk_bolum"):
            ilk[adlar.get(x["ilk_bolum"], "?")] += 1
        if x.get("hedef"):
            hedef[adlar.get(x["hedef"], "?")] += 1
    # hedefi, son tamamlanan turdaki ilk 10 önerisinde olan öğrenciler (sıralama: tercih.uyumlar ile aynı — toplam uyum)
    hedefli = {i: x["hedef"] for i, x in il.items() if x.get("hedef") and i in tamam_tur}
    ilk10 = 0
    if hedefli:
        ust = defaultdict(set)
        for r in db.execute(text("""
            SELECT ogrenci_id, bolum_id FROM (
                SELECT s.ogrenci_id, s.bolum_id, row_number() OVER (PARTITION BY s.tur_id ORDER BY s.toplam_uyum DESC) AS sira
                  FROM ogrenci_bolum_uyum_skorlari s WHERE s.tur_id = ANY(:t)) x WHERE sira <= 10"""),
                {"t": [tamam_tur[i] for i in hedefli]}).all():
            ust[r.ogrenci_id].add(r.bolum_id)
        ilk10 = sum(1 for i, b in hedefli.items() if b in ust[i])
    v["bolum"] = {
        "ilk_oneri": [{"ad": b, "deger": n} for b, n in ilk.most_common(15)],
        "hedef": [{"ad": b, "deger": n} for b, n in hedef.most_common(15)],
        "k5_alan": [{"ad": a, "deger": n} for a, n in k5_alan.most_common()],
        "uyum": {"hedefli_tamamlayan": len(hedefli), "ilk10": ilk10, "oran": _yuzde(ilk10, len(hedefli))},
        "simulasyon": None,
    }
    if kocluk_acik and ids:   # meslek simülasyonu Koçluk modülüyle gelir
        sim = db.execute(text(f"""
            SELECT meslek_ad, count(*) AS n, count(DISTINCT ogrenci_id) AS ogr, avg(keyif) AS keyif FROM simulasyon_sonuclari
             WHERE ogrenci_id = ANY(:ids) AND {TRG % 'olusturulma_zamani'} BETWEEN :bas AND :bit
             GROUP BY meslek_ad ORDER BY n DESC, meslek_ad LIMIT 15"""), p).all()
        v["bolum"]["simulasyon"] = [{"ad": r.meslek_ad, "deger": r.n, "ogrenci": r.ogr,
                                     "keyif": round(float(r.keyif), 1) if r.keyif is not None else None} for r in sim]

    # ------------------------------------------------------------------ Koçluk ve çalışma
    if kocluk_acik:
        gorev_rows = db.execute(text("""
            SELECT ogrenci_id, hafta_baslangic, count(*) AS top, count(*) FILTER (WHERE durum = 'tamamlandi') AS biten
              FROM ogrenci_haftalik_gorev WHERE ogrenci_id = ANY(:ids) GROUP BY 1, 2"""), p).all() if ids else []
        hafta = defaultdict(lambda: [0, 0, set()])
        ogr_haftalar = defaultdict(set)
        for r in gorev_rows:
            if r.biten:
                ogr_haftalar[r.ogrenci_id].add(r.hafta_baslangic)
            if bas - timedelta(days=6) <= r.hafta_baslangic <= bit:
                hafta[r.hafta_baslangic][0] += r.top
                hafta[r.hafta_baslangic][1] += r.biten
                hafta[r.hafta_baslangic][2].add(r.ogrenci_id)
        # seri: bu haftadan (ya da geçen haftadan) geriye, en az bir görev tamamlanan ardışık hafta sayısı
        bu_hafta = datetime.now(TR).date()
        bu_hafta -= timedelta(days=bu_hafta.weekday())
        seriler = []
        for i, hs in ogr_haftalar.items():
            h = bu_hafta if bu_hafta in hs else bu_hafta - timedelta(days=7)
            n = 0
            while h in hs:
                n += 1
                h -= timedelta(days=7)
            seriler.append(n)
        adim_say = db.execute(text(f"""
            SELECT count(*), count(DISTINCT ogrenci_id) FROM ogrenci_gelisim_adim_durumu
             WHERE ogrenci_id = ANY(:ids) AND durum = 'tamamlandi' AND {TRG % 'guncelleme_zamani'} BETWEEN :bas AND :bit"""), p).first() if ids else (0, 0)
        v["kocluk"] = {
            "adim": {"tamamlanan": adim_say[0] or 0, "ogrenci": adim_say[1] or 0, "hic_adim_yapan": len(adim_ogr)},
            "gorev_haftalik": [{"hafta": h.isoformat(), "toplam": hafta[h][0], "tamamlanan": hafta[h][1],
                                "oran": _yuzde(hafta[h][1], hafta[h][0]), "ogrenci": len(hafta[h][2])}
                               for h in haftalar],
            "gorev_toplam": {"toplam": sum(hafta[h][0] for h in haftalar), "tamamlanan": sum(hafta[h][1] for h in haftalar)},
            "seri": {"ortalama": _ort([s for s in seriler if s > 0]), "en_uzun": max(seriler, default=0),
                     "aktif_seri": sum(1 for s in seriler if s > 0), "dagilim": [
                         {"ad": ad, "deger": sum(1 for s in seriler if lo <= s <= hi)}
                         for ad, lo, hi in (("1 hafta", 1, 1), ("2–3 hafta", 2, 3), ("4–7 hafta", 4, 7), ("8+ hafta", 8, 10**6))]},
        }
    if "filiz" in moduller and ids:
        f = db.execute(text(f"""
            SELECT count(DISTINCT o.id) AS oturum, count(DISTINCT o.ogrenci_id) AS ogrenci,
                   count(m.id) FILTER (WHERE m.rol = 'ogrenci') AS mesaj
              FROM ogrenci_koclugu_oturumlari o LEFT JOIN ogrenci_koclugu_mesajlari m ON m.oturum_id = o.id
             WHERE o.ogrenci_id = ANY(:ids) AND {TRG % 'o.baslama_zamani'} BETWEEN :bas AND :bit"""), p).first()
        v["filiz"] = {"oturum": f.oturum or 0, "ogrenci": f.ogrenci or 0, "mesaj": f.mesaj or 0,
                      "oran": _yuzde(f.ogrenci or 0, toplam)}
    if "calisma" in moduller and ids:
        c = db.execute(text(f"""
            SELECT count(DISTINCT ogrenci_id) AS ogrenci, coalesce(sum(sure_dk), 0) AS dk, coalesce(sum(soru), 0) AS soru,
                   coalesce(sum(dogru), 0) AS dogru, coalesce(sum(yanlis), 0) AS yanlis
              FROM calisma_kayitlari WHERE ogrenci_id = ANY(:ids) AND tarih BETWEEN :bas AND :bit"""), p).first()
        dersler = db.execute(text("""
            SELECT ders, coalesce(sum(sure_dk), 0) AS dk, coalesce(sum(soru), 0) AS soru FROM calisma_kayitlari
             WHERE ogrenci_id = ANY(:ids) AND tarih BETWEEN :bas AND :bit GROUP BY ders ORDER BY dk DESC LIMIT 12"""), p).all()
        program = db.execute(text("SELECT count(DISTINCT ogrenci_id) FROM calisma_programi WHERE ogrenci_id = ANY(:ids)"), p).scalar()
        from app.api.calisma import DERSLER as CALISMA_DERSLERI
        v["calisma"] = {"ogrenci": c.ogrenci or 0, "saat": round((c.dk or 0) / 60, 1), "soru": int(c.soru or 0),
                        "isabet": _yuzde(int(c.dogru or 0), int(c.dogru or 0) + int(c.yanlis or 0)),
                        "program_yapan": program or 0,
                        "dersler": [{"ad": CALISMA_DERSLERI.get(r.ders, r.ders), "saat": round(r.dk / 60, 1), "soru": int(r.soru)} for r in dersler]}

    # ------------------------------------------------------------------ Akademik (okul denemeleri)
    if "okul_denemeleri" in moduller and ids:
        from app.core.sinav_yapisi import TESTLER
        rows = db.execute(text("""
            SELECT od.id, od.ad, od.tarih, od.oturum, d.toplam_net, d.dersler, o.sinif, o.sube
              FROM okul_denemeleri od JOIN ogrenci_denemeleri d ON d.okul_deneme_id = od.id JOIN ogrenciler o ON o.id = d.ogrenci_id
             WHERE od.okul_id = :ok AND d.ogrenci_id = ANY(:ids) AND od.tarih BETWEEN :bas AND :bit
             ORDER BY od.tarih, od.id"""), {**p, "ok": okul_id}).all()
        deneme = {}
        sube_net = defaultdict(lambda: defaultdict(list))
        ders_net = defaultdict(lambda: defaultdict(list))
        for r in rows:
            d = deneme.setdefault(r.id, {"id": r.id, "ad": r.ad, "tarih": r.tarih.isoformat(), "oturum": r.oturum, "netler": []})
            d["netler"].append(float(r.toplam_net))
            et = sube_etiketi(r.sinif, r.sube) if r.sube else (r.sinif or "—")
            sube_net[r.oturum][et].append(float(r.toplam_net))
            for k, x in (r.dersler or {}).items():
                if isinstance(x, dict) and x.get("net") is not None:
                    ders_net[r.oturum][k].append(float(x["net"]))
        v["akademik"] = {
            "denemeler": [{k: d[k] for k in ("id", "ad", "tarih", "oturum")} | {"katilim": len(d["netler"]), "ortalama": _ort(d["netler"])}
                          for d in deneme.values()],
            "oturumlar": sorted(sube_net, key=lambda x: ["TYT", "AYT", "YDT"].index(x) if x in ("TYT", "AYT", "YDT") else 9),
            "subeler": {ot: sorted([{"ad": s, "deger": _ort(l), "katilim": len(l)} for s, l in d.items()], key=lambda x: -(x["deger"] or 0))
                        for ot, d in sube_net.items()},
            "dersler": {ot: [{"ad": TESTLER[k][1] if k in TESTLER else k, "deger": _ort(l), "soru": TESTLER[k][2] if k in TESTLER else None}
                             for k, l in sorted(d.items(), key=lambda kv: list(TESTLER).index(kv[0]) if kv[0] in TESTLER else 99)]
                        for ot, d in ders_net.items()},
        }
    if "net_takibi" in moduller and ids:
        v["net_kullanim"] = db.execute(text("""
            SELECT count(DISTINCT ogrenci_id) FROM ogrenci_denemeleri
             WHERE ogrenci_id = ANY(:ids) AND okul_deneme_id IS NULL AND tarih BETWEEN :bas AND :bit"""), p).scalar() or 0

    # ------------------------------------------------------------------ Rehberlik
    reh = {}
    if "rehberlik" in moduller:
        from app.api.rehberlik import KURALLAR, KONULAR, TURLER, uyarilari_hesapla
        secili = {str(i) for i in ids}
        uy = [s for s in uyarilari_hesapla(db, okul_id) if s["ogrenci_id"] in secili]
        kural = Counter(u["kural"] for s in uy for u in s["uyarilar"])
        reh["erken_uyari"] = {
            "toplam": len(uy), "seviyeler": [{"ad": a, "kod": k, "deger": sum(1 for s in uy if s["seviye"] == k)}
                                             for k, a in (("yuksek", "Yüksek"), ("orta", "Orta"), ("dusuk", "Düşük"))],
            "kurallar": [{"ad": a, "kod": k, "deger": kural.get(k, 0)} for k, a in KURALLAR.items() if kural.get(k)],
            "gorusulmemis_yuksek": sum(1 for s in uy if s["seviye"] == "yuksek" and not s["son_gorusme"] and not s["yaklasan_gorusme"]),
        }
        g = db.execute(text(f"""
            SELECT g.tur, g.konu, g.durum, to_char(g.zaman AT TIME ZONE 'Europe/Istanbul', 'YYYY-MM') AS ay
              FROM rehberlik_gorusmeleri g WHERE g.ogrenci_id = ANY(:ids) AND {TRG % 'g.zaman'} BETWEEN :bas AND :bit"""), p).all() if ids else []
        yapilan = [x for x in g if x.durum == "yapildi"]
        aylar = Counter(x.ay for x in yapilan)
        reh["gorusmeler"] = {
            "toplam": len(g), "yapildi": len(yapilan), "planlandi": sum(1 for x in g if x.durum == "planlandi"),
            "turler": [{"ad": a, "deger": sum(1 for x in yapilan if x.tur == k)} for k, a in TURLER.items()],
            "konular": sorted([{"ad": a, "deger": sum(1 for x in yapilan if x.konu == k)} for k, a in KONULAR.items()], key=lambda x: -x["deger"]),
            "aylik": [{"ay": a, "deger": aylar[a]} for a in sorted(aylar)],
        }
    if "anketler" in moduller and okul_id:
        from app.api.anketler import _hedefte_mi, _j
        anketler = db.execute(text("SELECT id, baslik, durum, anonim, hedef FROM anketler WHERE okul_id = :ok AND durum <> 'taslak' "
                                   "ORDER BY olusturulma_zamani DESC"), {"ok": okul_id}).mappings().all()
        katilim = Counter()
        if anketler and ids:
            for aid, n in db.execute(text("SELECT anket_id, count(DISTINCT ogrenci_id) FROM anket_katilim WHERE anket_id = ANY(:a) "
                                          "AND ogrenci_id = ANY(:ids) GROUP BY 1"), {**p, "a": [a["id"] for a in anketler]}).all():
                katilim[aid] = n
        liste = []
        for a in anketler:
            hd = sum(1 for o in ogr if _hedefte_mi({"hedef": _j(a["hedef"], {})}, o))
            liste.append({"id": a["id"], "ad": a["baslik"], "durum": a["durum"], "anonim": a["anonim"], "hedef": hd,
                          "katilim": katilim[a["id"]], "oran": _yuzde(katilim[a["id"]], hd)})
        reh["anketler"] = liste
    if "tercih" in moduller and okul_id:
        from app.api.tercih import DURUMLAR as TERCIH_DURUM
        son = [o for o in ogr if o.sinif in ("12. Sınıf", "Mezun")]
        d = dict(db.execute(text("SELECT ogrenci_id, durum FROM tercih_listesi l WHERE ogrenci_id = ANY(:i) "
                                 "AND EXISTS (SELECT 1 FROM tercihler t WHERE t.ogrenci_id = l.ogrenci_id)"),
                            {"i": [o.id for o in son]}).all()) if son else {}
        say = Counter(d.values())
        reh["tercih"] = {"hedef": len(son), "liste_var": len(d),
                         "durumlar": [{"ad": "Liste yok", "kod": "yok", "deger": len(son) - len(d)}]
                         + [{"ad": a, "kod": k, "deger": say.get(k, 0)} for k, a in TERCIH_DURUM.items()]}
    if "mezun_takibi" in moduller and okul_id:
        yillar = {}
        for r in db.execute(text("SELECT yil, durum, hedefle_ayni, oneri_sirasi, ogrenci_id, bolum_id FROM mezun_yerlesmeleri WHERE okul_id = :ok"),
                            {"ok": okul_id}).all():
            y = yillar.setdefault(r.yil, {"yil": r.yil, "kayit": 0, "yerlesti": 0, "hedef_bilinen": 0, "hedefle_ayni": 0,
                                          "oneri_bilinen": 0, "ilk10": 0})
            y["kayit"] += 1
            if r.durum == "yerlesti":
                y["yerlesti"] += 1
                if r.hedefle_ayni is not None:
                    y["hedef_bilinen"] += 1
                    y["hedefle_ayni"] += bool(r.hedefle_ayni)
                if r.ogrenci_id and r.bolum_id:
                    y["oneri_bilinen"] += 1
                    y["ilk10"] += bool(r.oneri_sirasi and r.oneri_sirasi <= 10)
        reh["mezun"] = sorted(yillar.values(), key=lambda y: -y["yil"])
    if "kulupler" in moduller and okul_id:
        kul = db.execute(text("""
            SELECT k.ad, count(DISTINCT u.ogrenci_id) FILTER (WHERE u.durum = 'onaylandi') AS uye,
                   count(DISTINCT u.ogrenci_id) FILTER (WHERE u.durum = 'bekliyor') AS bekleyen
              FROM okul_kulupleri k LEFT JOIN kulup_uyelikleri u ON u.kulup_id = k.id AND u.ogrenci_id = ANY(:ids)
             WHERE k.okul_id = :ok AND k.aktif GROUP BY k.id, k.ad ORDER BY uye DESC, k.ad"""), {**p, "ok": okul_id}).all()
        uye_ogr = db.execute(text("""SELECT count(DISTINCT u.ogrenci_id) FROM kulup_uyelikleri u JOIN okul_kulupleri k ON k.id = u.kulup_id
                                      WHERE k.okul_id = :ok AND u.durum = 'onaylandi' AND u.ogrenci_id = ANY(:ids)"""), {**p, "ok": okul_id}).scalar() if ids else 0
        reh["kulupler"] = {"kulupler": [{"ad": r.ad, "deger": r.uye, "bekleyen": r.bekleyen} for r in kul],
                           "uye_ogrenci": uye_ogr or 0, "oran": _yuzde(uye_ogr or 0, toplam)}
    v["rehberlik"] = reh
    return v


# ============================================================================= uç noktalar
@router.get("/okul/{okul_id}/istatistik")
def okul_istatistik(okul_id: int, sinif: str | None = None, sube: str | None = None,
                    bas: date | None = Query(None), bit: date | None = Query(None),
                    db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    _okul_kapsami(db, yon, okul_id)
    return okul_istatistik_verisi(db, okul_id, sinif, sube, bas, bit)


def _sayfa(wb, ad: str, kolonlar: list[str], satirlar: list[list], not_metni: str | None = None):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet(ad[:31])
    ilk = 1
    if not_metni:
        ws.cell(row=1, column=1, value=not_metni).font = Font(italic=True, color="7A5C00")
        ilk = 3
    for j, k in enumerate(kolonlar, 1):
        c = ws.cell(row=ilk, column=j, value=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5D8A")
        c.alignment = Alignment(vertical="center")
    for i, r in enumerate(satirlar, ilk + 1):
        for j, x in enumerate(r, 1):
            ws.cell(row=i, column=j, value=x)
    for j, k in enumerate(kolonlar, 1):
        uz = max([len(str(k))] + [len(str(r[j - 1] if r[j - 1] is not None else "")) for r in satirlar[:500]])
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = min(48, uz + 3)
    ws.freeze_panes = ws.cell(row=ilk + 1, column=1)


def istatistik_xlsx(v: dict, okul_adi: str) -> bytes:
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    f, k = v["filtre"], v["kpi"]
    ust = f"{okul_adi} · {f['etiket']} · {f['bas']} – {f['bit']} · test hesapları hariç"
    _sayfa(wb, "Özet", ["Gösterge", "Değer"], [
        ["Öğrenci", k["ogrenci"]], ["Giriş yapan (hiç)", k["giris_yapan"]], ["Son 7 gün aktif", k["aktif_7"]],
        ["Aralıkta giriş yapan", k["aralikta_giris"]], ["Teste başlayan", k["teste_baslayan"]],
        ["Testi tamamlayan", k["tamamlayan"]], ["Tamamlama oranı (%)", k["tamamlama_orani"]],
        ["K5 (alan derinleşme) tamamlayan", k["k5_tamamlayan"]], ["Hedef bölüm seçen", k["hedef_secen"]],
        ["Ortalama güven puanı", k["ort_guven"]], ["Doğrulanamayan tur", k["gecersiz_tur"]],
        *([["Koçlukta aktif öğrenci (aralıkta)", k["kocluk_aktif"]]] if k.get("kocluk_aktif") is not None else []),
    ], ust)
    ka = v["katilim"]
    _sayfa(wb, "Haftalık katılım", ["Hafta (pazartesi)", "Giriş yapan öğrenci", "Testi tamamlayan"],
           [[h["hafta"], h["giris"], h["tamamlama"]] for h in ka["haftalik"]])
    _sayfa(wb, "Huni", ["Adım", "Öğrenci", "İlk adıma göre %"],
           [[h["ad"], h["deger"], _yuzde(h["deger"], ka["huni"][0]["deger"])] for h in ka["huni"]])
    _sayfa(wb, "Şubeler", ["Şube", "Öğrenci"] + [a for _, a in DURUM_ADLARI] + ["Tamamlama %"],
           [[s["etiket"], s["toplam"]] + [s[kd] for kd, _ in DURUM_ADLARI] + [_yuzde(s["tamamlandi"], s["toplam"])] for s in ka["subeler"]])
    pr = v["profil"]
    kodlar = [x["kod"] for x in pr["katmanlar"]]
    if pr["katman_ort"]:
        _sayfa(wb, "Katman ortalamaları", ["Kapsam"] + kodlar, [[g["ad"]] + [g["degerler"].get(c) for c in kodlar] for g in pr["katman_ort"]])
        _sayfa(wb, "Şube × katman", ["Şube"] + kodlar, [[g["ad"]] + [g["degerler"].get(c) for c in kodlar] for g in pr["sube_katman"]])
        _sayfa(wb, "Özellikler", ["Öne çıkanlar", "Katman", "Ortalama", "", "En az öne çıkanlar", "Katman", "Ortalama"],
               [[*(lambda a: [a["ad"], a["katman"], a["ortalama"]] if a else ["", "", ""])(pr["guclu"][i] if i < len(pr["guclu"]) else None), "",
                 *(lambda a: [a["ad"], a["katman"], a["ortalama"]] if a else ["", "", ""])(pr["zayif"][i] if i < len(pr["zayif"]) else None)]
                for i in range(max(len(pr["guclu"]), len(pr["zayif"])))])
    b = v["bolum"]
    _sayfa(wb, "Bölümler", ["1. sırada önerilen bölüm", "Öğrenci", "", "Hedeflenen bölüm", "Öğrenci"],
           [[*(lambda a: [a["ad"], a["deger"]] if a else ["", ""])(b["ilk_oneri"][i] if i < len(b["ilk_oneri"]) else None), "",
             *(lambda a: [a["ad"], a["deger"]] if a else ["", ""])(b["hedef"][i] if i < len(b["hedef"]) else None)]
            for i in range(max(len(b["ilk_oneri"]), len(b["hedef"])))],
           f"Hedefi ilk 10 önerisinde olan: {b['uyum']['ilk10']} / {b['uyum']['hedefli_tamamlayan']} (%{b['uyum']['oran'] or 0})")
    if b["k5_alan"]:
        _sayfa(wb, "K5 alanları", ["Alan", "Tamamlayan öğrenci"], [[a["ad"], a["deger"]] for a in b["k5_alan"]])
    if b.get("simulasyon"):
        _sayfa(wb, "Simülasyonlar", ["Meslek", "Simülasyon", "Öğrenci", "Ortalama keyif"],
               [[s["ad"], s["deger"], s["ogrenci"], s["keyif"]] for s in b["simulasyon"]])
    if v.get("kocluk"):
        kc = v["kocluk"]
        _sayfa(wb, "Koçluk", ["Hafta (pazartesi)", "Görev", "Tamamlanan", "Oran %", "Öğrenci"],
               [[h["hafta"], h["toplam"], h["tamamlanan"], h["oran"], h["ogrenci"]] for h in kc["gorev_haftalik"]],
               f"Aralıkta tamamlanan gelişim adımı: {kc['adim']['tamamlanan']} ({kc['adim']['ogrenci']} öğrenci) · "
               f"Seri: ortalama {kc['seri']['ortalama'] or 0} hafta, en uzun {kc['seri']['en_uzun']}")
    if v.get("calisma"):
        c = v["calisma"]
        _sayfa(wb, "Çalışma", ["Ders", "Saat", "Soru"], [[d["ad"], d["saat"], d["soru"]] for d in c["dersler"]],
               f"{c['ogrenci']} öğrenci · {c['saat']} saat · {c['soru']} soru · isabet %{c['isabet'] or 0} · program yapan {c['program_yapan']}")
    if v.get("akademik"):
        ak = v["akademik"]
        _sayfa(wb, "Okul denemeleri", ["Tarih", "Deneme", "Oturum", "Katılım", "Ortalama net"],
               [[d["tarih"], d["ad"], d["oturum"], d["katilim"], d["ortalama"]] for d in ak["denemeler"]])
        _sayfa(wb, "Deneme şube ort.", ["Oturum", "Şube", "Ortalama net", "Sonuç sayısı"],
               [[ot, s["ad"], s["deger"], s["katilim"]] for ot, l in ak["subeler"].items() for s in l])
        _sayfa(wb, "Deneme ders ort.", ["Oturum", "Ders", "Ortalama net", "Soru"],
               [[ot, d["ad"], d["deger"], d["soru"]] for ot, l in ak["dersler"].items() for d in l])
    r = v.get("rehberlik") or {}
    if r.get("erken_uyari"):
        e = r["erken_uyari"]
        _sayfa(wb, "Erken uyarı", ["Seviye / kural", "Öğrenci"],
               [[s["ad"] + " seviye", s["deger"]] for s in e["seviyeler"]] + [["", None]] + [[x["ad"], x["deger"]] for x in e["kurallar"]])
    if r.get("gorusmeler"):
        g = r["gorusmeler"]
        _sayfa(wb, "Görüşmeler", ["Tür / konu / ay", "Yapılan görüşme"],
               [[x["ad"], x["deger"]] for x in g["turler"]] + [["", None]] + [[x["ad"], x["deger"]] for x in g["konular"]]
               + [["", None]] + [[x["ay"], x["deger"]] for x in g["aylik"]],
               f"Toplam {g['toplam']} kayıt · {g['yapildi']} yapıldı · {g['planlandi']} planlandı")
    if r.get("anketler"):
        _sayfa(wb, "Anketler", ["Anket", "Durum", "Hedef öğrenci", "Katılan", "Katılım %"],
               [[a["ad"], a["durum"], a["hedef"], a["katilim"], a["oran"]] for a in r["anketler"]])
    if r.get("tercih"):
        _sayfa(wb, "Tercih", ["Durum", "Öğrenci"], [[x["ad"], x["deger"]] for x in r["tercih"]["durumlar"]],
               f"12. sınıf ve mezun: {r['tercih']['hedef']} öğrenci")
    if r.get("mezun"):
        _sayfa(wb, "Mezunlar", ["Yıl", "Kayıt", "Yerleşti", "Hedefi bilinen", "Hedefiyle aynı", "Önerisi bilinen", "İlk 10 öneride"],
               [[y["yil"], y["kayit"], y["yerlesti"], y["hedef_bilinen"], y["hedefle_ayni"], y["oneri_bilinen"], y["ilk10"]] for y in r["mezun"]])
    if r.get("kulupler"):
        _sayfa(wb, "Kulüpler", ["Kulüp", "Üye", "Bekleyen talep"], [[x["ad"], x["deger"], x["bekleyen"]] for x in r["kulupler"]["kulupler"]])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


@router.get("/okul/{okul_id}/istatistik/excel")
def okul_istatistik_excel(okul_id: int, sinif: str | None = None, sube: str | None = None,
                          bas: date | None = Query(None), bit: date | None = Query(None),
                          db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    from app.api.raporlar import _cevap, _dosya_adi
    okul = _okul_kapsami(db, yon, okul_id)
    v = okul_istatistik_verisi(db, okul_id, sinif, sube, bas, bit)
    okul_adi = okul.ad if okul else "Okul harici"
    icerik = istatistik_xlsx(v, okul_adi)
    denetim_yaz(db, yon, "rapor_indir", "okullar", okul_id,
                f"İstatistikler Excel ({v['filtre']['etiket']}, {v['filtre']['bas']} – {v['filtre']['bit']})", okul_id or None)
    db.commit()
    return _cevap(icerik, "xlsx", _dosya_adi("Istatistikler", okul_adi, v["filtre"]["etiket"] if v["filtre"]["sinif"] else None,
                                             datetime.now(TR).strftime("%Y%m%d")))
