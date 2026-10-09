# -*- coding: utf-8 -*-
"""
[2026-10-09] Okul bazlı yönetim — 3 yetki seviyesi:
  1) super_admin    : tüm okullar + okul harici öğrenciler (okul_id = 0) + okul yetkilisi hesapları
  2) okul_yetkilisi : YALNIZCA kendi okulu (öğrenci hesapları, istatistik, kayıtlar)
  3) öğrenci        : bu uç noktalara erişemez

Okul anahtarı: /yonetim/okul/{okul_id} — 0 = okul harici (bireysel) öğrenciler, yalnızca süper admin.

GET    /yonetim/ben                              — oturumdaki yönetim hesabı (+ okulu)
POST   /yonetim/ben/ilk-sifre                    — geçici şifreyle girişte yeni şifre belirleme
GET    /yonetim/okullar                          — erişilebilen okullar ve özet sayılar
GET    /yonetim/okul/{okul_id}/ozet              — okul istatistikleri
GET    /yonetim/okul/{okul_id}/ogrenciler        — öğrenci listesi (ilerleme, son giriş, öneri)
GET    /yonetim/sablon                           — toplu yükleme Excel şablonu (base64)
POST   /yonetim/okul/{okul_id}/onizle            — Excel/CSV dosyasını okuyup satırları doğrular
POST   /yonetim/okul/{okul_id}/ogrenciler        — hesapları oluşturur, geçici şifreleri BİR KEZ döner (+ Excel)
GET    /yonetim/ogrenci/{id}                     — öğrenci detayı (hesap, ilerleme, sonuç, istatistik, kayıtlar)
PUT    /yonetim/ogrenci/{id}                     — ad / e-posta / sınıf / şube düzelt
POST   /yonetim/ogrenci/{id}/okul                — (süper admin) başka okula / okul harici aktar
POST   /yonetim/ogrenciler/sifre-sifirla         — seçilenlere yeni geçici şifre (+ Excel)
POST   /yonetim/ogrenciler/sil                   — seçilenleri tüm verileriyle sil
GET    /yonetim/okul/{okul_id}/yetkililer        — okul yetkilileri
POST   /yonetim/okul/{okul_id}/yetkililer        — (süper admin) okul yetkilisi ekle → geçici şifre
POST   /yonetim/yetkili/{id}/sifre-sifirla       — (süper admin) yetkiliye yeni geçici şifre
DELETE /yonetim/yetkili/{id}                     — (süper admin) yetkiliyi sil
GET    /yonetim/okul/{okul_id}/kayitlar          — okulun işlem kayıtları (yönetim işlemleri + öğrenci giriş/şifre olayları)
PUT    /yonetim/okul/{okul_id}/tema             — okul rengi (#RRGGBB; boş = Filizyol rengi) — okul yetkilisi + süper admin
GET    /yonetim/okul/{okul_id}/bilgi             — okul tanıtım bilgileri (kuruluş yılı, öğrenci sayısı, tanıtım, iletişim, kadro)
PUT    /yonetim/okul/{okul_id}/bilgi             — tanıtım bilgilerini güncelle (okul yetkilisi + süper admin)
"""
import base64
import csv
import io
import re
import uuid
from collections import Counter, defaultdict
from datetime import datetime, timedelta

import bcrypt
from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel
from sqlalchemy import func, text
from sqlalchemy.orm import Session

from app.api.deps import get_mevcut_yonetim
from app.core.database import get_db
from app.core.hesap_yonetimi import (
    denetim_yaz, gecici_sifre, ogrenciyi_sil, olay_yaz, simdi, yapan_etiketi,
)
from app.core.security import sifre_dogrula, sifre_hashle
from app.models import AdminKullanici, AuditLog, Bolum, Ogrenci, OgrenciHesapOlayi, Okul

router = APIRouter(prefix="/yonetim", tags=["Okul yönetimi"])

EN_FAZLA_SATIR = 2000
SINIFLAR = ["9. Sınıf", "10. Sınıf", "11. Sınıf", "12. Sınıf", "Mezun"]
_EPOSTA = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]{2,}$")

OLAY_ETIKET = {
    "hesap_acildi": "Hesap açıldı",
    "giris": "Giriş yaptı",
    "giris_hatali": "Hatalı şifre denemesi",
    "sifre_sifirlandi": "Şifre sıfırlandı (yönetim)",
    "sifre_degisti": "Şifresini değiştirdi",
    "bilgi_guncellendi": "Bilgileri güncellendi",
    "okul_degisti": "Okulu değiştirildi",
    "katman_basladi": "Test bölümüne başladı",
    "katman_bitti": "Test bölümünü tamamladı",
}


# ----------------------------------------------------------------------------- yetki / kapsam
def _super(yon: AdminKullanici) -> None:
    if yon.rol != "super_admin":
        raise HTTPException(status_code=403, detail="Bu işlem yalnızca süper admine açık.")


def _okul_kapsami(db: Session, yon: AdminKullanici, okul_id: int) -> Okul | None:
    """okul_id=0 → okul harici (yalnızca süper admin). Okul yetkilisi yalnızca kendi okuluna erişir."""
    if yon.rol == "okul_yetkilisi" and okul_id != yon.okul_id:
        raise HTTPException(status_code=403, detail="Yalnızca kendi okulunuzun bilgilerine erişebilirsiniz.")
    if okul_id == 0:
        _super(yon)
        return None
    okul = db.get(Okul, okul_id)
    if okul is None:
        raise HTTPException(status_code=404, detail="Okul bulunamadı.")
    return okul


def _ogrenci_kapsami(db: Session, yon: AdminKullanici, ogrenci_id) -> Ogrenci:
    try:
        oid = uuid.UUID(str(ogrenci_id))
    except ValueError:
        raise HTTPException(status_code=400, detail="Geçersiz öğrenci.")
    ogr = db.get(Ogrenci, oid)
    if ogr is None:
        raise HTTPException(status_code=404, detail="Öğrenci bulunamadı.")
    if yon.rol == "okul_yetkilisi" and ogr.okul_id != yon.okul_id:
        raise HTTPException(status_code=403, detail="Bu öğrenci sizin okulunuzda değil.")
    return ogr


def _okul_filtresi(okul_id: int) -> tuple[str, dict]:
    return ("o.okul_id IS NULL", {}) if okul_id == 0 else ("o.okul_id = :okul_id", {"okul_id": okul_id})


def _hizli_hash(sifre: str) -> str:
    # Geçici şifreler ilk girişte zorunlu değiştirildiği için toplu açılışta daha düşük maliyet (10) kullanılır;
    # doğrulama aynı bcrypt fonksiyonuyla yapılır.
    return bcrypt.hashpw(sifre.encode("utf-8"), bcrypt.gensalt(rounds=10)).decode("utf-8")


def _eposta_kullanimda(db: Session, email: str, haric: uuid.UUID | None = None) -> bool:
    e = email.strip().lower()
    q = db.query(Ogrenci.id).filter(func.lower(Ogrenci.email) == e)
    if haric is not None:
        q = q.filter(Ogrenci.id != haric)
    return q.first() is not None or db.query(AdminKullanici.id).filter(func.lower(AdminKullanici.email) == e).first() is not None


# ----------------------------------------------------------------------------- sınıf / şube ayrıştırma
def sinif_ayir(sinif_ham: str | None, sube_ham: str | None = None) -> tuple[str | None, str | None]:
    """'11-B', '11/B', '11 B', '11B', '11. sınıf', 'Mezun' → ('11. Sınıf', 'B')"""
    s = (str(sinif_ham or "")).strip()
    sube = (str(sube_ham or "")).strip().upper() or None
    if not s:
        return None, sube
    if "mezun" in s.lower():
        return "Mezun", sube
    m = re.match(r"^\s*(9|10|11|12)\s*(?:\.?\s*s[ıi]n[ıi]f[ıi]?)?\s*[-/ .]?\s*([A-Za-zÇĞİÖŞÜçğıöşü]{1,2})?\s*$", s, re.I)
    if m:
        return f"{m.group(1)}. Sınıf", (sube or (m.group(2).upper() if m.group(2) else None))
    m = re.match(r"^\s*(9|10|11|12)\b", s)
    if m:
        return f"{m.group(1)}. Sınıf", sube
    return s[:30], sube


def _sinif_metni(o: Ogrenci) -> str:
    if not o.sinif:
        return ""
    return f"{o.sinif.replace('. Sınıf', '')}-{o.sube}" if o.sube and o.sinif != "Mezun" else o.sinif


# ----------------------------------------------------------------------------- dosya okuma (Excel / CSV)
def _norm(s) -> str:
    t = str(s or "").strip().lower()
    for a, b in (("ı", "i"), ("ğ", "g"), ("ü", "u"), ("ş", "s"), ("ö", "o"), ("ç", "c"), ("i̇", "i")):
        t = t.replace(a, b)
    return re.sub(r"[^a-z0-9]", "", t)


_BASLIK = {
    "ad_soyad": {"adsoyad", "adisoyadi", "adsoyadi", "adisoyad", "ogrenci", "ogrenciadi", "ogrenciadisoyadi", "isimsoyisim", "advesoyad"},
    "ad": {"ad", "adi", "isim"},
    "soyad": {"soyad", "soyadi", "soyisim"},
    "sinif": {"sinif", "sinifi", "sinifduzeyi", "duzey", "sinifsube", "sinifvesube"},
    "sube": {"sube", "subesi"},
    "email": {"eposta", "email", "mail", "epostaadresi", "emailadresi", "elektronikposta"},
}


def _dosyadan_satirlar(dosya_adi: str, icerik_b64: str) -> list[list]:
    try:
        veri = base64.b64decode(icerik_b64.split(",", 1)[-1])
    except Exception:
        raise HTTPException(status_code=400, detail="Dosya okunamadı.")
    if len(veri) > 5_000_000:
        raise HTTPException(status_code=400, detail="Dosya çok büyük (en fazla 5 MB).")
    ad = (dosya_adi or "").lower()
    if ad.endswith((".xlsx", ".xlsm")) or veri[:2] == b"PK":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(veri), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            return [[("" if c is None else c) for c in r] for r in ws.iter_rows(values_only=True)]
        except Exception:
            raise HTTPException(status_code=400, detail="Excel dosyası açılamadı. .xlsx olarak kaydedip tekrar deneyin.")
    if ad.endswith(".xls"):
        raise HTTPException(status_code=400, detail="Eski .xls biçimi desteklenmiyor; dosyayı .xlsx olarak kaydedin.")
    for kod in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            metin = veri.decode(kod)
            break
        except UnicodeDecodeError:
            continue
    ilk = metin.splitlines()[0] if metin else ""
    ayrac = max([";", ",", "\t"], key=ilk.count)
    return [r for r in csv.reader(io.StringIO(metin), delimiter=ayrac)]


def _satirlari_coz(satirlar: list[list]) -> list[dict]:
    satirlar = [r for r in satirlar if any(str(c).strip() for c in r)]
    if not satirlar:
        raise HTTPException(status_code=400, detail="Dosyada satır bulunamadı.")
    # başlık satırı: ilk 6 satırda e-posta sütunu tanınan ilk satır (şablondaki açıklama satırı atlanır)
    bas_no, kolon = 0, {}
    for n, aday in enumerate(satirlar[:6]):
        k = {}
        for i, b in enumerate(_norm(c) for c in aday):
            for alan, adlar in _BASLIK.items():
                if b in adlar and alan not in k:
                    k[alan] = i
        if "email" in k:
            bas_no, kolon = n, k
            break
    satirlar = satirlar[bas_no:]
    if "email" not in kolon or not ({"ad_soyad"} <= kolon.keys() or {"ad", "soyad"} <= kolon.keys() or "ad" in kolon):
        raise HTTPException(status_code=400, detail="Başlık satırı tanınmadı. Şablondaki sütunları kullanın: Ad Soyad, Sınıf, Şube, E-posta.")
    al = lambda r, k: (str(r[kolon[k]]).strip() if k in kolon and kolon[k] < len(r) and r[kolon[k]] is not None else "")
    sonuc = []
    for no, r in enumerate(satirlar[1:EN_FAZLA_SATIR + 1], start=2):
        ad_soyad = al(r, "ad_soyad") or " ".join(x for x in (al(r, "ad"), al(r, "soyad")) if x)
        sinif_ham = al(r, "sinif")
        if isinstance(sinif_ham, str) and sinif_ham.endswith(".0"):
            sinif_ham = sinif_ham[:-2]
        sinif, sube = sinif_ayir(sinif_ham, al(r, "sube"))
        sonuc.append({"satir": no, "ad_soyad": re.sub(r"\s+", " ", ad_soyad).strip(), "sinif": sinif, "sube": sube,
                      "email": al(r, "email").replace(" ", "")})
    if len(satirlar) - 1 > EN_FAZLA_SATIR:
        raise HTTPException(status_code=400, detail=f"Tek seferde en fazla {EN_FAZLA_SATIR} öğrenci yüklenebilir.")
    return sonuc


def _dogrula(db: Session, satirlar: list[dict]) -> list[dict]:
    gorulen = Counter(s["email"].lower() for s in satirlar if s["email"])
    for s in satirlar:
        hatalar, uyarilar = [], []
        if len(s["ad_soyad"]) < 3:
            hatalar.append("Ad soyad eksik")
        if not s["email"]:
            hatalar.append("E-posta eksik")
        elif not _EPOSTA.match(s["email"]):
            hatalar.append("E-posta geçersiz")
        elif gorulen[s["email"].lower()] > 1:
            hatalar.append("E-posta dosyada birden fazla kez geçiyor")
        elif s["email"].lower().endswith("@ornek.com"):
            hatalar.append("Şablondaki örnek satır — silin")
        elif _eposta_kullanimda(db, s["email"]):
            hatalar.append("Bu e-posta ile zaten hesap var")
        if not s["sinif"]:
            uyarilar.append("Sınıf boş — öğrenci ilk girişte seçer")
        elif s["sinif"] not in SINIFLAR:
            uyarilar.append(f"Sınıf tanınmadı ('{s['sinif']}'), olduğu gibi kaydedilecek")
        s["hata"] = "; ".join(hatalar) or None
        s["uyari"] = "; ".join(uyarilar) or None
    return satirlar


# ----------------------------------------------------------------------------- Excel çıktısı
def _xlsx_b64(baslik: str, kolonlar: list[str], satirlar: list[list], not_metni: str | None = None) -> str:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = baslik[:31]
    ilk = 1
    if not_metni:
        ws.cell(row=1, column=1, value=not_metni).font = Font(italic=True, color="7A5C00")
        ilk = 3
    for j, k in enumerate(kolonlar, start=1):
        c = ws.cell(row=ilk, column=j, value=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5D8A")
        c.alignment = Alignment(vertical="center")
    for i, r in enumerate(satirlar, start=ilk + 1):
        for j, v in enumerate(r, start=1):
            ws.cell(row=i, column=j, value=v)
    for j, k in enumerate(kolonlar, start=1):
        uz = max([len(str(k))] + [len(str(r[j - 1] or "")) for r in satirlar[:500]])
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = min(45, uz + 3)
    ws.freeze_panes = ws.cell(row=ilk + 1, column=1)
    b = io.BytesIO()
    wb.save(b)
    return base64.b64encode(b.getvalue()).decode()


# ----------------------------------------------------------------------------- ilerleme / durum
def _ilerleme(db: Session, okul_id: int | None = None, ogrenci_id: uuid.UUID | None = None) -> dict:
    """{ogrenci_id: {tur_id, tur_durumu, biten, toplam, baslayan, ilk_bolum, ilk_uyum, hedef, tamamlanma}}"""
    if ogrenci_id is not None:
        kosul, p = "o.id = :oid", {"oid": ogrenci_id}
    else:
        kosul, p = _okul_filtresi(okul_id)
    satirlar = db.execute(text(f"""
        WITH son AS (
            SELECT DISTINCT ON (t.ogrenci_id) t.id, t.ogrenci_id, t.durum, t.tamamlanma_zamani
              FROM ogrenci_degerlendirme_turu t JOIN ogrenciler o ON o.id = t.ogrenci_id
             WHERE {kosul}
             ORDER BY t.ogrenci_id, t.tur_no DESC
        )
        SELECT son.ogrenci_id, son.id AS tur_id, son.durum, son.tamamlanma_zamani,
               (SELECT count(*) FROM ogrenci_katman_oturumlari k JOIN katmanlar km ON km.id = k.katman_id
                 WHERE k.tur_id = son.id AND k.durum = 'tamamlandi' AND NOT km.kosullu_mu) AS biten,
               (SELECT count(*) FROM ogrenci_katman_oturumlari k JOIN katmanlar km ON km.id = k.katman_id
                 WHERE k.tur_id = son.id AND k.durum <> 'baslamadi' AND NOT km.kosullu_mu) AS baslayan,
               (SELECT s.bolum_id FROM ogrenci_bolum_uyum_skorlari s WHERE s.tur_id = son.id
                 ORDER BY s.toplam_uyum DESC LIMIT 1) AS ilk_bolum
          FROM son
    """), p).mappings().all()
    hedefler = db.execute(text(f"""
        SELECT hb.ogrenci_id, hb.bolum_id FROM ogrenci_hedef_bolum hb JOIN ogrenciler o ON o.id = hb.ogrenci_id
         WHERE hb.aktif_mi AND {kosul}
    """), p).mappings().all()
    toplam = db.execute(text("SELECT count(*) FROM katmanlar WHERE NOT kosullu_mu")).scalar() or 4
    sonuc = {r["ogrenci_id"]: {**dict(r), "toplam": toplam, "hedef": None} for r in satirlar}
    for h in hedefler:
        sonuc.setdefault(h["ogrenci_id"], {"tur_id": None, "durum": None, "biten": 0, "baslayan": 0, "ilk_bolum": None,
                                            "toplam": toplam, "tamamlanma_zamani": None})["hedef"] = h["bolum_id"]
    return sonuc


def _durum(o: Ogrenci, il: dict | None) -> tuple[str, str]:
    """(kod, etiket)"""
    if il and il.get("durum") == "tamamlandi":
        return "tamamlandi", "Testi tamamladı"
    if il and (il.get("baslayan") or 0) > 0:
        return "devam", f"Devam ediyor ({il.get('biten', 0)}/{il.get('toplam', 4)})"
    if o.son_giris_zamani is None and o.sifre_degistirmeli:
        return "giris_yok", "Henüz giriş yapmadı"
    return "baslamadi", "Teste başlamadı"


# ----------------------------------------------------------------------------- şemalar
class IlkSifreIstek(BaseModel):
    yeni_sifre: str


class DosyaIstek(BaseModel):
    dosya_adi: str
    icerik_base64: str


class OgrenciSatir(BaseModel):
    ad_soyad: str
    email: str
    sinif: str | None = None
    sube: str | None = None


class TopluOlusturIstek(BaseModel):
    ogrenciler: list[OgrenciSatir]


class OgrenciDuzenleIstek(BaseModel):
    ad_soyad: str | None = None
    email: str | None = None
    sinif: str | None = None
    sube: str | None = None


class IdlerIstek(BaseModel):
    idler: list[str]


class OkulAtaIstek(BaseModel):
    okul_id: int   # 0 = okul harici


class YetkiliEkleIstek(BaseModel):
    ad_soyad: str
    email: str


# ----------------------------------------------------------------------------- oturum sahibi
@router.get("/ben")
def ben(db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    okul = db.get(Okul, yon.okul_id) if yon.okul_id else None
    return {"id": str(yon.id), "ad_soyad": yon.ad_soyad, "email": yon.email, "rol": yon.rol,
            "rol_adi": "Süper Admin" if yon.rol == "super_admin" else "Okul Yetkilisi",
            "okul_id": yon.okul_id, "okul_ad": okul.ad if okul else None, "okul_logo": okul.logo if okul else None,
            "okul_renk": okul.tema_renk if okul else None,
            "sifre_degistirmeli": bool(yon.sifre_degistirmeli)}


@router.post("/ben/ilk-sifre", status_code=204)
def ben_ilk_sifre(istek: IlkSifreIstek, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    if len(istek.yeni_sifre or "") < 8:
        raise HTTPException(status_code=400, detail="Yeni şifre en az 8 karakter olmalı.")
    if sifre_dogrula(istek.yeni_sifre, yon.sifre_hash):
        raise HTTPException(status_code=400, detail="Yeni şifre geçici şifreyle aynı olamaz.")
    yon.sifre_hash = sifre_hashle(istek.yeni_sifre)
    yon.sifre_degistirmeli = False
    denetim_yaz(db, yon, "sifre_belirledi", "admin_kullanicilar", yon.id, "İlk girişte kendi şifresini belirledi", yon.okul_id)
    db.commit()


# ----------------------------------------------------------------------------- okullar / özet
@router.get("/okullar")
def erisilen_okullar(db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    q = db.query(Okul)
    if yon.rol == "okul_yetkilisi":
        q = q.filter(Okul.id == yon.okul_id)
    sayilar = dict(db.query(Ogrenci.okul_id, func.count(Ogrenci.id)).group_by(Ogrenci.okul_id).all())
    sonuc = [{"id": o.id, "ad": o.ad, "logo": o.logo, "ogrenci_sayisi": sayilar.get(o.id, 0)} for o in q.order_by(Okul.ad).all()]
    if yon.rol == "super_admin":
        sonuc.append({"id": 0, "ad": "Okul harici öğrenciler", "logo": None, "ogrenci_sayisi": sayilar.get(None, 0)})
    return sonuc


def _okul_basligi(okul: Okul | None) -> dict:
    return {"id": okul.id if okul else 0, "ad": okul.ad if okul else "Okul harici öğrenciler",
            "alt_baslik": okul.alt_baslik if okul else "Bir okula bağlı olmayan bireysel hesaplar",
            "logo": okul.logo if okul else None}


@router.get("/okul/{okul_id}/ozet")
def okul_ozeti(okul_id: int, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    okul = _okul_kapsami(db, yon, okul_id)
    kosul, p = _okul_filtresi(okul_id)
    ogrenciler = db.query(Ogrenci).filter(Ogrenci.okul_id.is_(None) if okul_id == 0 else Ogrenci.okul_id == okul_id).all()
    il = _ilerleme(db, okul_id)
    adlar = {b.id: b.ad for b in db.query(Bolum.id, Bolum.ad).all()}

    durum_say = Counter()
    sinif = defaultdict(lambda: {"ogrenci": 0, "tamamlayan": 0, "devam": 0, "giris_yapan": 0})
    ilk_bolum, hedef = Counter(), Counter()
    for o in ogrenciler:
        kod, _ = _durum(o, il.get(o.id))
        durum_say[kod] += 1
        anahtar = o.sinif or "Belirtilmedi"
        sinif[anahtar]["ogrenci"] += 1
        sinif[anahtar]["tamamlayan"] += kod == "tamamlandi"
        sinif[anahtar]["devam"] += kod == "devam"
        sinif[anahtar]["giris_yapan"] += o.son_giris_zamani is not None
        x = il.get(o.id) or {}
        if x.get("durum") == "tamamlandi" and x.get("ilk_bolum"):
            ilk_bolum[adlar.get(x["ilk_bolum"], "?")] += 1
        if x.get("hedef"):
            hedef[adlar.get(x["hedef"], "?")] += 1

    toplam = len(ogrenciler)
    giris_yapan = sum(1 for o in ogrenciler if o.son_giris_zamani is not None)
    sira = {s: i for i, s in enumerate(SINIFLAR)}
    bugun = simdi().date()
    gunluk = {(bugun - timedelta(days=i)).isoformat(): 0 for i in range(13, -1, -1)}
    for g, n in db.execute(text(f"""
        SELECT (e.zaman AT TIME ZONE 'Europe/Istanbul')::date AS gun, count(DISTINCT e.ogrenci_id)
          FROM ogrenci_hesap_olaylari e JOIN ogrenciler o ON o.id = e.ogrenci_id
         WHERE e.olay = 'giris' AND e.zaman >= now() - interval '14 days' AND {kosul}
         GROUP BY 1"""), p).all():
        if g.isoformat() in gunluk:
            gunluk[g.isoformat()] = n
    return {
        "okul": _okul_basligi(okul),
        "toplam": toplam,
        "giris_yapan": giris_yapan,
        "teste_baslayan": durum_say["devam"] + durum_say["tamamlandi"],
        "tamamlayan": durum_say["tamamlandi"],
        "hedef_secen": sum(1 for o in ogrenciler if (il.get(o.id) or {}).get("hedef")),
        "durumlar": [{"kod": k, "etiket": e, "sayi": durum_say[k]} for k, e in
                     (("giris_yok", "Henüz giriş yapmadı"), ("baslamadi", "Teste başlamadı"),
                      ("devam", "Devam ediyor"), ("tamamlandi", "Testi tamamladı"))],
        "siniflar": [{"sinif": k, **v} for k, v in sorted(sinif.items(), key=lambda kv: sira.get(kv[0], 99))],
        "en_cok_onerilen": [{"bolum": b, "sayi": n} for b, n in ilk_bolum.most_common(8)],
        "en_cok_hedeflenen": [{"bolum": b, "sayi": n} for b, n in hedef.most_common(8)],
        "gunluk_giris": [{"gun": g, "sayi": n} for g, n in gunluk.items()],
        "yetkili_sayisi": db.query(AdminKullanici).filter(AdminKullanici.okul_id == okul_id).count() if okul_id else 0,
    }


@router.get("/okul/{okul_id}/ogrenciler")
def okul_ogrencileri(okul_id: int, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    _okul_kapsami(db, yon, okul_id)
    sira = {s: i for i, s in enumerate(SINIFLAR)}
    ogrenciler = sorted(db.query(Ogrenci).filter(Ogrenci.okul_id.is_(None) if okul_id == 0 else Ogrenci.okul_id == okul_id).all(),
                        key=lambda o: (sira.get(o.sinif, 98 if o.sinif else 99), o.sube or "", o.ad_soyad.lower()))
    il = _ilerleme(db, okul_id)
    adlar = {b.id: b.ad for b in db.query(Bolum.id, Bolum.ad).all()}
    sonuc = []
    for o in ogrenciler:
        x = il.get(o.id) or {}
        kod, etiket = _durum(o, x)
        sonuc.append({
            "id": str(o.id), "ad_soyad": o.ad_soyad, "email": o.email, "sinif": o.sinif, "sube": o.sube,
            "sinif_metni": _sinif_metni(o), "olusturulma_zamani": o.olusturulma_zamani, "son_giris_zamani": o.son_giris_zamani,
            "sifre_degistirmeli": bool(o.sifre_degistirmeli), "durum": kod, "durum_etiket": etiket,
            "biten_katman": x.get("biten", 0), "toplam_katman": x.get("toplam", 4),
            "ilk_bolum": adlar.get(x.get("ilk_bolum")) if x.get("durum") == "tamamlandi" else None,
            "hedef_bolum": adlar.get(x.get("hedef")),
        })
    return sonuc


# ----------------------------------------------------------------------------- toplu hesap açma
@router.get("/sablon")
def sablon(yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    return {"dosya_adi": "ogrenci_yukleme_sablonu.xlsx", "icerik_base64": _xlsx_b64(
        "Öğrenciler", ["Ad Soyad", "Sınıf", "Şube", "E-posta"],
        [["Ayşe Yılmaz", 11, "A", "ayse.yilmaz@ornek.com"], ["Mehmet Demir", 12, "B", "mehmet.demir@ornek.com"]],
        "Örnek satırları silip kendi öğrencilerinizi yazın. Sınıf: 9, 10, 11, 12 veya Mezun. Şube isteğe bağlı.")}


@router.post("/okul/{okul_id}/onizle")
def onizle(okul_id: int, istek: DosyaIstek, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    _okul_kapsami(db, yon, okul_id)
    satirlar = _satirlari_coz(_dosyadan_satirlar(istek.dosya_adi, istek.icerik_base64))
    return {"satirlar": _dogrula(db, satirlar)}


@router.post("/okul/{okul_id}/ogrenciler")
def toplu_olustur(okul_id: int, istek: TopluOlusturIstek, db: Session = Depends(get_db),
                  yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    okul = _okul_kapsami(db, yon, okul_id)
    if not istek.ogrenciler:
        raise HTTPException(status_code=400, detail="Oluşturulacak öğrenci yok.")
    if len(istek.ogrenciler) > EN_FAZLA_SATIR:
        raise HTTPException(status_code=400, detail=f"Tek seferde en fazla {EN_FAZLA_SATIR} öğrenci.")
    satirlar = []
    for i, s in enumerate(istek.ogrenciler, start=1):
        sinif, sube = sinif_ayir(s.sinif, s.sube)
        satirlar.append({"satir": i, "ad_soyad": re.sub(r"\s+", " ", s.ad_soyad).strip(), "email": s.email.strip(),
                         "sinif": sinif, "sube": sube})
    _dogrula(db, satirlar)
    olusan, hatali = [], []
    yapan = yapan_etiketi(yon)
    for s in satirlar:
        if s["hata"]:
            hatali.append(s)
            continue
        sifre = gecici_sifre()
        o = Ogrenci(ad_soyad=s["ad_soyad"], email=s["email"], sifre_hash=_hizli_hash(sifre), sinif=s["sinif"], sube=s["sube"],
                    okul_id=okul.id if okul else None, okul=okul.ad if okul else None, sifre_degistirmeli=True,
                    olusturulma_zamani=simdi())
        db.add(o)
        db.flush()
        olay_yaz(db, o.id, "hesap_acildi", f"{okul.ad if okul else 'Okul harici'} — toplu hesap açma", yapan)
        olusan.append({**s, "id": str(o.id), "gecici_sifre": sifre})
    denetim_yaz(db, yon, "ogrenci_toplu_ekle", "ogrenciler", okul_id, f"{len(olusan)} hesap açıldı, {len(hatali)} satır atlandı",
                okul_id or None)
    db.commit()
    return {
        "olusturulan": olusan, "hatali": hatali,
        "dosya_adi": f"giris_bilgileri_{(okul.ad if okul else 'okul_harici').replace(' ', '_')}_{simdi():%Y%m%d_%H%M}.xlsx",
        "icerik_base64": _giris_listesi_xlsx(okul, olusan) if olusan else None,
    }


def _giris_listesi_xlsx(okul: Okul | None, kayitlar: list[dict]) -> str:
    return _xlsx_b64("Giriş Bilgileri", ["Ad Soyad", "Sınıf", "Şube", "E-posta (kullanıcı adı)", "Geçici Şifre"],
                     [[k["ad_soyad"], k.get("sinif") or "", k.get("sube") or "", k["email"], k["gecici_sifre"]] for k in kayitlar],
                     f"{okul.ad if okul else 'Okul harici'} — giriş adresi: sitenin /giris sayfası. Geçici şifre ilk girişte "
                     f"değiştirilir. Bu dosyayı güvenli saklayın; şifreler sistemde tekrar gösterilmez.")


# ----------------------------------------------------------------------------- öğrenci detayı
@router.get("/ogrenci/{ogrenci_id}")
def ogrenci_detay(ogrenci_id: str, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    o = _ogrenci_kapsami(db, yon, ogrenci_id)
    adlar = {b.id: b.ad for b in db.query(Bolum.id, Bolum.ad).all()}
    il = _ilerleme(db, ogrenci_id=o.id).get(o.id)
    kod, etiket = _durum(o, il)

    turlar = db.execute(text("""
        SELECT t.id, t.tur_no, t.durum, t.baslama_zamani, t.tamamlanma_zamani, t.guven_skoru, t.sonuc_gecerli_mi
          FROM ogrenci_degerlendirme_turu t WHERE t.ogrenci_id = :id ORDER BY t.tur_no DESC"""), {"id": o.id}).mappings().all()
    katmanlar = db.execute(text("""
        SELECT k.tur_id, km.kod, km.ad, km.kosullu_mu, k.durum, k.baslama_zamani, k.tamamlanma_zamani,
               (SELECT count(*) FROM ogrenci_cevaplar c JOIN sorular s ON s.id = c.soru_id
                 WHERE c.tur_id = k.tur_id AND c.ogrenci_id = k.ogrenci_id AND s.katman_id = k.katman_id) AS cevap
          FROM ogrenci_katman_oturumlari k JOIN katmanlar km ON km.id = k.katman_id
         WHERE k.ogrenci_id = :id ORDER BY km.sira"""), {"id": o.id}).mappings().all()
    dallar = db.execute(text("""
        SELECT d.tur_id, dl.ad, d.durum, d.baslama_zamani, d.tamamlanma_zamani
          FROM ogrenci_dal_oturumlari d JOIN dallar dl ON dl.id = d.dal_id
         WHERE d.ogrenci_id = :id ORDER BY d.id"""), {"id": o.id}).mappings().all()
    guvenlik = db.execute(text("SELECT count(*) FROM guvenlik_olaylari WHERE ogrenci_id = :id AND olay_tipi IN "
                               "('tam_ekrandan_cikti','sekme_degisti','pencere_odagi_kaybedildi')"), {"id": o.id}).scalar() or 0

    def _sure(b, t):
        return round((t - b).total_seconds() / 60, 1) if b and t else None

    tur_listesi = []
    for t in turlar:
        k_list = [{"kod": k["kod"], "ad": k["ad"], "durum": k["durum"], "baslama": k["baslama_zamani"],
                   "tamamlanma": k["tamamlanma_zamani"], "sure_dk": _sure(k["baslama_zamani"], k["tamamlanma_zamani"]),
                   "cevap": k["cevap"]} for k in katmanlar if k["tur_id"] == t["id"]]
        tur_listesi.append({
            "tur_no": t["tur_no"], "durum": t["durum"], "baslama": t["baslama_zamani"], "tamamlanma": t["tamamlanma_zamani"],
            "guven_skoru": float(t["guven_skoru"]) if t["guven_skoru"] is not None else None,
            "sonuc_gecerli_mi": t["sonuc_gecerli_mi"], "katmanlar": k_list,
            "dallar": [{"ad": d["ad"], "durum": d["durum"], "sure_dk": _sure(d["baslama_zamani"], d["tamamlanma_zamani"])}
                       for d in dallar if d["tur_id"] == t["id"]],
            "toplam_sure_dk": round(sum(k["sure_dk"] or 0 for k in k_list), 1),
            "cevap_sayisi": sum(k["cevap"] or 0 for k in k_list),
        })

    # Sonuç: öğrencinin gördüğü öneri listesiyle aynı sıralama (K1-K4 bittiyse)
    oneriler, sonuc_notu = [], None
    if turlar and turlar[0]["durum"] == "tamamlandi":
        try:
            from app.core.skor_motoru import siralama_getir
            from app.models import OgrenciDegerlendirmeTuru
            tur = db.get(OgrenciDegerlendirmeTuru, turlar[0]["id"])
            oneriler = [{"sira": i + 1, "bolum": adlar.get(s.bolum_id, "?"), "uyum": round(float(s.toplam_uyum), 1),
                         "alan": getattr(s, "alan", None)} for i, s in enumerate(siralama_getir(db, o, tur, ilk_n=10))]
            if any(d["durum"] != "tamamlandi" for d in tur_listesi[0]["dallar"]):
                sonuc_notu = "Alan (K5) soruları bitmediği için liste kesinleşmedi."
        except Exception:
            db.rollback()
            # yedek: kayıtlı uyum skorlarından ham sıralama (K5 düzeltmesi ve alan sırası olmadan)
            ham = db.execute(text("SELECT bolum_id, toplam_uyum FROM ogrenci_bolum_uyum_skorlari WHERE tur_id = :t "
                                  "ORDER BY toplam_uyum DESC LIMIT 10"), {"t": turlar[0]["id"]}).all()
            oneriler = [{"sira": i + 1, "bolum": adlar.get(b, "?"), "uyum": round(float(u), 1), "alan": None} for i, (b, u) in enumerate(ham)]
            sonuc_notu = "Ham uyum sıralaması gösteriliyor (öğrencinin ekranındaki alan düzenlemesi uygulanmadı)." if ham else "Sonuç listesi hesaplanamadı."
    else:
        sonuc_notu = "Öğrenci K1–K4 bölümlerini tamamladığında öneri listesi burada görünür."
    hedef = None
    if il and il.get("hedef"):
        uyum = db.execute(text("SELECT toplam_uyum FROM ogrenci_bolum_uyum_skorlari WHERE ogrenci_id = :o AND bolum_id = :b "
                               "ORDER BY tur_id DESC LIMIT 1"), {"o": o.id, "b": il["hedef"]}).scalar()
        hedef = {"bolum": adlar.get(il["hedef"]), "uyum": round(float(uyum), 1) if uyum is not None else None}

    okul = db.get(Okul, o.okul_id) if o.okul_id else None
    return {
        "hesap": {"id": str(o.id), "ad_soyad": o.ad_soyad, "email": o.email, "sinif": o.sinif, "sube": o.sube,
                  "sinif_metni": _sinif_metni(o), "okul_id": o.okul_id or 0, "okul_ad": okul.ad if okul else "Okul harici",
                  "olusturulma_zamani": o.olusturulma_zamani, "son_giris_zamani": o.son_giris_zamani,
                  "sifre_degistirmeli": bool(o.sifre_degistirmeli), "durum": kod, "durum_etiket": etiket,
                  "dogum_tarihi": o.dogum_tarihi, "ilgi_alanlari": o.ilgi_alanlari},
        "istatistik": {
            "tur_sayisi": len(turlar),
            "biten_katman": (il or {}).get("biten", 0), "toplam_katman": (il or {}).get("toplam", 4),
            "toplam_sure_dk": tur_listesi[0]["toplam_sure_dk"] if tur_listesi else 0,
            "cevap_sayisi": tur_listesi[0]["cevap_sayisi"] if tur_listesi else 0,
            "guven_skoru": tur_listesi[0]["guven_skoru"] if tur_listesi else None,
            "dikkat_dagilmasi": guvenlik,
            "giris_sayisi": db.query(OgrenciHesapOlayi).filter(OgrenciHesapOlayi.ogrenci_id == o.id,
                                                              OgrenciHesapOlayi.olay == "giris").count(),
        },
        "turlar": tur_listesi,
        "oneriler": oneriler, "sonuc_notu": sonuc_notu, "hedef": hedef,
        "kayitlar": _ogrenci_kayitlari(db, o, katmanlar),
    }


def _ogrenci_kayitlari(db: Session, o: Ogrenci, katmanlar) -> list[dict]:
    k = [{"zaman": e.zaman, "olay": e.olay, "etiket": OLAY_ETIKET.get(e.olay, e.olay), "aciklama": e.aciklama,
          "yapan": e.yapan, "tur": "hesap"}
         for e in db.query(OgrenciHesapOlayi).filter(OgrenciHesapOlayi.ogrenci_id == o.id)
         .order_by(OgrenciHesapOlayi.zaman.desc()).limit(300).all()]
    for x in katmanlar:
        if x["baslama_zamani"]:
            k.append({"zaman": x["baslama_zamani"], "olay": "katman_basladi", "etiket": OLAY_ETIKET["katman_basladi"],
                      "aciklama": f"{x['kod']} — {x['ad']}", "yapan": "Öğrenci", "tur": "test"})
        if x["tamamlanma_zamani"]:
            k.append({"zaman": x["tamamlanma_zamani"], "olay": "katman_bitti", "etiket": OLAY_ETIKET["katman_bitti"],
                      "aciklama": f"{x['kod']} — {x['ad']}", "yapan": "Öğrenci", "tur": "test"})
    if not any(x["olay"] == "hesap_acildi" for x in k):
        k.append({"zaman": o.olusturulma_zamani, "olay": "hesap_acildi", "etiket": OLAY_ETIKET["hesap_acildi"],
                  "aciklama": None, "yapan": None, "tur": "hesap"})
    _z = lambda x: x["zaman"].timestamp() if x["zaman"] else 0
    return sorted(k, key=_z, reverse=True)


@router.put("/ogrenci/{ogrenci_id}")
def ogrenci_duzenle(ogrenci_id: str, istek: OgrenciDuzenleIstek, db: Session = Depends(get_db),
                    yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    o = _ogrenci_kapsami(db, yon, ogrenci_id)
    degisen = []
    if istek.ad_soyad is not None and istek.ad_soyad.strip() != o.ad_soyad:
        if len(istek.ad_soyad.strip()) < 3:
            raise HTTPException(status_code=400, detail="Ad soyad eksik.")
        o.ad_soyad = re.sub(r"\s+", " ", istek.ad_soyad).strip()
        degisen.append("ad soyad")
    if istek.email is not None and istek.email.strip().lower() != o.email.lower():
        e = istek.email.strip()
        if not _EPOSTA.match(e):
            raise HTTPException(status_code=400, detail="E-posta geçersiz.")
        if _eposta_kullanimda(db, e, haric=o.id):
            raise HTTPException(status_code=400, detail="Bu e-posta ile zaten hesap var.")
        o.email = e
        degisen.append("e-posta")
    if istek.sinif is not None or istek.sube is not None:
        sinif, sube = sinif_ayir(istek.sinif if istek.sinif is not None else o.sinif,
                                 istek.sube if istek.sube is not None else o.sube)
        if (sinif, sube) != (o.sinif, o.sube):
            o.sinif, o.sube = sinif, sube
            degisen.append("sınıf/şube")
    if degisen:
        olay_yaz(db, o.id, "bilgi_guncellendi", ", ".join(degisen), yapan_etiketi(yon))
        denetim_yaz(db, yon, "ogrenci_duzenle", "ogrenciler", o.id, f"{o.ad_soyad}: {', '.join(degisen)}", o.okul_id)
        db.commit()
    return {"tamam": True}


@router.post("/ogrenci/{ogrenci_id}/okul")
def ogrenci_okul_ata(ogrenci_id: str, istek: OkulAtaIstek, db: Session = Depends(get_db),
                     yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    _super(yon)
    o = _ogrenci_kapsami(db, yon, ogrenci_id)
    hedef = _okul_kapsami(db, yon, istek.okul_id)
    eski = o.okul or "Okul harici"
    o.okul_id, o.okul = (hedef.id, hedef.ad) if hedef else (None, None)
    olay_yaz(db, o.id, "okul_degisti", f"{eski} → {hedef.ad if hedef else 'Okul harici'}", yapan_etiketi(yon))
    denetim_yaz(db, yon, "ogrenci_okul_degistir", "ogrenciler", o.id,
                f"{o.ad_soyad}: {eski} → {hedef.ad if hedef else 'Okul harici'}", hedef.id if hedef else None)
    db.commit()
    return {"okul_id": o.okul_id or 0, "okul_ad": o.okul or "Okul harici"}


@router.post("/ogrenciler/sifre-sifirla")
def toplu_sifre_sifirla(istek: IdlerIstek, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    if not istek.idler:
        raise HTTPException(status_code=400, detail="Öğrenci seçilmedi.")
    ogrenciler = [_ogrenci_kapsami(db, yon, i) for i in istek.idler[:EN_FAZLA_SATIR]]
    yapan, kayitlar = yapan_etiketi(yon), []
    from app.models import GuvenilirCihaz
    for o in ogrenciler:
        sifre = gecici_sifre()
        o.sifre_hash = _hizli_hash(sifre)
        o.sifre_degistirmeli = True
        db.query(GuvenilirCihaz).filter(GuvenilirCihaz.kullanici_tipi == "ogrenci", GuvenilirCihaz.kullanici_id == o.id).delete()
        olay_yaz(db, o.id, "sifre_sifirlandi", "Yeni geçici şifre verildi", yapan)
        kayitlar.append({"id": str(o.id), "ad_soyad": o.ad_soyad, "email": o.email, "sinif": o.sinif, "sube": o.sube,
                         "gecici_sifre": sifre})
    okul_idleri = {o.okul_id for o in ogrenciler}
    denetim_yaz(db, yon, "ogrenci_sifre_sifirla", "ogrenciler", ",".join(str(o.id) for o in ogrenciler[:5]),
                f"{len(ogrenciler)} öğrenci: " + ", ".join(o.ad_soyad for o in ogrenciler[:10]) + ("…" if len(ogrenciler) > 10 else ""),
                okul_idleri.pop() if len(okul_idleri) == 1 else None)
    db.commit()
    okul = db.get(Okul, ogrenciler[0].okul_id) if ogrenciler[0].okul_id else None
    return {"kayitlar": kayitlar, "dosya_adi": f"yeni_sifreler_{simdi():%Y%m%d_%H%M}.xlsx",
            "icerik_base64": _giris_listesi_xlsx(okul, kayitlar)}


@router.post("/ogrenciler/sil")
def toplu_sil(istek: IdlerIstek, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    if not istek.idler:
        raise HTTPException(status_code=400, detail="Öğrenci seçilmedi.")
    ogrenciler = [_ogrenci_kapsami(db, yon, i) for i in istek.idler[:EN_FAZLA_SATIR]]
    gruplar = defaultdict(list)
    for o in ogrenciler:
        gruplar[o.okul_id].append(f"{o.ad_soyad} <{o.email}>")
        ogrenciyi_sil(db, o)
    for okul_id, adlar in gruplar.items():
        denetim_yaz(db, yon, "ogrenci_sil", "ogrenciler", okul_id or 0,
                    f"{len(adlar)} öğrenci silindi: " + "; ".join(adlar[:20]) + ("…" if len(adlar) > 20 else ""), okul_id)
    db.commit()
    return {"silinen": len(ogrenciler)}


# ----------------------------------------------------------------------------- okul yetkilileri
def _yetkili_out(y: AdminKullanici) -> dict:
    return {"id": str(y.id), "ad_soyad": y.ad_soyad, "email": y.email, "olusturulma_zamani": y.olusturulma_zamani,
            "son_giris_zamani": y.son_giris_zamani, "sifre_degistirmeli": bool(y.sifre_degistirmeli), "aktif_mi": y.aktif_mi}


@router.get("/okul/{okul_id}/yetkililer")
def yetkilileri_listele(okul_id: int, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    _okul_kapsami(db, yon, okul_id)
    return [_yetkili_out(y) for y in db.query(AdminKullanici).filter(AdminKullanici.okul_id == okul_id)
            .order_by(AdminKullanici.ad_soyad).all()]


@router.post("/okul/{okul_id}/yetkililer", status_code=201)
def yetkili_ekle(okul_id: int, istek: YetkiliEkleIstek, db: Session = Depends(get_db),
                 yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    _super(yon)
    okul = _okul_kapsami(db, yon, okul_id)
    if okul is None:
        raise HTTPException(status_code=400, detail="Okul yetkilisi bir okula bağlı olmalı.")
    ad, email = re.sub(r"\s+", " ", istek.ad_soyad).strip(), istek.email.strip()
    if len(ad) < 3:
        raise HTTPException(status_code=400, detail="Ad soyad eksik.")
    if not _EPOSTA.match(email):
        raise HTTPException(status_code=400, detail="E-posta geçersiz.")
    if _eposta_kullanimda(db, email):
        raise HTTPException(status_code=400, detail="Bu e-posta ile zaten hesap var.")
    sifre = gecici_sifre()
    y = AdminKullanici(ad_soyad=ad, email=email, sifre_hash=sifre_hashle(sifre), rol="okul_yetkilisi", okul_id=okul.id,
                       aktif_mi=True, sifre_degistirmeli=True, olusturulma_zamani=simdi())
    db.add(y)
    db.flush()
    denetim_yaz(db, yon, "okul_yetkilisi_ekle", "admin_kullanicilar", y.id, f"{ad} <{email}>", okul.id)
    db.commit()
    return {**_yetkili_out(y), "gecici_sifre": sifre}


def _yetkili(db: Session, yetkili_id: str) -> AdminKullanici:
    try:
        y = db.get(AdminKullanici, uuid.UUID(yetkili_id))
    except ValueError:
        y = None
    if y is None or y.rol != "okul_yetkilisi":
        raise HTTPException(status_code=404, detail="Okul yetkilisi bulunamadı.")
    return y


@router.post("/yetkili/{yetkili_id}/sifre-sifirla")
def yetkili_sifre_sifirla(yetkili_id: str, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    _super(yon)
    y = _yetkili(db, yetkili_id)
    sifre = gecici_sifre()
    y.sifre_hash = sifre_hashle(sifre)
    y.sifre_degistirmeli = True
    from app.models import GuvenilirCihaz
    db.query(GuvenilirCihaz).filter(GuvenilirCihaz.kullanici_tipi == "yonetim", GuvenilirCihaz.kullanici_id == y.id).delete()
    denetim_yaz(db, yon, "okul_yetkilisi_sifre_sifirla", "admin_kullanicilar", y.id, y.ad_soyad, y.okul_id)
    db.commit()
    return {"gecici_sifre": sifre}


@router.delete("/yetkili/{yetkili_id}", status_code=204)
def yetkili_sil(yetkili_id: str, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    _super(yon)
    y = _yetkili(db, yetkili_id)
    denetim_yaz(db, yon, "okul_yetkilisi_sil", "admin_kullanicilar", y.id, f"{y.ad_soyad} <{y.email}>", y.okul_id)
    db.delete(y)
    db.commit()


# ----------------------------------------------------------------------------- okul kayıtları (log)
ISLEM_ETIKET = {
    "ogrenci_toplu_ekle": "Öğrenci hesapları açıldı", "ogrenci_sil": "Öğrenci silindi",
    "ogrenci_sifre_sifirla": "Öğrenci şifresi sıfırlandı", "ogrenci_duzenle": "Öğrenci bilgisi düzeltildi",
    "ogrenci_okul_degistir": "Öğrenci okulu değiştirildi", "okul_yetkilisi_ekle": "Okul yetkilisi eklendi",
    "okul_yetkilisi_sil": "Okul yetkilisi silindi", "okul_yetkilisi_sifre_sifirla": "Okul yetkilisinin şifresi sıfırlandı",
    "okul_ekle": "Okul eklendi", "okul_guncelle": "Okul bilgisi güncellendi", "okul_sil": "Okul silindi",
    "sifre_belirledi": "Yetkili kendi şifresini belirledi",
    "okul_bilgi_guncelle": "Okul tanıtım bilgileri güncellendi",
}


@router.get("/okul/{okul_id}/kayitlar")
def okul_kayitlari(okul_id: int, gun: int = 30, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    _okul_kapsami(db, yon, okul_id)
    sinir = simdi() - timedelta(days=max(1, min(gun, 365)))
    kayit = []
    q = db.query(AuditLog).filter(AuditLog.zaman >= sinir)
    q = q.filter(AuditLog.okul_id.is_(None), AuditLog.hedef_tablo == "ogrenciler") if okul_id == 0 else q.filter(AuditLog.okul_id == okul_id)
    for a in q.order_by(AuditLog.zaman.desc()).limit(500).all():
        kayit.append({"zaman": a.zaman, "tur": "yonetim", "etiket": ISLEM_ETIKET.get(a.islem, a.islem),
                      "aciklama": a.gerekce, "yapan": a.yapan_ad or "—", "ogrenci_id": None})
    kosul, p = _okul_filtresi(okul_id)
    for r in db.execute(text(f"""
        SELECT e.zaman, e.olay, e.aciklama, e.yapan, o.ad_soyad, o.id
          FROM ogrenci_hesap_olaylari e JOIN ogrenciler o ON o.id = e.ogrenci_id
         WHERE {kosul} AND e.zaman >= :sinir AND e.olay IN ('giris','giris_hatali','sifre_degisti')
         ORDER BY e.zaman DESC LIMIT 1000"""), {**p, "sinir": sinir}).mappings().all():
        kayit.append({"zaman": r["zaman"], "tur": "ogrenci", "etiket": OLAY_ETIKET.get(r["olay"], r["olay"]),
                      "aciklama": r["aciklama"], "yapan": r["ad_soyad"], "ogrenci_id": str(r["id"])})
    kayit.sort(key=lambda x: x["zaman"].timestamp() if x["zaman"] else 0, reverse=True)
    return kayit[:1200]


# ----------------------------------------------------------------------------- okul tanıtım bilgileri
class KadroGirdi(BaseModel):
    gorev: str
    ad: str
    eposta: str | None = None
    telefon: str | None = None


class OkulBilgiIstek(BaseModel):
    kurulus_yili: int | None = None
    ogrenci_sayisi: int | None = None
    tanitim: str | None = None
    adres: str | None = None
    telefon: str | None = None
    eposta: str | None = None
    web: str | None = None
    kadro: list[KadroGirdi] = []


GOREVLER = ["Okul Müdürü", "Müdür Başyardımcısı", "Müdür Yardımcısı", "Rehber Öğretmen / Psikolojik Danışman",
            "Psikolog", "Kariyer Danışmanı", "Sınıf Öğretmeni", "Öğretmen", "Okul Sekreteri", "Diğer"]


def _bilgi_out(okul: Okul) -> dict:
    return {"id": okul.id, "ad": okul.ad, "alt_baslik": okul.alt_baslik, "logo": okul.logo,
            "kurulus_yili": okul.kurulus_yili, "ogrenci_sayisi": okul.ogrenci_sayisi, "tanitim": okul.tanitim,
            "adres": okul.adres, "telefon": okul.telefon, "eposta": okul.eposta, "web": okul.web,
            "kadro": okul.kadro or [], "bilgi_guncelleme_zamani": okul.bilgi_guncelleme_zamani, "gorevler": GOREVLER,
            "tema_renk": okul.tema_renk}


def _kirp(v: str | None, n: int) -> str | None:
    v = (v or "").strip()
    return v[:n] or None


@router.get("/okul/{okul_id}/bilgi")
def okul_bilgi(okul_id: int, db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    okul = _okul_kapsami(db, yon, okul_id)
    if okul is None:
        raise HTTPException(status_code=400, detail="Okul harici grubun tanıtım bilgisi yok.")
    return _bilgi_out(okul)


@router.put("/okul/{okul_id}/bilgi")
def okul_bilgi_guncelle(okul_id: int, istek: OkulBilgiIstek, db: Session = Depends(get_db),
                        yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    okul = _okul_kapsami(db, yon, okul_id)
    if okul is None:
        raise HTTPException(status_code=400, detail="Okul harici grubun tanıtım bilgisi yok.")
    yil = simdi().year
    if istek.kurulus_yili is not None and not (1800 <= istek.kurulus_yili <= yil):
        raise HTTPException(status_code=400, detail=f"Kuruluş yılı 1800–{yil} arasında olmalı.")
    if istek.ogrenci_sayisi is not None and not (0 <= istek.ogrenci_sayisi <= 50000):
        raise HTTPException(status_code=400, detail="Öğrenci sayısı geçersiz.")
    if len(istek.tanitim or "") > 3000:
        raise HTTPException(status_code=400, detail="Tanıtım metni en fazla 3000 karakter olabilir.")
    eposta = _kirp(istek.eposta, 120)
    if eposta and not _EPOSTA.match(eposta):
        raise HTTPException(status_code=400, detail="Okulun e-posta adresi geçersiz.")
    web = _kirp(istek.web, 200)
    if web and not re.match(r"^https?://", web):
        web = "https://" + web
    if len(istek.kadro) > 40:
        raise HTTPException(status_code=400, detail="Kadroya en fazla 40 kişi eklenebilir.")
    kadro = []
    for i, k in enumerate(istek.kadro, start=1):
        ad, gorev = _kirp(k.ad, 80), _kirp(k.gorev, 60)
        if not ad and not (k.eposta or "").strip():
            continue  # boş satır
        if not ad or not gorev:
            raise HTTPException(status_code=400, detail=f"Kadro {i}. satır: görev ve ad soyad gerekli.")
        e = _kirp(k.eposta, 120)
        if e and not _EPOSTA.match(e):
            raise HTTPException(status_code=400, detail=f"Kadro {i}. satır ({ad}): e-posta geçersiz.")
        kadro.append({"gorev": gorev, "ad": ad, "eposta": e, "telefon": _kirp(k.telefon, 30)})
    okul.kurulus_yili, okul.ogrenci_sayisi = istek.kurulus_yili, istek.ogrenci_sayisi
    okul.tanitim = _kirp(istek.tanitim, 3000)
    okul.adres, okul.telefon, okul.eposta, okul.web = _kirp(istek.adres, 300), _kirp(istek.telefon, 30), eposta, web
    okul.kadro = kadro
    okul.bilgi_guncelleme_zamani = simdi()
    denetim_yaz(db, yon, "okul_bilgi_guncelle", "okullar", okul.id, f"{okul.ad}: tanıtım/iletişim, kadroda {len(kadro)} kişi", okul.id)
    db.commit()
    db.refresh(okul)
    return _bilgi_out(okul)


# ----------------------------------------------------------------------------- okul rengi (görünüm)
_HEX = re.compile(r"^#[0-9A-Fa-f]{6}$")


class TemaIstek(BaseModel):
    renk: str | None = None          # "#1F3A93" — boş/None = Filizyol'un kendi rengi


@router.put("/okul/{okul_id}/tema")
def okul_tema_guncelle(okul_id: int, istek: TemaIstek, db: Session = Depends(get_db),
                       yon: AdminKullanici = Depends(get_mevcut_yonetim)):
    """[2026-10-09] Okul rengi: öğrenci ve okul paneli arayüzündeki ayırıcı/vurgu çizgilerinin rengi."""
    okul = _okul_kapsami(db, yon, okul_id)
    if okul is None:
        raise HTTPException(status_code=400, detail="Okul harici grubun okul rengi yok.")
    renk = (istek.renk or "").strip() or None
    if renk and not _HEX.match(renk):
        raise HTTPException(status_code=400, detail="Renk #RRGGBB biçiminde olmalı (ör. #1F3A93).")
    okul.tema_renk = renk.upper() if renk else None
    denetim_yaz(db, yon, "okul_tema_guncelle", "okullar", okul.id, f"{okul.ad}: okul rengi {okul.tema_renk or 'varsayılan'}", okul.id)
    db.commit()
    return {"tema_renk": okul.tema_renk}
