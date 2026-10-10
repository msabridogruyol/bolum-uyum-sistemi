# -*- coding: utf-8 -*-
"""
[2026-10-10] Okul karşılaştırması (yalnızca süper admin) — birden çok okulu olan kurumlara rapor vermek için.
Ayrı bir "kurum yöneticisi" rolü yoktur; karşılaştırma süper adminde, Excel olarak kuruma iletilebilir.

GET /yonetim/okul-karsilastirma            — tüm okullar, metrikler
GET /yonetim/okul-karsilastirma/excel?okullar=1,2,3
"""
import io
from datetime import date, datetime, timedelta, timezone

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.api.deps import get_mevcut_super_admin
from app.api.okul_yonetimi import _durum, _ilerleme
from app.core.database import get_db
from app.core.hesap_yonetimi import denetim_yaz
from app.core.kucuk_grup import DIPNOT, EN_AZ_GRUP, GIZLI_METIN, gizle, yeterli
from app.core.paketler import okul_modulleri
from app.models import AdminKullanici, Ogrenci, Okul

router = APIRouter(prefix="/yonetim", tags=["Okul karşılaştırması"])

METRIKLER = [
    {"k": "ogrenci", "ad": "Öğrenci", "tur": "sayi"},
    {"k": "giris", "ad": "Giriş yapan", "tur": "yuzde"},
    {"k": "aktif30", "ad": "Son 30 günde aktif", "tur": "yuzde"},
    {"k": "tamamlama", "ad": "Testi tamamlayan", "tur": "yuzde"},
    {"k": "hedef", "ad": "Hedef bölüm seçen", "tur": "yuzde"},
    {"k": "tyt", "ad": "Ort. son TYT neti", "tur": "net", "modul": "net_takibi"},
    {"k": "gorusme", "ad": "Rehberlik görüşmesi (90 gün)", "tur": "sayi", "modul": "rehberlik"},
    {"k": "risk", "ad": "Yüksek uyarılı öğrenci", "tur": "sayi", "modul": "rehberlik"},
    {"k": "yerlesme", "ad": "Yerleşme (son yıl)", "tur": "yuzde", "modul": "mezun_takibi"},
]


def _okul_metrikleri(db: Session, okul: Okul) -> dict:
    ogr = [o for o in db.query(Ogrenci).filter(Ogrenci.okul_id == okul.id).all() if not getattr(o, "test_hesabi", False)]
    n = len(ogr)
    il = _ilerleme(db, okul.id)
    sinir = datetime.now(timezone.utc) - timedelta(days=30)
    oran = lambda x: round(100 * x / n) if n else None  # noqa: E731
    moduller = set(okul_modulleri(db, okul.id))
    r_gizli = False
    m = {
        "ogrenci": n,
        "giris": oran(sum(1 for o in ogr if o.son_giris_zamani is not None)),
        "aktif30": oran(sum(1 for o in ogr if o.son_giris_zamani and (o.son_giris_zamani if o.son_giris_zamani.tzinfo else o.son_giris_zamani.replace(tzinfo=timezone.utc)) >= sinir)),
        "tamamlama": oran(sum(1 for o in ogr if _durum(o, il.get(o.id) or {})[0] == "tamamlandi")),
        "hedef": oran(sum(1 for o in ogr if (il.get(o.id) or {}).get("hedef"))),
        "tyt": None, "gorusme": None, "risk": None, "yerlesme": None,
    }
    ids = [o.id for o in ogr]
    if ids and "net_takibi" in moduller:
        v, vn = db.execute(text("""
            SELECT avg(toplam_net), count(*) FROM (
                SELECT DISTINCT ON (ogrenci_id) toplam_net FROM ogrenci_denemeleri
                 WHERE ogrenci_id = ANY(:i) AND oturum = 'TYT' AND tarih >= :b ORDER BY ogrenci_id, tarih DESC, id DESC) x
        """), {"i": ids, "b": date.today() - timedelta(days=90)}).one()
        m["tyt"] = gizle(round(float(v), 1), vn) if v is not None else None   # [2026-10-10] 5'ten az deneme girene dayanıyorsa gizli
        m["tyt_n"] = vn
    if "rehberlik" in moduller:
        try:
            m["gorusme"] = db.execute(text("SELECT count(*) FROM rehberlik_gorusmeleri WHERE okul_id = :o AND durum = 'yapildi' AND zaman >= now() - interval '90 days'"),
                                      {"o": okul.id}).scalar() or 0
            from app.api.rehberlik import uyarilari_hesapla
            r = sum(1 for s in uyarilari_hesapla(db, okul.id) if s["seviye"] == "yuksek")
            r_gizli = not (r == 0 or yeterli(r))
            m["risk"] = None if r_gizli else r   # [2026-10-10] 1-4 kişilik hücre → "5'ten az öğrenci"
        except Exception:
            db.rollback()
    if "mezun_takibi" in moduller:
        r = db.execute(text("""SELECT yil, count(*) AS n, count(*) FILTER (WHERE durum = 'yerlesti') AS y FROM mezun_yerlesmeleri
                               WHERE okul_id = :o GROUP BY yil ORDER BY yil DESC LIMIT 1"""), {"o": okul.id}).first()
        if r and r.n:
            m["yerlesme"] = gizle(round(100 * r.y / r.n), r.n)   # [2026-10-10] 5'ten az mezun kaydı → gizli
            m["yerlesme_yil"] = r.yil
            m["yerlesme_n"] = r.n
    # [2026-10-10] KVKK küçük grup: 5'ten az öğrencili okulun metrikleri gösterilmez (öğrenci sayısı kalır)
    m["gizli"] = not yeterli(n)
    gizli = {k for k in ("tyt", "yerlesme") if m.get(k) is None and 0 < (m.get(k + "_n") or 0) < EN_AZ_GRUP}
    if m["risk"] is None and r_gizli:
        gizli.add("risk")
    if m["gizli"]:
        for k in list(m):
            if k not in ("ogrenci", "gizli"):
                m[k] = None
        gizli = {x["k"] for x in METRIKLER if x["k"] != "ogrenci"}
    m["gizli_alanlar"] = sorted(gizli)   # arayüz bu hücrelerde "5'ten az öğrenci" yazar
    return m


def _veri(db: Session) -> dict:
    paket = dict(db.execute(text("SELECT kod, ad FROM paketler")).all())
    okullar = []
    for ok in db.query(Okul).order_by(Okul.ad).all():
        okullar.append({"id": ok.id, "ad": ok.ad, "logo": ok.logo, "paket": paket.get(getattr(ok, "paket", None), getattr(ok, "paket", None)),
                        "moduller": okul_modulleri(db, ok.id), **_okul_metrikleri(db, ok)})
    return {"okullar": okullar, "metrikler": METRIKLER,
            "kucuk_grup": {"esik": EN_AZ_GRUP, "metin": GIZLI_METIN, "dipnot": DIPNOT}}


@router.get("/okul-karsilastirma")
def karsilastirma(db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_super_admin)):
    return _veri(db)


@router.get("/okul-karsilastirma/excel")
def karsilastirma_excel(okullar: str | None = Query(None), db: Session = Depends(get_db), yon: AdminKullanici = Depends(get_mevcut_super_admin)):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    v = _veri(db)
    secili = {int(x) for x in (okullar or "").split(",") if x.strip().isdigit()}
    liste = [o for o in v["okullar"] if not secili or o["id"] in secili]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Karşılaştırma"
    kol = ["Okul", "Paket"] + [m["ad"] + (" (%)" if m["tur"] == "yuzde" else "") for m in METRIKLER]
    for j, k in enumerate(kol, 1):
        c = ws.cell(row=1, column=j, value=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5D8A")
    for i, o in enumerate(liste, 2):
        for j, val in enumerate([o["ad"], o["paket"]] + [GIZLI_METIN if m["k"] in (o.get("gizli_alanlar") or []) else o.get(m["k"]) for m in METRIKLER], 1):
            ws.cell(row=i, column=j, value=val)
    ws.cell(row=len(liste) + 3, column=1, value=f"Filizyol · {date.today().strftime('%d.%m.%Y')} · Boş hücre: modül okulun paketinde yok, veri yok "
                                                f"ya da {GIZLI_METIN} içeriyor.").font = Font(italic=True, color="7A5C00")
    ws.cell(row=len(liste) + 4, column=1, value=DIPNOT).font = Font(italic=True, color="7A5C00")
    for j in range(1, len(kol) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 28 if j == 1 else 16
    b = io.BytesIO()
    wb.save(b)
    denetim_yaz(db, yon, "okul_karsilastirma_excel", "okullar", ",".join(str(o["id"]) for o in liste), f"{len(liste)} okul")
    db.commit()
    return Response(b.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="okul_karsilastirma_{date.today().isoformat()}.xlsx"', "Cache-Control": "no-store"})
