# -*- coding: utf-8 -*-
"""
[2026-10-11] Veli toplantısı sunumu (.pptx, python-pptx). Veri: app/api/veli_sunumu.veli_sunumu_verisi.

16:9, Arial (Office'te ve LibreOffice'te — Liberation Sans — aynı genişlikte). Vurgu rengi okulun tema rengi; okulun logosu
PNG/JPEG ise kapakta yer alır (SVG logo eklenmez). Grafikler PowerPoint'in kendi grafikleridir (düzenlenebilir).
Her slaytın notlar bölümünde öğretmen için kısa konuşma notu vardır.

Gizlilik veri katmanında uygulanır (öğrenci bazlı veri buraya hiç gelmez; 5'ten az öğrenciye dayanan değerler None gelir).
Bu dosya None değerleri "—" yazar ya da ilgili grafiği hiç çizmez.
"""
from __future__ import annotations

import base64
import io
from datetime import date, datetime

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

FONT = "Arial"
MUREKKEP = "2B2620"
GRI = "6B6257"
BEYAZ = "FFFFFF"
AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
W, H = 13.333, 7.5
SOL = 0.6
GEN = W - 2 * SOL


# ============================================================================= renk
def _hex(h: str | None, varsayilan="E8804A") -> str:
    h = (h or "").strip().lstrip("#")
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    try:
        int(h, 16)
        return h.upper() if len(h) == 6 else varsayilan
    except ValueError:
        return varsayilan


def _karistir(a: str, b: str, t: float) -> str:
    """a'dan b'ye t oranında."""
    x = [int(a[i:i + 2], 16) for i in (0, 2, 4)]
    y = [int(b[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(p + (q - p) * t):02X}" for p, q in zip(x, y))


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def _parlaklik(h: str) -> float:
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


class _Renk:
    def __init__(self, ana: str):
        ana = _hex(ana)
        # çok açık okul renklerinde beyaz yazı okunmaz → kapak için koyulaştır
        self.ana = ana if _parlaklik(ana) < 0.62 else _karistir(ana, "000000", 0.35)
        self.koyu = _karistir(self.ana, "000000", 0.35)
        self.acik = _karistir(self.ana, "FFFFFF", 0.88)
        self.orta = _karistir(self.ana, "FFFFFF", 0.45)
        self.seri = [self.ana, _karistir(self.ana, "FFFFFF", 0.35), "8C8378", _karistir(self.ana, "000000", 0.3),
                     _karistir(self.ana, "FFFFFF", 0.6), "C9C1B5", "4E8A4F"]


# ============================================================================= biçim yardımcıları
def _tr(x, kesir=0) -> str:
    if x is None:
        return "—"
    s = f"{x:.{kesir}f}"
    if kesir:
        s = s.rstrip("0").rstrip(".")
    return s.replace(".", ",")


def _yuzde(x) -> str:
    return "—" if x is None else f"%{_tr(x)}"


def _tarih(d) -> str:
    if isinstance(d, datetime):
        d = d.date()
    if not isinstance(d, date):
        return str(d or "")
    return f"{d.day} {AYLAR[d.month - 1]} {d.year}"


def _metin(sl, x, y, w, h, parcalar, boyut=16, renk=MUREKKEP, kalin=False, hiza=PP_ALIGN.LEFT, dikey=MSO_ANCHOR.TOP,
           satir_arasi=None, ad=None):
    """parcalar: str | [paragraf] — paragraf: str ya da [(metin, {boyut, renk, kalin, italik})]."""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if ad:
        tb.name = ad
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = dikey
    paragraflar = parcalar if isinstance(parcalar, list) else [parcalar]
    for i, par in enumerate(paragraflar):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = hiza
        if satir_arasi:
            p.space_after = Pt(satir_arasi)
        runs = par if isinstance(par, list) else [(par, {})]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("boyut", boyut))
            f.bold = o.get("kalin", kalin)
            f.italic = o.get("italik", False)
            f.color.rgb = _rgb(o.get("renk", renk))
    return tb


def _kart(sl, x, y, w, h, dolgu, ad=None, yuvarlak=0.06):
    s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = yuvarlak
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(dolgu)
    s.line.fill.background()
    s.shadow.inherit = False
    if ad:
        s.name = ad
    return s


def _daire(sl, x, y, cap, dolgu, yazi, yazi_renk=BEYAZ, boyut=16):
    s = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(cap), Inches(cap))
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(dolgu)
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = yazi
    r.font.name, r.font.size, r.font.bold = FONT, Pt(boyut), True
    r.font.color.rgb = _rgb(yazi_renk)
    return s


def _arka(sl, renk):
    bg = sl.background.fill
    bg.solid()
    bg.fore_color.rgb = _rgb(renk)


def _not(sl, metin: str):
    sl.notes_slide.notes_text_frame.text = metin


class _Deste:
    def __init__(self, v: dict):
        self.v = v
        self.r = _Renk(v["okul"].get("renk"))
        self.p = Presentation()
        self.p.slide_width, self.p.slide_height = Inches(W), Inches(H)
        self.bos = self.p.slide_layouts[6]
        self.no = 0
        self.alt = f"{v['okul']['ad']} · {v['kapsam']} · Veli toplantısı"

    # ---------------------------------------------------------------- ortak iskelet
    def slayt(self, ust: str, baslik: str, koyu=False):
        sl = self.p.slides.add_slide(self.bos)
        self.no += 1
        _arka(sl, self.r.ana if koyu else BEYAZ)
        _metin(sl, SOL, 0.42, GEN, 0.3, ust.upper(), 12, BEYAZ if koyu else self.r.ana, kalin=True, ad="Üst başlık")
        _metin(sl, SOL, 0.72, GEN, 0.8, baslik, 32, BEYAZ if koyu else MUREKKEP, kalin=True, ad="Başlık")
        renk = _karistir(self.r.ana, "FFFFFF", 0.7) if koyu else GRI
        _metin(sl, SOL, H - 0.5, GEN - 1, 0.25, self.alt, 10, renk, ad="Alt bilgi")
        _metin(sl, W - SOL - 1, H - 0.5, 1, 0.25, str(self.no), 10, renk, hiza=PP_ALIGN.RIGHT, ad="Slayt no")
        return sl

    def stat(self, sl, x, y, w, deger: str, etiket: str, dolgu=None, h=1.55):
        _kart(sl, x, y, w, h, dolgu or self.r.acik)
        _metin(sl, x + 0.25, y + 0.18, w - 0.5, 0.8, deger, 40, self.r.koyu, kalin=True)
        _metin(sl, x + 0.25, y + 0.98, w - 0.5, h - 1.05, etiket, 13, GRI)

    def grafik_bicim(self, ch, baslik: str | None = None, yuzde=False, eksen_max=None):
        ch.font.name = FONT
        ch.font.size = Pt(12)
        ch.font.color.rgb = _rgb(MUREKKEP)
        if baslik:
            ch.has_title = True
            ch.chart_title.text_frame.text = baslik
            f = ch.chart_title.text_frame.paragraphs[0].runs[0].font
            f.name, f.size, f.bold = FONT, Pt(14), True
            f.color.rgb = _rgb(MUREKKEP)
        else:
            ch.has_title = False
        try:
            va = ch.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = _rgb("E6DDCF")
            va.format.line.fill.background()
            va.tick_labels.font.size = Pt(11)
            va.tick_labels.font.color.rgb = _rgb(GRI)
            va.minimum_scale = 0
            if eksen_max:
                va.maximum_scale = eksen_max
            ca = ch.category_axis
            ca.tick_labels.font.size = Pt(12)
            ca.tick_labels.font.color.rgb = _rgb(MUREKKEP)
            ca.format.line.color.rgb = _rgb("C9C1B5")
            ca.has_major_gridlines = False
        except (ValueError, AttributeError):
            pass

    def cubuk(self, sl, x, y, w, h, kategoriler, seriler: list[tuple[str, list]], yatay=True, eksen_max=None, sonek=""):
        cd = CategoryChartData()
        cd.categories = kategoriler
        for ad, degerler in seriler:
            cd.add_series(ad, degerler)
        tur = XL_CHART_TYPE.BAR_CLUSTERED if yatay else XL_CHART_TYPE.COLUMN_CLUSTERED
        ch = sl.shapes.add_chart(tur, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
        self.grafik_bicim(ch, eksen_max=eksen_max)
        if yatay:
            ch.category_axis.reverse_order = True     # ilk kategori üstte
            try:
                from pptx.enum.chart import XL_TICK_LABEL_POSITION
                ch.value_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
            except Exception:
                pass
        pl = ch.plots[0]
        pl.gap_width = 60
        pl.overlap = -10 if len(seriler) > 1 else 0
        for i, s in enumerate(pl.series):
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = _rgb(self.r.ana if i == 0 else "B8AFA2")
            s.invert_if_negative = False
        pl.has_data_labels = True
        dl = pl.data_labels
        dl.font.size, dl.font.name, dl.font.bold = Pt(12), FONT, True
        dl.font.color.rgb = _rgb(MUREKKEP)
        dl.number_format = f'0"{sonek}"' if sonek else "0"
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        ch.has_legend = len(seriler) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
            ch.legend.font.size, ch.legend.font.name = Pt(12), FONT
        return ch

    def halka(self, sl, x, y, w, h, kategoriler, degerler, baslik=None):
        cd = CategoryChartData()
        cd.categories = kategoriler
        cd.add_series("Öğrenci", degerler)
        ch = sl.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
        self.grafik_bicim(ch, baslik)
        pl = ch.plots[0]
        for i, pt in enumerate(pl.series[0].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = _rgb(self.r.seri[i % len(self.r.seri)])
            pt.format.line.color.rgb = _rgb(BEYAZ)
        pl.has_data_labels = True
        dl = pl.data_labels
        # yüzde yerine öğrenci sayısı: gizlenen küçük gruplar paydadan düştüğü için yüzde yanıltıcı olur
        dl.show_percentage, dl.show_value, dl.show_category_name = False, True, False
        dl.number_format, dl.number_format_is_linked = "0", False
        dl.font.size, dl.font.name, dl.font.bold = Pt(12), FONT, True
        dl.font.color.rgb = _rgb(BEYAZ)
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.RIGHT
        ch.legend.include_in_layout = False
        ch.legend.font.size, ch.legend.font.name = Pt(12), FONT
        try:   # halka kalınlığı
            from lxml import etree
            hole = pl._element.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}holeSize")
            if hole is not None:
                hole.set("val", "55")
        except Exception:
            pass
        return ch


# ============================================================================= logo
def _logo_akisi(logo: str | None):
    if not logo or not logo.startswith("data:image/") or "svg" in logo[:30]:
        return None
    try:
        return io.BytesIO(base64.b64decode(logo.split(",", 1)[1]))
    except Exception:
        return None


# ============================================================================= slaytlar
def _kapak(d: _Deste):
    v = d.v
    sl = d.p.slides.add_slide(d.bos)
    d.no += 1
    _arka(sl, d.r.ana)
    # sağda iki büyük yumuşak daire (motif: yuvarlak biçimler)
    for (x, y, c, t) in ((8.9, -1.2, 5.6, 0.18), (10.6, 3.6, 4.0, 0.3)):
        s = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(c), Inches(c))
        s.fill.solid()
        s.fill.fore_color.rgb = _rgb(_karistir(d.r.ana, "FFFFFF", t))
        s.line.fill.background()
        s.shadow.inherit = False
    y = 1.3
    logo = _logo_akisi(v["okul"].get("logo"))
    if logo is not None:
        try:
            sl.shapes.add_picture(logo, Inches(SOL), Inches(0.7), height=Inches(0.9))
            y = 1.9
        except Exception:
            pass
    _metin(sl, SOL, y + 0.3, 8, 0.4, "VELİ TOPLANTISI", 14, _karistir(d.r.ana, "FFFFFF", 0.75), kalin=True)
    _metin(sl, SOL, y + 0.8, 8.2, 1.9, v["okul"]["ad"], 44, BEYAZ, kalin=True, dikey=MSO_ANCHOR.TOP)
    _metin(sl, SOL, y + 2.75, 8, 0.9, [
        [("Bölüm ve meslek yönelimi çalışmalarımız", {"boyut": 22, "kalin": True})],
        [(f"{v['kapsam']} · {_tarih(v['tarih'])}", {"boyut": 16, "renk": _karistir(d.r.ana, "FFFFFF", 0.75)})],
    ], 18, BEYAZ, satir_arasi=6)
    _metin(sl, SOL, H - 0.75, 8, 0.3, "Filizyol · Bölüm Uyum ve Kariyer Rehberliği", 11, _karistir(d.r.ana, "FFFFFF", 0.7))
    _not(sl, f"Hoş geldiniz. Bu sunumda {v['kapsam'].lower() if v['kapsam_tur'] == 'okul' else v['kapsam']} öğrencilerimizin "
             "bölüm ve meslek yönelimi çalışmalarını toplu olarak paylaşacağız. Sunumda hiçbir öğrencinin adı ya da kişisel sonucu "
             "yer almıyor; çocuğunuzun kendi raporunu öğrencinin hesabından ya da rehberlik servisinden isteyebilirsiniz.")


def _nedir(d: _Deste, katmanlar: list[tuple[str, str]]):
    sl = d.slayt("Filizyol nedir", "Çocuğunuzu tanımaya yardımcı olan bir yol arkadaşı")
    _metin(sl, SOL, 1.75, 5.3, 2.6, [
        "Filizyol, öğrencilerin kendi değerlerini, çalışma tarzlarını, becerilerini ve ilgi alanlarını tanımalarına "
        "yardımcı olan bir değerlendirme ve rehberlik sistemidir.",
        "Öğrenci soruları kendi hızında yanıtlar; sistem bu yanıtları 300'den fazla üniversite bölümünün gerektirdiği "
        "özelliklerle karşılaştırır ve öğrenciye en uyumlu bölümleri nedenleriyle birlikte gösterir.",
    ], 16, MUREKKEP, satir_arasi=12)
    # katman kartları (2 x 2) + derinleşme
    x0, y0, kw, kh = 6.35, 1.75, 3.1, 1.15
    for i, (kod, ad) in enumerate(katmanlar[:4]):
        x = x0 + (i % 2) * (kw + 0.2)
        y = y0 + (i // 2) * (kh + 0.2)
        _kart(sl, x, y, kw, kh, d.r.acik)
        _daire(sl, x + 0.2, y + 0.3, 0.55, d.r.ana, str(i + 1), boyut=16)
        _metin(sl, x + 0.9, y + 0.15, kw - 1.05, kh - 0.3, ad, 14, MUREKKEP, kalin=True, dikey=MSO_ANCHOR.MIDDLE)
    _kart(sl, x0, y0 + 2 * (kh + 0.2), 2 * kw + 0.2, 0.8, d.r.acik)
    _daire(sl, x0 + 0.2, y0 + 2 * (kh + 0.2) + 0.13, 0.55, d.r.ana, "5", boyut=16)
    _metin(sl, x0 + 0.9, y0 + 2 * (kh + 0.2) + 0.1, 2 * kw - 0.8, 0.6,
           "Alan derinleşmesi: öne çıkan alanlarda ek sorularla daha ayrıntılı bakış", 14, MUREKKEP, dikey=MSO_ANCHOR.MIDDLE)
    # karar desteği uyarısı
    _kart(sl, SOL, 5.45, GEN, 1.2, _karistir(d.r.ana, "FFFFFF", 0.94))
    _metin(sl, SOL + 0.3, 5.6, GEN - 0.6, 0.9, [
        [("Sonuçlar bir karar desteğidir. ", {"kalin": True, "renk": d.r.koyu}),
         ("Bir tanı, test notu ya da kesin hüküm değildir; öğrencinin bugünkü yanıtlarını yansıtır ve zamanla değişebilir. "
          "Son kararı her zaman öğrenci, aile ve rehber öğretmen birlikte verir.", {})],
    ], 15, MUREKKEP, dikey=MSO_ANCHOR.MIDDLE)
    if not d.v["toplu"]:
        _metin(sl, SOL, 6.72, GEN, 0.25, f"Bu kapsamda {d.v['esik']}'ten az öğrenci olduğu için toplu sonuçlar paylaşılmamaktadır "
                                          "(kişisel verilerin korunması).", 11, GRI)
    _not(sl, "Sistemi kısaca tanıtın: öğrenci dört katmanda sorular yanıtlıyor; sonra öne çıkan alanlarda derinleşme soruları geliyor. "
             "Özellikle vurgulayın: bu bir sınav değil, doğru ya da yanlış cevap yok. Sonuçlar bir yön gösterir, kesin karar vermez. "
             "Psikolojik bir tanı aracı da değildir.")


def _katilim(d: _Deste):
    k = d.v["katilim"]
    sl = d.slayt("Katılım", "Öğrencilerimizin değerlendirmeye katılımı")
    d.stat(sl, SOL, 1.75, 3.3, str(k["toplam"]), "öğrenci bu kapsamda")
    d.stat(sl, SOL, 3.5, 3.3, _yuzde(k["tamamlama_orani"]), "değerlendirmeyi tamamladı")
    d.stat(sl, SOL, 5.25, 3.3, _yuzde(k["baslama_orani"]), "değerlendirmeye başladı", h=1.45)
    kir = k.get("kirilim") or []
    gen = 4.2 if kir else 8.3
    dag = k["dagilim"]
    if len(dag) >= 2 and sum(x["sayi"] for x in dag) == k["toplam"]:
        d.halka(sl, 4.25, 1.7, gen, 4.9, [x["ad"] for x in dag], [x["sayi"] for x in dag], "Değerlendirme durumu (öğrenci)")
    elif k["tamamlama_orani"] is not None:
        _ilerleme(d, sl, 4.35, 2.6, gen - 0.2, k["tamamlama_orani"], "Tamamlama")
        _ilerleme(d, sl, 4.35, 4.1, gen - 0.2, k["baslama_orani"], "Başlama")
    if kir:
        d.cubuk(sl, 8.65, 1.95, 4.1, 4.65, [x["ad"] for x in kir], [("Tamamlama %", [x["oran"] for x in kir])],
                yatay=False, eksen_max=100, sonek="%")
        _metin(sl, 8.65, 1.7, 4.1, 0.3, f"{k['kirilim_adi']} bazında tamamlama", 14, MUREKKEP, kalin=True)
        if k.get("kirilim_gizli"):
            _metin(sl, 8.65, 6.62, 4.1, 0.25, f"{d.v['esik']}'ten az öğrencili gruplar gösterilmez.", 10, GRI)
    _not(sl, f"Bu kapsamda {k['toplam']} öğrencimiz var; tamamlama oranı {_yuzde(k['tamamlama_orani'])}. Değerlendirmeyi henüz "
             "bitirmeyen öğrencilerimizi evde nazikçe teşvik edebilirsiniz; acele etmeden, dinlenmiş bir zamanda tamamlamaları en "
             "doğru sonucu verir. Grafikte 5'ten az öğrenciye dayanan gruplar gösterilmez.")


def _ilerleme(d: _Deste, sl, x, y, w, oran, etiket):
    _metin(sl, x, y, w, 0.35, [[(etiket, {"kalin": True}), (f"  {_yuzde(oran)}", {"renk": d.r.koyu, "kalin": True})]], 16)
    _kart(sl, x, y + 0.5, w, 0.45, d.r.acik, yuvarlak=0.5)
    if oran:
        _kart(sl, x, y + 0.5, max(0.45, w * min(100, oran) / 100), 0.45, d.r.ana, yuvarlak=0.5)


def _guclu(d: _Deste):
    g = d.v.get("guclu") or []
    sl = d.slayt("Öne çıkan yönler", "Öğrencilerimizde öne çıkan özellikler")
    if g:
        okul_var = any(x["okul"] is not None for x in g)
        seriler = [(d.v["kapsam"], [x["ort"] for x in g])]
        if okul_var:
            seriler.append(("Okul geneli", [x["okul"] for x in g]))
        d.cubuk(sl, SOL, 1.7, 7.9, 5.0, [x["ad"] for x in g], seriler, eksen_max=100)
    else:
        _metin(sl, SOL, 2.2, 7.9, 1, "Bu kapsamda gösterilecek kadar toplu veri yok.", 16, GRI)
    _kart(sl, 8.85, 1.85, 3.88, 4.75, d.r.acik)
    _metin(sl, 9.15, 2.1, 3.3, 4.3, [
        [("Nasıl okunur?", {"kalin": True, "boyut": 16, "renk": d.r.koyu})],
        "Çubuklar, değerlendirmeyi tamamlayan öğrencilerin bu özelliklerdeki ortalama puanıdır (0–100).",
        "Grubun en belirgin 5 ortak özelliği gösterilir. Her öğrencinin kendi profili farklıdır.",
        [("Öğrenciler kişilik ya da değer puanlarına göre sıralanmaz; ", {"kalin": True}),
         ("bu özellikler iyi–kötü değil, farklı yönlerdir.", {})],
    ], 14, MUREKKEP, satir_arasi=10)
    _not(sl, "Bu grafik grubun ortak eğilimini gösterir, tek tek öğrencileri değil. Bu özelliklerin hiçbiri 'iyi' ya da 'kötü' "
             "değildir. Çocuğunuzun kendi güçlü yönlerini onun raporunda birlikte okumanızı öneririz. "
             + ("Gri çubuklar okul geneli ortalamasıdır." if any(x["okul"] is not None for x in g) else ""))


def _alanlar(d: _Deste):
    a = [x for x in (d.v.get("alanlar") or []) if x.get("sayi")]
    if not any(not x.get("diger") for x in a):
        a = []
    b = d.v.get("bolumler") or []
    sl = d.slayt("Alanlar ve bölümler", "En çok önerilen alanlar ve bölümler")
    if a:
        d.halka(sl, SOL, 1.65, 5.6, 5.0, [x["ad"] for x in a], [x["sayi"] for x in a], "En uyumlu bölümün alanı (öğrenci)")
    if b:
        _metin(sl, 6.6, 1.7, 6.1, 0.3, "Öğrencilerin ilk 10 önerisinde en sık yer alan bölümler", 14, MUREKKEP, kalin=True)
        d.cubuk(sl, 6.6, 2.05, 6.1, 4.55, [x["ad"] for x in b], [("Öğrencilerin %'si", [x["oran"] for x in b])], eksen_max=100, sonek="%")
    if not a and not b:
        _metin(sl, SOL, 2.2, GEN, 1, "Bu kapsamda gösterilecek kadar toplu veri yok.", 16, GRI)
    _metin(sl, SOL, 6.7, GEN, 0.25, f"Toplu değerler; {d.v['esik']}'ten az öğrencinin seçtiği alan ve bölümler “Diğer” altında birleştirilir ya da gösterilmez.", 10, GRI)
    _not(sl, "Solda, öğrencilerin en uyumlu çıkan bölümlerinin hangi alanlarda toplandığını görüyorsunuz. Sağda ise öğrencilerin "
             "ilk 10 önerisinde en sık yer alan bölümler var. Bu liste bir yönlendirme değil, grubun ilgi haritasıdır. "
             "Her öğrencinin önerileri kendine özgüdür.")


def _hedef(d: _Deste):
    h = d.v.get("hedef") or {}
    sl = d.slayt("Hedef bölüm", "Hedef belirleyen öğrencilerimiz")
    d.stat(sl, SOL, 1.75, 3.7, _yuzde(h.get("oran")), "öğrencimiz bir hedef bölüm seçti", h=1.7)
    if h.get("ilk10_oran") is not None:
        d.stat(sl, SOL, 3.7, 3.7, _yuzde(h.get("ilk10_oran")), "hedef seçenlerin hedefi kendi ilk 10 önerisinde", h=1.7)
    al = h.get("alanlar") or []
    if not any(not x.get("diger") for x in al):
        al = []
    if al:
        d.halka(sl, 4.7, 1.65, 4.3, 4.6, [x["ad"] for x in al], [x["sayi"] for x in al], "Hedeflerin alanı (öğrenci)")
    x = 9.25 if al else 4.75
    _kart(sl, x, 1.75, W - SOL - x, 3.65, d.r.acik)
    _metin(sl, x + 0.3, 2.0, W - SOL - x - 0.6, 3.2, [
        [("Hedef ne demek?", {"kalin": True, "boyut": 16, "renk": d.r.koyu})],
        "Öğrenci, sonuçlarını inceledikten sonra kendine bir hedef bölüm seçebilir. Hedef, çalışma planını ve koçluk "
        "adımlarını yönlendirir.",
        "Hedef kesin değildir; öğrenci tanıdıkça değiştirebilir. Önerilerin dışında bir hedef seçmek de mümkündür ve "
        "birlikte konuşmak için iyi bir fırsattır.",
    ], 14, MUREKKEP, satir_arasi=10)
    _not(sl, "Hedef seçmek, öğrencinin motivasyonu için önemli bir adım. Çocuğunuzun hedefini sorun; neden seçtiğini dinleyin. "
             "Hedefin önerilerle örtüşmemesi bir sorun değildir, yalnızca birlikte konuşmak için iyi bir başlangıçtır.")


def _etkinlik(d: _Deste):
    e = d.v.get("etkinlik") or {}
    sl = d.slayt("Okulda neler yapıyoruz", "Koçluk ve İş Hayatı çalışmaları")
    kartlar = []
    if "kocluk" in e:
        k = e["kocluk"] or {}
        sayilar = []
        if k.get("adim") is not None:
            sayilar.append(f"{k['adim']} öğrenci kişisel gelişim adımı tamamladı")
        if k.get("gorev") is not None:
            sayilar.append(f"{k['gorev']} öğrenci haftalık görevlerini yaptı")
        if k.get("simulasyon") is not None:
            sayilar.append(f"{k['simulasyon']} meslek simülasyonu (“Bir günümü yaşa”) oynandı")
        kartlar.append(("Koçluk ve görevler", [
            "Kişisel yol haritası: güçlü yönlere ve hedefe göre küçük, uygulanabilir adımlar",
            "Haftalık görevler: bölüm keşfi, okuma, meslek araştırması",
            "Meslek simülasyonu: bir mesleğin gününü deneyimleyerek tanıma",
        ], sayilar))
    if "is_hayati" in e:
        kartlar.append(("İş Hayatı", [
            "Mezunların istihdam oranı ve iş bulma süresi (resmî kaynaklar)",
            "Mesleğe giden yol: eğitim, staj, sınav ve belge adımları",
            "Kazanç ve yaşam giderleri hakkında gerçekçi bilgi",
            "CV ve mülakat atölyeleri",
        ], []))
    n = len(kartlar)
    kw = (GEN - 0.3 * (n - 1)) / n
    kh = 4.85 if any(k[2] for k in kartlar) else 3.7
    for i, (baslik, maddeler, sayilar) in enumerate(kartlar):
        x = SOL + i * (kw + 0.3)
        _kart(sl, x, 1.75, kw, kh, d.r.acik)
        _daire(sl, x + 0.3, 2.0, 0.6, d.r.ana, str(i + 1), boyut=18)
        _metin(sl, x + 1.1, 2.05, kw - 1.4, 0.5, baslik, 20, MUREKKEP, kalin=True, dikey=MSO_ANCHOR.MIDDLE)
        pars = [[("•  ", {"renk": d.r.ana, "kalin": True}), (m, {})] for m in maddeler]
        _metin(sl, x + 0.3, 2.85, kw - 0.6, 2.4, pars, 15, MUREKKEP, satir_arasi=8)
        if sayilar:
            _metin(sl, x + 0.3, 5.25, kw - 0.6, 1.2, [[(s, {"kalin": True, "renk": d.r.koyu})] for s in sayilar], 14, satir_arasi=4)
    if kh < 4.85:
        _metin(sl, SOL, 5.75, GEN, 0.6, "Öğrencilerimiz bu çalışmalara Filizyol hesaplarından istedikleri zaman ulaşabilir. "
               f"Katılım sayıları {d.v['esik']}'ten az öğrenciye dayandığında paylaşılmaz.", 14, GRI)
    _not(sl, "Okulumuzda Filizyol yalnızca bir test değil; sonrasında koçluk adımları ve iş hayatı bilgileriyle devam ediyor. "
             "Çocuğunuza bu hafta hangi görevi yaptığını, hangi mesleği simülasyonda denediğini sorabilirsiniz. "
             "Sayılar yalnızca toplu olarak verilir; 5'ten az öğrenciye dayanan sayılar gösterilmez.")


DESTEK = [
    ("Dinleyin", "Ne istediğini, neden istediğini merakla sorun; hemen yargılamayın."),
    ("Baskı yapmayın", "Puan, sıralama ya da “şu meslek iyi para kazandırır” baskısı kararı zorlaştırır."),
    ("Gerçekçi bilgi verin", "Mesleklerin zorlu yanlarını da konuşun; iş hayatı bilgilerini birlikte inceleyin."),
    ("Meslek sahipleriyle tanıştırın", "Çevrenizdeki farklı meslek sahipleriyle kısa bir sohbet çok şey öğretir."),
    ("Keşfe alan açın", "Kulüp, gönüllülük, kampüs gezisi ve yaz okulları ilgileri somutlaştırır."),
    ("Rehberlik servisiyle birlikte", "Sonuçları çocuğunuzla birlikte okuyun; sorularınızı rehber öğretmenimize iletin."),
]


def _destek(d: _Deste):
    sl = d.slayt("Velilerimize", "Çocuğunuza nasıl destek olabilirsiniz")
    kw, kh = (GEN - 0.6) / 3, 2.3
    for i, (b, a) in enumerate(DESTEK):
        x = SOL + (i % 3) * (kw + 0.3)
        y = 1.75 + (i // 3) * (kh + 0.3)
        _kart(sl, x, y, kw, kh, d.r.acik)
        _daire(sl, x + 0.3, y + 0.3, 0.6, d.r.ana, str(i + 1), boyut=18)
        _metin(sl, x + 1.1, y + 0.3, kw - 1.35, 0.6, b, 18, MUREKKEP, kalin=True, dikey=MSO_ANCHOR.MIDDLE)
        _metin(sl, x + 0.3, y + 1.1, kw - 0.6, kh - 1.25, a, 15, MUREKKEP)
    _not(sl, "Bu slayt sunumun en önemli kısmı. Ergenlik döneminde meslek kararı çoğu zaman değişir; bu doğaldır. En büyük destek, "
             "baskı yerine merakla dinlemek ve çocuğa farklı meslekleri tanıma fırsatı sunmaktır. Bildiğiniz bir meslek sahibini "
             "okula konuşmacı olarak davet etmek isterseniz bizimle iletişime geçebilirsiniz.")


VARSAYILAN_ADIMLAR = [
    ("Değerlendirmeyi tamamlama", "Henüz bitirmeyen öğrenciler kendi hızında tamamlar."),
    ("Sonuçları birlikte okuma", "Öğrenci raporunu evde birlikte inceleyin."),
    ("Hedef bölüm seçimi", "Öğrenci önerilere bakarak bir hedef belirler; sonra değiştirebilir."),
    ("Rehberlik görüşmesi", "Gerekirse rehber öğretmenimizle bireysel görüşme planlanır."),
]


def _takvim(d: _Deste):
    t = d.v.get("takvim") or []
    sl = d.slayt("Takvim", "Yaklaşan etkinlikler ve sonraki adımlar" if t else "Sonraki adımlar")
    ogeler = [(_tarih(e["tarih"]) + (f" · {e['saat']}" if e.get("saat") else ""), e["baslik"]) for e in t[:6]]
    if len(ogeler) < 3:
        ogeler += VARSAYILAN_ADIMLAR[:4 - len(ogeler)]
    n = len(ogeler)
    # yatay zaman çizelgesi: daireler + altlarında metin
    y = 3.0
    cizgi = sl.shapes.add_connector(1, Inches(SOL + 0.35), Inches(y + 0.3), Inches(W - SOL - 0.35), Inches(y + 0.3))
    cizgi.line.color.rgb = _rgb(d.r.orta)
    cizgi.line.width = Pt(2)
    aralik = GEN / n
    for i, (ust, alt) in enumerate(ogeler):
        x = SOL + i * aralik
        _daire(sl, x + aralik / 2 - 0.3, y, 0.6, d.r.ana, str(i + 1), boyut=16)
        _metin(sl, x + 0.1, y + 0.85, aralik - 0.2, 0.7, ust, 15, d.r.koyu, kalin=True, hiza=PP_ALIGN.CENTER)
        _metin(sl, x + 0.1, y + 1.5, aralik - 0.2, 1.6, alt, 14, MUREKKEP, hiza=PP_ALIGN.CENTER)
    if t:
        _metin(sl, SOL, 1.8, GEN, 0.6, "Okulumuzun ve genel takvimdeki yaklaşan etkinlikler ile öğrencilerimizle atacağımız adımlar. "
               "Etkinlikler öğrencinin Filizyol takviminde de görünür.",
               15, GRI)
    else:
        _metin(sl, SOL, 1.8, GEN, 0.6, "Önümüzdeki haftalarda öğrencilerimizle birlikte atacağımız adımlar.", 15, GRI)
    _not(sl, ("Yaklaşan etkinlikleri hatırlatın; tarihler öğrencilerin Filizyol takviminde de yer alıyor. " if t else
              "Sonraki adımları kısaca anlatın. ")
         + "Velilerin katılabileceği etkinlikler (veli görüşme günleri, kariyer günleri) varsa özellikle vurgulayın.")


def _iletisim(d: _Deste):
    il = d.v.get("iletisim") or {}
    sl = d.slayt("İletişim", "Sorularınız için rehberlik servisimiz", koyu=True)
    reh = il.get("rehberler") or []
    yazi = BEYAZ
    soluk = _karistir(d.r.ana, "FFFFFF", 0.75)
    y = 1.9
    if reh:
        for k in reh:
            _metin(sl, SOL, y, 7.5, 1.3, [
                [(k["ad"], {"boyut": 24, "kalin": True})],
                [(k.get("gorev") or "", {"boyut": 15, "renk": soluk})],
                [(" · ".join(x for x in (k.get("eposta"), k.get("telefon")) if x), {"boyut": 15})],
            ], 15, yazi, satir_arasi=2)
            y += 1.45
    else:
        _metin(sl, SOL, y, 7.5, 1, "Rehberlik servisimiz okul saatlerinde görüşme için hazırdır.", 20, yazi)
        y += 1.2
    genel = [x for x in (il.get("telefon"), il.get("eposta"), il.get("web")) if x]
    if genel:
        _metin(sl, SOL, max(y, 4.6), 7.5, 0.9, [[("Okul", {"kalin": True, "renk": soluk})], [(" · ".join(genel), {})]], 15, yazi, satir_arasi=2)
    _metin(sl, 8.6, 2.2, 4.1, 2.5, [[("Teşekkür ederiz", {"boyut": 34, "kalin": True})],
                                    [("Çocuğunuzun raporu, öğrenci hesabındaki Raporlarım bölümünden ya da rehberlik servisinden "
                                      "alınabilir.", {"boyut": 15, "renk": soluk})]], 15, yazi, satir_arasi=10)
    _not(sl, "Rehber öğretmenimizin iletişim bilgilerini paylaşın ve soruları alın. Bireysel sonuçlarla ilgili soruları toplantıda "
             "değil, birebir görüşmede konuşmayı önerin; böylece her öğrencinin gizliliği korunur.")


def veli_sunumu_pptx(v: dict) -> bytes:
    from app.core.katman_adlari import KATMAN_ADI
    d = _Deste(v)
    _kapak(d)
    _nedir(d, [(k, KATMAN_ADI[k]) for k in ("K1", "K2", "K3", "K4")])
    if v.get("toplu"):
        _katilim(d)
        if v.get("profil_yeterli"):
            if v.get("guclu"):
                _guclu(d)
            if any(not x.get("diger") and x.get("sayi") for x in v.get("alanlar") or []) or v.get("bolumler"):
                _alanlar(d)
        if (v.get("hedef") or {}).get("oran") is not None:
            _hedef(d)
        if v.get("etkinlik"):
            _etkinlik(d)
    _destek(d)
    _takvim(d)
    _iletisim(d)
    d.p.core_properties.title = f"Veli toplantısı — {v['okul']['ad']} ({v['kapsam']})"
    d.p.core_properties.author = "Filizyol"
    out = io.BytesIO()
    d.p.save(out)
    return out.getvalue()
