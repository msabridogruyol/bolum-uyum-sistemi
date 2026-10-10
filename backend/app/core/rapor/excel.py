# -*- coding: utf-8 -*-
"""[2026-10-10] Excel raporları: öğrenci (tüm puanlar, öneriler, SWOT) ve okul (özet, sınıflar, öğrenciler)."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASLIK = Font(bold=True, color="FFFFFF")
DOLGU = PatternFill("solid", fgColor="E8804A")


def _sayfa(wb, ad, basliklar, satirlar, genislik=None, renk="E8804A"):
    ws = wb.create_sheet(ad[:31])
    ws.append(basliklar)
    for c in ws[1]:
        c.font = BASLIK
        c.fill = PatternFill("solid", fgColor=renk.lstrip("#"))
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for s in satirlar:
        ws.append([(x.replace(tzinfo=None) if isinstance(x, datetime) and x.tzinfo else x) for x in s])
    for i, g in enumerate(genislik or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = g
    ws.freeze_panes = "A2"
    return ws


def ogrenci_xlsx(v: dict, netler: bool = True) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    renk = v["okul"].get("renk") or "#E8804A"
    k, tur = v["kisi"], v.get("tur") or {}
    _sayfa(wb, "Özet", ["Alan", "Değer"], [
        ["Öğrenci", k["ad_soyad"]], ["E-posta", k["email"]], ["Okul", v["okul"]["ad"]],
        ["Sınıf / şube", " / ".join(filter(None, [k.get("sinif"), k.get("sube")]))],
        ["Değerlendirme turu", tur.get("no")], ["Tamamlanma", tur.get("tamamlanma")],
        ["Güven puanı", tur.get("guven")], ["Geçerli", "Evet" if tur.get("gecerli", True) else "Hayır"],
        ["Hedef bölüm", (v.get("hedef") or {}).get("ad")], ["Listem", ", ".join(v["listem"])],
        ["Rapor tarihi", v["tarih"]],
    ], [24, 60], renk)
    _sayfa(wb, "Özellik puanları", ["Katman", "Özellik", "Puan", "Düzey", "Durum tespiti", "Öneri"],
           [[kt["kod"], x["ad"], x["puan"], x["seviye"], x.get("yorum") or "", x.get("oneri") or ""]
            for kt in v["katmanlar"] for x in kt["ozellikler"]] + [["K5", x["ad"], x["puan"], "", "", ""] for x in v["k5"]],
           [9, 34, 8, 14, 60, 60], renk)
    _sayfa(wb, "Bölüm önerileri", ["Sıra", "Bölüm", "Alan", "Uyum %", "Örtüşen özellikler", "Dikkat"],
           [[b["sira"], b["ad"], b.get("alan"), b["uyum"], ", ".join(b.get("ortusen") or []), b.get("dikkat") or ""] for b in v["bolumler"]],
           [6, 40, 26, 9, 50, 60], renk)
    s = v["swot"]
    n = max(len(s["S"]), len(s["W"]), len(s["O"]), len(s["T"]), 1)
    _sayfa(wb, "SWOT", ["Güçlü yönler (S)", "Gelişim alanları (W)", "Fırsatlar (O)", "Dikkat edilecekler (T)"],
           [[(s[x][i] if i < len(s[x]) else "") for x in "SWOT"] for i in range(n)], [45, 45, 45, 45], renk)
    nt = v.get("netler") if netler else None
    if nt and nt.get("denemeler"):   # [2026-10-10] deneme / net takibi
        from app.core.sinav_yapisi import TESTLER
        kodlar = [k for k in TESTLER if any(k in d["dersler"] for d in nt["denemeler"])]
        _sayfa(wb, "Denemeler", ["Tarih", "Oturum", "Deneme", *[f"{TESTLER[k][1]} net" for k in kodlar], "Toplam net"],
               [[d["tarih"], d["oturum"], d.get("ad") or "", *[(d["dersler"].get(k) or {}).get("net") for k in kodlar], d["toplam"]]
                for d in nt["denemeler"]], [12, 8, 28, *[14] * len(kodlar), 12], renk)
        k = nt.get("kiyas")
        if k:
            _sayfa(wb, "Hedef kıyas", ["Ders", "Öğrenci (son 3 ort.)", f"{k.get('yil')} son yerleşen", "Önceki yıl", "Yıllar ort.", "Fark"],
                   [[x["ad"], x["ben"], x["hedef"], x["hedef_onceki_yil"], x["hedef_ortalama"], x["fark"]] for x in k["satirlar"]]
                   + [["Hedef program", f"{k['hedef']['universite']} · {k['hedef']['program']}", "", "", "", ""]],
                   [26, 18, 18, 14, 14, 10], renk)
    t = io.BytesIO()
    wb.save(t)
    return t.getvalue()


def okul_xlsx(v: dict, netler: bool = True) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    renk = v["okul"].get("renk") or "#E8804A"
    o = v["ozet"]
    ks = v.get("kapsam") or {}
    kapsam_satir = [["Kapsam", ks.get("etiket")]] if ks.get("sinif") else []
    if (ks.get("ogretmen") or {}).get("ad"):
        kapsam_satir.append(["Sınıf öğretmeni", ks["ogretmen"]["ad"]])
    okul_ort = {k["kod"]: k["ortalama"] for k in v.get("okul_katman_ort") or []}
    _sayfa(wb, "Özet", ["Gösterge", "Değer"], [
        ["Okul", v["okul"]["ad"]], *kapsam_satir, ["Öğrenci", o.get("toplam")], ["Giriş yapan", o.get("giris_yapan")],
        ["Teste başlayan", o.get("teste_baslayan")], ["Tamamlayan", o.get("tamamlayan")], ["Hedef seçen", o.get("hedef_secen")],
        ["Güven eşiği altında", v["gecersiz"]], ["Rapor tarihi", v["tarih"]],
    ] + [[f"Katman ortalaması · {k['kod']} {k['ad']}", k["ortalama"]] for k in v["katman_ort"]]
      + [[f"Okul ortalaması · {k} ", x] for k, x in okul_ort.items()], [40, 30], renk)
    if o.get("subeler"):
        _sayfa(wb, "Şubeler", ["Şube", "Sınıf öğretmeni", "Öğrenci", "Giriş yapan", "Devam eden", "Tamamlayan", "Hedef seçen"],
               [[x["etiket"], (x.get("ogretmen") or {}).get("ad") or "", x["ogrenci"], x["giris_yapan"], x["devam"], x["tamamlayan"],
                 x.get("hedef_secen", 0)] for x in o["subeler"]], [12, 28, 10, 12, 12, 12, 12], renk)
    _sayfa(wb, "Sınıflar", ["Sınıf", "Öğrenci", "Giriş yapan", "Devam eden", "Tamamlayan"],
           [[s["sinif"], s["ogrenci"], s["giris_yapan"], s["devam"], s["tamamlayan"]] for s in o.get("siniflar", [])], [18, 10, 12, 12, 12], renk)
    net_bas = ["Deneme", "Son TYT", "Son AYT"] if netler else []
    _sayfa(wb, "Öğrenciler", ["No", "Ad soyad", "Sınıf", "Durum", "1. öneri", "Hedef", "Güven", "Son giriş", "Test hesabı", *net_bas],
           [[x.get("no") or "", x["ad_soyad"], x["sinif"], x["durum"], x["ilk_bolum"], x["hedef"], x["guven"], x["son_giris"],
             "Evet" if x["test"] else "", *([x.get("deneme") or 0, x.get("son_tyt"), x.get("son_ayt")] if netler else [])]
            for x in v["ogrenciler"]], [8, 28, 10, 22, 34, 34, 8, 18, 10, *([9, 10, 10] if netler else [])], renk)
    nt = v.get("net") or {}
    if netler and nt.get("dersler"):
        _sayfa(wb, "Net özeti", ["Ders", "Soru", "Son deneme ortalaması", "Öğrenci"],
               [[d["ad"], d["soru"], d["ort"], d["n"]] for d in nt["dersler"]], [34, 8, 22, 10], renk)
    _sayfa(wb, "Bölüm ve alan", ["Alan (1. öneri)", "Öğrenci", "", "En çok önerilen", "Sayı", "En çok hedeflenen", "Sayı"],
           [[(v["alanlar"][i]["alan"] if i < len(v["alanlar"]) else ""), (v["alanlar"][i]["sayi"] if i < len(v["alanlar"]) else ""), "",
             *((o["en_cok_onerilen"][i]["bolum"], o["en_cok_onerilen"][i]["sayi"]) if i < len(o.get("en_cok_onerilen", [])) else ("", "")),
             *((o["en_cok_hedeflenen"][i]["bolum"], o["en_cok_hedeflenen"][i]["sayi"]) if i < len(o.get("en_cok_hedeflenen", [])) else ("", ""))]
            for i in range(max(len(v["alanlar"]), len(o.get("en_cok_onerilen", [])), len(o.get("en_cok_hedeflenen", [])), 1))],
           [30, 9, 3, 36, 7, 36, 7], renk)
    t = io.BytesIO()
    wb.save(t)
    return t.getvalue()
