# -*- coding: utf-8 -*-
"""
[2026-10-17] Cevap ve uyum matrisleri — Excel (okul yetkilisi + süper admin; öğrenciye verilmez).

Sayfalar:
  Okuma notu        — puanların nasıl okunacağı
  Cevaplar          — soru × özellik: her cevabın her özelliğe katkısı (cevap matrisi)
  Özellikler        — özellik puanı ve hangi cevaplardan oluştuğu
  Uyum (fark)       — bölüm × özellik: öğrenci − bölüm beklentisi; yeşil uyumlu, turuncu fazla/az, kırmızı belirgin
  Katkılar          — bölüm × özellik: "bu özellik olmasaydı" puan değişimi (artı yükseltiyor, eksi düşürüyor)
  Bölüm beklentisi  — bölüm × özellik: bölümün beklediği düzey (katman içi ölçek)
  Yöntemler         — yalnızca süper admin: 10 ÇKKV yönteminin puanı + Kendall W
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.koclugu_servisi import fark_sinifi

BASLIK = Font(bold=True, color="FFFFFF")
UST = PatternFill("solid", fgColor="E8804A")
RENK = {"yesil": PatternFill("solid", fgColor="DCEFD6"), "turuncu": PatternFill("solid", fgColor="FBE3C2"),
        "kirmizi": PatternFill("solid", fgColor="F6CFCB")}
KATKI_ARTI = PatternFill("solid", fgColor="DCEFD6")
KATKI_EKSI = PatternFill("solid", fgColor="F6CFCB")


def _sayfa(wb, ad, basliklar, genislik=None):
    ws = wb.create_sheet(ad[:31])
    ws.append(basliklar)
    for c in ws[1]:
        c.font = BASLIK
        c.fill = UST
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, g in enumerate(genislik or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = g
    ws.freeze_panes = "C2"
    return ws


def _renk(sinif: str):
    return RENK["yesil" if sinif == "uyumlu" else "kirmizi" if sinif.startswith("belirgin") else "turuncu"]


def cevap_analizi_xlsx(db, o, tur, satirlar: list, ozellikler: list, tam: dict | None, yontem: bool, not_metni: str) -> bytes:
    from app.models import OgrenciBolumUyumSkoru
    wb = Workbook()
    wb.remove(wb.active)

    ws = _sayfa(wb, "Okuma notu", ["Bilgi", "Değer"], [26, 120])
    for a, b in (("Öğrenci", o.ad_soyad), ("Tur", tur.tur_no), ("Durum", tur.durum), ("Nasıl okunur", not_metni),
                 ("Gizlilik", "Bu dosya öğrencinin değerlendirme cevaplarını içerir (kişisel veri). Yalnızca rehberlik amacıyla "
                              "kullanın, paylaşmayın.")):
        ws.append([a, b])
    for r in ws.iter_rows(min_row=2):
        r[1].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    # Cevap matrisi: satır = soru, sütun = özellik
    kodlar = [x["kod"] for x in ozellikler]
    ws = _sayfa(wb, "Cevaplar", ["#", "Katman", "Tür", "Soru", "Seçilen cevap", "En az uyan"] + kodlar,
                [5, 8, 12, 60, 40, 30] + [7] * len(kodlar))
    for c in satirlar:
        secilen = next((x["metin"] for x in c["secenekler"] if x["secildi"]), "")
        en_az = next((x["metin"] for x in c["secenekler"] if x["en_az"]), "")
        katki = {k["kod"]: k["puan"] for k in c["katkilar"]}
        ws.append([c["sira"], c["katman"], c["tip_ad"] + (" (ters)" if c["ters"] else ""), c["soru"], secilen, en_az]
                  + [katki.get(k) for k in kodlar])

    ws = _sayfa(wb, "Özellikler", ["Kod", "Özellik", "Katman", "Puan", "Cevap sayısı", "Katkı veren sorular (#)"], [7, 34, 8, 8, 12, 50])
    for x in ozellikler:
        ws.append([x["kod"], x["ad"], x.get("katman"), x.get("puan"), x["cevap_sayisi"], ", ".join(str(i) for i in x["cevaplar"])])

    if tam is not None:
        g, K = tam["girdi"], tam["katki"]
        kod = g.degisken_kodlari
        bas = ["Sıra", "Bölüm", "Uyum"]
        gen = [6, 44, 8] + [7] * len(kod)
        ws_f = _sayfa(wb, "Uyum (fark)", bas + kod, gen)
        ws_k = _sayfa(wb, "Katkılar", bas + kod, gen)
        ws_b = _sayfa(wb, "Bölüm beklentisi", bas + kod, gen)
        ws_f.append(["", "Öğrenci (katman içi ölçek)", ""] + [round(float(v), 1) for v in g.ogrenci_olcekli])
        for n, s in enumerate(tam["sira"], 1):
            i = tam["satir_of"][s.bolum_id]
            ad = g.bolum_adlari.get(s.bolum_id)
            farklar = [round(float(g.ogrenci_olcekli[j] - g.bolum_olcekli[i, j]), 1) for j in range(len(kod))]
            ws_f.append([n, ad, round(float(s.toplam_uyum), 1)] + farklar)
            for j, f in enumerate(farklar):
                ws_f.cell(row=ws_f.max_row, column=4 + j).fill = _renk(fark_sinifi(f))
            katkilar = [round(float(K[i, j]), 2) for j in range(len(kod))]
            ws_k.append([n, ad, round(float(s.toplam_uyum), 1)] + katkilar)
            for j, v in enumerate(katkilar):
                if abs(v) >= 0.5:
                    ws_k.cell(row=ws_k.max_row, column=4 + j).fill = KATKI_ARTI if v > 0 else KATKI_EKSI
            ws_b.append([n, ad, round(float(s.toplam_uyum), 1)] + [round(float(g.bolum_olcekli[i, j]), 1) for j in range(len(kod))])

        if yontem:
            from app.core.skor_motoru import YONTEMLER
            adlar = list(YONTEMLER)
            ws = _sayfa(wb, "Yöntemler", ["Sıra", "Bölüm", "Uyum (ortalama)"] + adlar + ["Kendall W"], [6, 44, 10] + [9] * (len(adlar) + 1))
            kayit = {r.bolum_id: r for r in db.query(OgrenciBolumUyumSkoru).filter(
                OgrenciBolumUyumSkoru.ogrenci_id == o.id, OgrenciBolumUyumSkoru.tur_id == tur.id)}
            for n, s in enumerate(tam["sira"], 1):
                r = kayit.get(s.bolum_id)
                ys = (r.yontem_skorlari or {}) if r else {}
                ws.append([n, g.bolum_adlari.get(s.bolum_id), round(float(s.toplam_uyum), 1)]
                          + [ys.get(a) for a in adlar] + [float(r.kendall_w) if r and r.kendall_w is not None else None])

    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()
