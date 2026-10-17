"""
[2026-10-17] Cevap ve uyum matrisleri:
  * erişim: süper admin + okul yetkilisi (yalnızca kendi okulu); öğrenci erişemez;
  * yöntem matrisi yalnızca süper admine döner;
  * katkı = "bu özellik olmasaydı" (gösterilen puanla aynı hesap);
  * Excel: beklenen sayfalar, yöntem sayfası yalnızca süper admin.
"""
import io
import uuid

import numpy as np
from sqlalchemy import text

from app.core import uyum_aciklama as ua
from tests.conftest import db_gerekli


def test_katki_birakmali_ve_tutarli():
    rng = np.random.default_rng(7)
    X = rng.uniform(20, 100, size=(40, 6))
    w = rng.uniform(0.5, 2, size=6)
    K = ua.ozellik_katkilari(X, w)
    assert K.shape == X.shape
    tam = ua._ortalama(X, w)
    # j. sütun = tam puan − j. özellik çıkarılınca hesaplanan puan
    j = 3
    maske = np.arange(6) != j
    assert np.allclose(K[:, j], tam - ua._ortalama(X[:, maske], w[maske]))
    # bir bölümde bir özellik çok iyiyse (diğer bölümlere göre) o özelliğin katkısı o bölümde pozitif olmalı
    X2 = np.full((10, 4), 60.0) + rng.uniform(-3, 3, size=(10, 4))
    X2[0, 2] = 100.0
    X2[1:, 2] = 30.0
    K2 = ua.ozellik_katkilari(X2, np.ones(4))
    assert K2[0, 2] > 0
    assert K2[0, 2] == K2[0].max()


def _ogrenci(db, okul_id):
    oid = uuid.uuid4()
    db.execute(text("""INSERT INTO ogrenciler (id, ad_soyad, email, sifre_hash, okul_id, olusturulma_zamani, guncelleme_zamani)
                       VALUES (:i, 'Başka Okul Öğrencisi', :e, 'x', :o, now(), now())"""),
               {"i": oid, "e": f"bo.{oid.hex[:8]}@test.invalid", "o": okul_id})
    db.commit()
    return oid


def _kendi_ogrencisi(db, okul_id):
    return db.execute(text("SELECT id FROM ogrenciler WHERE okul_id=:o LIMIT 1"), {"o": okul_id}).scalar()


@db_gerekli
def test_yetkili_kendi_okulu_gorur_baska_okulu_goremez(istemci, db, kisiler):
    kendi = _kendi_ogrencisi(db, kisiler["okul_a"])
    r = istemci.get(f"/yonetim/ogrenci/{kendi}/cevap-analizi", headers=kisiler["yetkili"])
    assert r.status_code == 200, r.text
    assert r.json()["tur"] is None                                   # teste başlamamış öğrenci
    baska = _ogrenci(db, kisiler["okul_b"])
    for yol in ("cevap-analizi", "cevap-analizi.xlsx"):
        r = istemci.get(f"/yonetim/ogrenci/{baska}/{yol}", headers=kisiler["yetkili"])
        assert r.status_code == 403, (yol, r.text)
    assert istemci.get(f"/yonetim/ogrenci/{baska}/cevap-analizi", headers=kisiler["super"]).status_code == 200


@db_gerekli
def test_ogrenci_erisemez(istemci, db, kisiler):
    kendi = _kendi_ogrencisi(db, kisiler["okul_a"])
    for yol in ("cevap-analizi", "cevap-analizi.xlsx"):
        r = istemci.get(f"/yonetim/ogrenci/{kendi}/{yol}", headers=kisiler["ogrenci"])
        assert r.status_code in (401, 403), (yol, r.status_code)


def _tamamlanmis(db):
    """Son turu tamamlanmış ve uyum skoru kayıtlı, okula bağlı bir öğrenci."""
    return db.execute(text("""SELECT o.id, o.okul_id FROM ogrenciler o
                              JOIN LATERAL (SELECT id, durum FROM ogrenci_degerlendirme_turu t WHERE t.ogrenci_id = o.id
                                            ORDER BY t.tur_no DESC LIMIT 1) son ON TRUE
                              WHERE o.okul_id IS NOT NULL AND son.durum = 'tamamlandi'
                                AND EXISTS (SELECT 1 FROM ogrenci_bolum_uyum_skorlari s WHERE s.tur_id = son.id)
                              LIMIT 1""")).first()


def _yetkili_tokeni(db, okul_id):
    from app.core.security import erisim_tokeni_uret
    kid = uuid.uuid4()
    db.execute(text("""INSERT INTO admin_kullanicilar (id, ad_soyad, email, sifre_hash, rol, okul_id, aktif_mi, olusturulma_zamani)
                       VALUES (:i, 'Test yetkili', :e, 'x', 'okul_yetkilisi', :o, TRUE, now())"""),
               {"i": kid, "e": f"yet.{kid.hex[:8]}@test.invalid", "o": okul_id})
    db.commit()
    return {"Authorization": f"Bearer {erisim_tokeni_uret(kid, 'okul_yetkilisi')}"}


@db_gerekli
def test_matris_ve_yontem_kapsami(istemci, db, kisiler):
    satir = _tamamlanmis(db)
    if satir is None:
        return
    oid, okul = satir
    sa = istemci.get(f"/yonetim/ogrenci/{oid}/cevap-analizi", headers=kisiler["super"]).json()
    m = sa["matris"]
    assert m and m["yontem_adlari"] and len(m["yontem_adlari"]) == 10
    assert all("yontemler" in b for b in m["bolumler"])
    sira = [b["sira"] for b in m["bolumler"]]
    assert sira[:10] == list(range(1, 11))                          # varsayılan: ilk 10 (+ hedef) + son 3
    b0 = m["bolumler"][0]
    assert len(b0["hucreler"]) == len(m["ozellikler"])
    assert all(x["katki"] > 0 for x in b0["yukselten"]) and all(x["katki"] < 0 for x in b0["dusuren"])
    assert all(h["sinif"] in {"uyumlu", "fazla", "eksik", "belirgin_fazla", "belirgin_eksik"} for h in b0["hucreler"])

    yt = istemci.get(f"/yonetim/ogrenci/{oid}/cevap-analizi", headers=_yetkili_tokeni(db, okul)).json()
    assert yt["matris"] and yt["matris"]["yontem_adlari"] is None
    assert all("yontemler" not in b and "kendall_w" not in b for b in yt["matris"]["bolumler"])
    assert yt["cevaplar"] and yt["okuma_notu"]

    # bölüm seçimi
    secim = [m["bolumler"][-1]["bolum_id"], m["bolumler"][0]["bolum_id"]]
    r = istemci.get(f"/yonetim/ogrenci/{oid}/cevap-analizi?bolumler={secim[0]},{secim[1]}", headers=kisiler["super"]).json()
    assert [b["bolum_id"] for b in r["matris"]["bolumler"]] == secim


@db_gerekli
def test_excel_sayfalari(istemci, db, kisiler):
    from openpyxl import load_workbook
    satir = _tamamlanmis(db)
    if satir is None:
        return
    oid, okul = satir
    temel = ["Okuma notu", "Cevaplar", "Özellikler", "Uyum (fark)", "Katkılar", "Bölüm beklentisi"]
    r = istemci.get(f"/yonetim/ogrenci/{oid}/cevap-analizi.xlsx", headers=kisiler["super"])
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    assert load_workbook(io.BytesIO(r.content)).sheetnames == temel + ["Yöntemler"]
    r = istemci.get(f"/yonetim/ogrenci/{oid}/cevap-analizi.xlsx", headers=_yetkili_tokeni(db, okul))
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == temel
    assert wb["Uyum (fark)"].max_row > 100                          # öğrenci satırı + tüm bölümler
