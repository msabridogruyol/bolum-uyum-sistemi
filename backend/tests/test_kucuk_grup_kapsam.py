"""
[2026-10-14] Küçük grup gizleme kapsamı: yalnızca verinin yetkisiz kitleye gittiği yerlerde.
  * veli toplantısı sunumu — veli sürümü gizler, öğretmen sürümü gizlemez;
  * okul yetkilisi raporu (okul / şube PDF-Excel verisi) gizlemez;
  * kucuk_grup yardımcıları eşik parametresiyle çalışır (varsayılan 5).
"""
import io

from sqlalchemy import text

from app.core.kucuk_grup import gizle, grup_birlestir, kirilim_gizle, yeterli
from tests.conftest import db_gerekli


def test_esik_parametresi():
    assert gizle(3, 3) is None and gizle(3, 3, esik=0) == 3
    assert not yeterli(4) and yeterli(4, 0)
    assert any(x.get("diger") for x in grup_birlestir({"A": 10, "B": 2}))
    assert not any(x.get("diger") for x in grup_birlestir({"A": 10, "B": 2}, esik=0))
    s = [{"ad": "x", "ogrenci": 2, "oran": 50.0}]
    assert kirilim_gizle([dict(s[0])], ["oran"])[0]["oran"] is None
    assert kirilim_gizle([dict(s[0])], ["oran"], esik=0)[0]["oran"] == 50.0


def _kucuk_sube(db, okul_id):
    """3 öğrencili şube: 12. Sınıf / Z."""
    db.execute(text("UPDATE ogrenciler SET sinif = '12. Sınıf', sube = 'Z' WHERE okul_id = :o"), {"o": okul_id})
    for i in range(2):
        db.execute(text("""INSERT INTO ogrenciler (id, ad_soyad, email, sifre_hash, okul_id, sinif, sube, olusturulma_zamani, guncelleme_zamani)
                           VALUES (gen_random_uuid(), :a, :e, 'x', :o, '12. Sınıf', 'Z', now(), now())"""),
                   {"a": f"Kucuk {i}", "e": f"kucuk{i}.{okul_id}@test.invalid", "o": okul_id})
    db.commit()


@db_gerekli
def test_veli_sunumu_iki_surum(istemci, kisiler, db):
    from pptx import Presentation
    _kucuk_sube(db, kisiler["okul_a"])
    yol = f"/yonetim/okul/{kisiler['okul_a']}/veli-sunumu"
    p = {"sinif": "12. Sınıf", "sube": "Z"}
    veli = istemci.get(yol, params={**p, "kitle": "veli"}, headers=kisiler["super"])
    ogr = istemci.get(yol, params={**p, "kitle": "ogretmen"}, headers=kisiler["super"])
    assert veli.status_code == 200 and ogr.status_code == 200
    assert istemci.get(yol, params={"kitle": "x"}, headers=kisiler["super"]).status_code == 400
    metin = lambda r: " ".join(sh.text_frame.text for s in Presentation(io.BytesIO(r.content)).slides
                               for sh in s.shapes if sh.has_text_frame)
    mv, mo = metin(veli), metin(ogr)
    assert "KATILIM" not in mv and "5'ten az öğrenci olduğu için toplu sonuçlar paylaşılmamaktadır" in mv   # 3 kişi: gizli
    assert "KATILIM" in mo and "ÖĞRETMEN SÜRÜMÜ" in mo and "velilerle paylaşmayın" in mo                 # öğretmen: tüm sayılar
    assert "ÖĞRETMEN SÜRÜMÜ" not in mv


@db_gerekli
def test_yetkili_raporunda_gizleme_yok(kisiler, db):
    from app.core.rapor.veri import okul_raporu_verisi
    from app.models import AdminKullanici
    _kucuk_sube(db, kisiler["okul_a"])
    yon = db.query(AdminKullanici).filter(AdminKullanici.rol == "super_admin").first()
    v = okul_raporu_verisi(db, kisiler["okul_a"], yon, sinif="12. Sınıf", sube="Z")
    assert v["ozet"]["kapsam_gizli"] is False and v["kucuk_grup"]["uygulandi"] is False
    assert v["ozet"]["toplam"] == 3 and v["ozet"]["tamamlayan"] == 0     # gerçek sayı, None değil


@db_gerekli
def test_veli_sunumu_pdf_bicimleri(istemci, kisiler, db):
    _kucuk_sube(db, kisiler["okul_a"])
    yol = f"/yonetim/okul/{kisiler['okul_a']}/veli-sunumu"
    for bicim, uzanti in (("pdf", "pdf"), ("pdf_notlu", "pdf"), ("pptx", "pptx")):
        r = istemci.get(yol, params={"kitle": "ogretmen", "bicim": bicim}, headers=kisiler["super"])
        assert r.status_code == 200 and f'.{uzanti}"' in r.headers["content-disposition"]
        if uzanti == "pdf":
            assert r.content[:5] == b"%PDF-" and b"LiberationSans" in r.content and b"/Helvetica" not in r.content
    assert istemci.get(yol, params={"bicim": "docx"}, headers=kisiler["super"]).status_code == 400
