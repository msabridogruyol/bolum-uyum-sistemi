# -*- coding: utf-8 -*-
r"""
B5 — Veriye dayalı K5 iş türü bağları (2026-10-18)

Neden: bolum_k5_baglari (bölüm → iş türü, 0-1) elle hazırlanmıştı; bazı bölümlerde ana iş türü yanlıştı (ör. Bilgisayar
Mühendisliği → donanım). Kural: elle düzeltme yok; bağlar veriden üretilir.

YÖNTEM
  1) İş türü metni: her iş türünün (129, veri/k5_eksenleri.csv) K5 soru bankasındaki 3 etkinlik cümlesi
     (veri/k5_cumleleri.csv; app/data/k5_soru_bankasi.json'dan aktarılır). Her cümle 3 modelle "passage" olarak
     gömülür (Model C için cümlenin İngilizcesi; yoksa Model C atlanır); iş türü vektörü = 3 cümlenin ortalaması.
  2) Meslek ↔ iş türü: mesleklerin önbellekteki vektörleri (m4 adımı, "query") ile kosinüs benzerliği. Her model için
     önce sütun (iş türü) bazında z-puanı (bazı iş türü cümleleri her meslekle biraz daha benzer çıkar), sonra satır
     (meslek) bazında merkezleme; modellerin ortalaması alınır.
  3) Bölüm → meslek: b4 ile aynı kural [2026-10-21]: bölüm metni (resmî ders planındaki ders adları; yoksa kısa açıklama)
     ↔ ESCO meslek metni, en benzer 40 meslek; öğretmenlik dışı bölümde öğretmen meslekleri (ISCO 23) sayılmaz.
  4) Bölüm skoru = %70 meslek kanıtı (bölümün meslek ağırlıklarıyla ortalama) + %30 bölüm açıklaması kanıtı
     (açıklama vektörü ↔ iş türü). Her kanıt, iş türü bazında bölümler arasında z-puanına çevrilir (göreli: "bu iş
     bu bölümde diğer bölümlere göre ne kadar öne çıkıyor") ve bölüm içinde z-puanıyla birleştirilir (yarı yarıya).
  5) Seçim: en yüksek iş türü her zaman (bağ 1,0); diğerleri skoru en az ESIK_Z ve en yüksekten en çok FARK_Z
     gerideyse, en fazla EN_FAZLA_BAG bağ. Bağ = 0,25 + 0,75 × yakınlık (0,05 adımlı).
ÇIKTI
  veri/bolum_k5_bag_veri.csv              — üretilen bağlar (bolum;alan;eksen;bag)
  cikti/k5_baglari_veri.sql               — Supabase SQL'i (bolum_k5_baglari'nı bu bağlarla değiştirir)
  cikti/rapor/b5_k5_baglari_raporu.csv    — bölüm başına ilk 5 iş türü + skor + eski elle bağlar + uyum
  cikti/rapor/b5_ozet.txt                 — eski elle bağlarla uyum özeti (ana iş türü aynı mı, kesişim)
Kullanım:  C:\pipeline\venv\Scripts\python.exe C:\pipeline\src\b5_k5_baglari.py
           (A8_TFIDF=1 ortam değişkeniyle model yerine karakter n-gram TF-IDF kullanılır — yalnızca deneme için)
"""
import collections
import csv
import os
import sys
from pathlib import Path

import numpy as np

KOK = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1]
ONB, VERI, ONET, CIKTI = KOK / "onbellek", KOK / "veri", KOK / "onet", KOK / "cikti"
RAPOR = CIKTI / "rapor"
MODELLER = ("model_a", "model_b", "model_c")
MESLEK_PAYI = 0.70
ESIK_Z = 1.0          # ikincil bağ için en düşük birleşik skor (z)
FARK_Z = 0.8          # ikincil bağ en yüksek skordan en çok bu kadar geride olabilir
EN_FAZLA_BAG = 5     # [2026-10-21] 3 → 5: bölüm düzeyinde daha ayrıntılı bağ (ör. Müzik Öğretmenliği: sanat öğretmek + çalgı + şarkı)


def oku_csv(yol, ayirici=","):
    with open(yol, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=ayirici))


def zs(x, eksen):
    return (x - x.mean(eksen, keepdims=True)) / (x.std(eksen, keepdims=True) + 1e-9)


def normalize(x):
    return x / (np.linalg.norm(x, axis=1, keepdims=True) + 1e-12)


def xlsx_satirlar(yol, atla):
    from openpyxl import load_workbook
    ws = load_workbook(yol, read_only=True).worksheets[0]
    return list(ws.iter_rows(values_only=True))[atla:]


def is_turu_vektorleri(kodlar, cumle, model):
    """kod -> 3 cümlenin ortalama vektörü (normalize). Model C İngilizce cümle ister."""
    from embedding_saglayici import embed_et
    metinler, sahip = [], []
    for k in kodlar:
        for tr, en in cumle[k]:
            m = en if model == "model_c" else tr
            if not m:
                return None
            metinler.append(m); sahip.append(k)
    V = embed_et(metinler, model, onek="passage")
    out = np.zeros((len(kodlar), V.shape[1]), np.float32)
    idx = {k: i for i, k in enumerate(kodlar)}
    for v, k in zip(V, sahip):
        out[idx[k]] += v
    return normalize(out)


def tfidf_benzerlik(meslek_metin, bolum_metin, is_metin):
    from sklearn.feature_extraction.text import TfidfVectorizer
    vek = TfidfVectorizer(analyzer="char_wb", ngram_range=(4, 6), min_df=2, sublinear_tf=True)
    vek.fit(meslek_metin + bolum_metin + is_metin)
    M, B, I = (vek.transform(x) for x in (meslek_metin, bolum_metin, is_metin))
    return (M @ I.T).toarray(), (B @ I.T).toarray()


def main():
    RAPOR.mkdir(parents=True, exist_ok=True)
    print("[1/5] Girdiler")
    eksen = oku_csv(VERI / "k5_eksenleri.csv", ";")
    kodlar = [r["kod"].strip() for r in eksen]
    eksen_ad = {r["kod"].strip(): r["ad"].strip() for r in eksen}
    eksen_alan = {r["kod"].strip(): r["alan"].strip() for r in eksen}
    cumle = collections.defaultdict(list)
    for r in oku_csv(VERI / "k5_cumleleri.csv", ";"):
        cumle[r["kod"].strip()].append((r["metin"].strip(), (r.get("metin_en") or "").strip()))
    eksik = [k for k in kodlar if len(cumle.get(k, [])) == 0]
    if eksik:
        raise SystemExit(f"veri/k5_cumleleri.csv'de cümlesi olmayan iş türü var: {eksik[:10]}")
    meslek = oku_csv(VERI / "meslekler_esco.csv")
    bolumler = [r["ad"] for r in oku_csv(VERI / "bolumler.csv")]
    bolum_aciklama = {r["ad"]: r.get("kisa_aciklama") or r["ad"] for r in oku_csv(VERI / "bolumler.csv")}
    if (VERI / "bolum_metni.csv").exists():                 # b2_bolum_ders_metni.py çıktısı (embedding'e giren metin)
        bolum_aciklama.update({r["bolum"]: r["metin"] for r in oku_csv(VERI / "bolum_metni.csv", ";")})
    bad = list(np.load(ONB / "bolum_isimleri.npy", allow_pickle=True))
    bad_i = {b: i for i, b in enumerate(bad)}

    print("[2/5] Meslek ↔ iş türü ve açıklama ↔ iş türü benzerlikleri")
    Z_m, Z_b, kullanilan = [], [], []
    if os.environ.get("A8_TFIDF") == "1":
        Sm, Sb = tfidf_benzerlik([f"{o['ad']}: {o['aciklama']}" for o in meslek], [bolum_aciklama.get(b, b) for b in bad],
                                 [" ".join(t for t, _ in cumle[k]) for k in kodlar])
        Z_m.append(zs(zs(Sm, 0), 1)); Z_b.append(zs(Sb, 0)); kullanilan.append("tfidf")
    else:
        for m in MODELLER:
            I = is_turu_vektorleri(kodlar, cumle, m)
            if I is None:
                print(f"  {m}: İngilizce cümle yok — atlandı"); continue
            M = normalize(np.load(ONB / f"meslekler_{m}.npy")); B = normalize(np.load(ONB / f"bolumler_{m}.npy"))
            Z_m.append(zs(zs(M @ I.T, 0), 1)); Z_b.append(zs(B @ I.T, 0)); kullanilan.append(m)
    if not Z_m:
        raise SystemExit("Hiçbir model kullanılamadı.")
    Zm, Zb = np.mean(Z_m, 0), np.mean(Z_b, 0)
    print(f"  kullanılan: {', '.join(kullanilan)}")

    print("[3/5] Bölüm → meslek ağırlıkları (b4 kuralı: bölüm metni ↔ meslek metni)")
    from bolum_meslek import Kurallar, benzerlik       # [2026-11-01] b4 ile aynı bölüm → meslek kuralı (tek yerde)
    benz = benzerlik(ONB, MODELLER)
    kural = Kurallar(meslek)

    def agirliklar(b):
        if b not in bad_i:
            return {}
        return kural.agirliklar(b, benz[bad_i[b]], 40)

    print("[4/5] Bölüm skorları ve bağ seçimi")
    n, k = len(bolumler), len(kodlar)
    Bm, Ba, var_m = np.zeros((n, k)), np.zeros((n, k)), np.zeros(n, bool)
    for bi, b in enumerate(bolumler):
        w = agirliklar(b)
        if w:
            ii = np.array(list(w)); ww = np.array([w[i] for i in ii])
            Bm[bi] = (Zm[ii] * ww[:, None]).sum(0) / ww.sum(); var_m[bi] = True
        if b in bad_i:
            Ba[bi] = Zb[bad_i[b]]

    def birlesik(X):
        return 0.5 * zs(X, 0) + 0.5 * zs(X, 1)
    Sk = np.where(var_m[:, None], MESLEK_PAYI * birlesik(Bm) + (1 - MESLEK_PAYI) * birlesik(Ba), birlesik(Ba))

    eski = collections.defaultdict(dict)
    if (VERI / "bolum_k5_bag.csv").exists():
        for r in oku_csv(VERI / "bolum_k5_bag.csv", ";"):
            eski[r["bolum"].strip()][r["eksen"].strip()] = float(r["bag"].replace(",", "."))
    alan_of = {r["bolum"].strip(): r["alan"].strip() for r in oku_csv(VERI / "bolum_k5_bag.csv", ";")} \
        if (VERI / "bolum_k5_bag.csv").exists() else {}

    secim, rapor = {}, []
    for bi, b in enumerate(bolumler):
        sira = np.argsort(-Sk[bi])
        ust = float(Sk[bi, sira[0]])
        sec = {kodlar[sira[0]]: 1.0}
        for j in sira[1:EN_FAZLA_BAG]:
            s = float(Sk[bi, j])
            if s >= ESIK_Z and ust - s <= FARK_Z:
                sec[kodlar[j]] = round((0.25 + 0.75 * (1 - (ust - s) / FARK_Z)) * 20) / 20
        secim[b] = sec
        e = eski.get(b, {})
        e_ana = max(e, key=e.get) if e else ""
        rapor.append([b, alan_of.get(b, ""), "meslek+açıklama" if var_m[bi] else "açıklama",
                      " | ".join(f"{kodlar[j]} {eksen_ad[kodlar[j]]} ({Sk[bi, j]:.2f})" for j in sira[:5]),
                      " | ".join(f"{x}:{g:g}" for x, g in sec.items()),
                      " | ".join(f"{x}:{g:g}" for x, g in sorted(e.items(), key=lambda t: -t[1])),
                      "evet" if e_ana and e_ana == kodlar[sira[0]] else ("eski yok" if not e_ana else "hayır"),
                      "evet" if e_ana and e_ana in sec else ""])

    print("[5/5] Çıktılar")
    q = lambda x: "'" + str(x).replace("'", "''") + "'"
    with open(VERI / "bolum_k5_bag_veri.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";"); w.writerow(["bolum", "alan", "eksen", "bag"])
        for b, sec in secim.items():
            for x, g in sec.items():
                w.writerow([b, eksen_alan.get(x, ""), x, str(g).replace(".", ",")])
    degerler = ",\n".join(f"({q(b)}, {q(x)}, {g:g})" for b, sec in secim.items() for x, g in sec.items())
    sql = ["-- Filizyol: veriye dayalı K5 iş türü bağları (b5_k5_baglari.py; veri/bolum_k5_bag_veri.csv)",
           f"-- Kullanılan modeller: {', '.join(kullanilan)}. Supabase → SQL Editor'da çalıştırın; tekrar çalıştırmak güvenlidir.",
           "BEGIN;",
           "DELETE FROM bolum_k5_baglari WHERE degisken_id IN (SELECT id FROM degiskenler WHERE kod ~ '^U[0-9]{2}_');",
           "INSERT INTO bolum_k5_baglari (bolum_id, degisken_id, bag)\nSELECT b.id, g.id, v.bag FROM (VALUES\n" + degerler
           + "\n) v(ad, deg, bag) JOIN bolumler b ON b.ad = v.ad JOIN degiskenler g ON g.kod = v.deg;",
           "COMMIT;"]
    (CIKTI / "k5_baglari_veri.sql").write_text("\n".join(sql) + "\n", encoding="utf-8")
    with open(RAPOR / "b5_k5_baglari_raporu.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["bolum", "eski_alan", "kanit", "ilk5_is_turu (skor)", "yeni_bag", "eski_elle_bag", "ana_is_turu_ayni", "eski_ana_yeni_bagda"])
        w.writerows(rapor)
    var = [r for r in rapor if r[6] != "eski yok"]
    ayni = sum(r[6] == "evet" for r in var); icinde = sum(r[7] == "evet" for r in var)
    alan_ayni = sum(1 for b, sec in secim.items() if alan_of.get(b) and eksen_alan.get(max(sec, key=sec.get)) == alan_of[b])
    ozet = (f"B5 K5 bağları — modeller: {', '.join(kullanilan)}\n"
            f"Bölüm: {len(secim)} · bağ: {sum(len(s) for s in secim.values())} (bölüm başına {sum(len(s) for s in secim.values()) / len(secim):.2f})\n"
            f"Eski elle bağı olan bölüm: {len(var)}\n"
            f"  ana iş türü aynı: {ayni} (%{100 * ayni / max(1, len(var)):.0f})\n"
            f"  eski ana iş türü yeni bağlarda var: {icinde} (%{100 * icinde / max(1, len(var)):.0f})\n"
            f"  yeni ana iş türü eski üst alanla aynı alanda: {alan_ayni} (%{100 * alan_ayni / max(1, len(var)):.0f})\n"
            f"Farklı çıkan bölümler raporda (ana_is_turu_ayni = hayır) — elle düzeltilmez; yöntem/veri sorunu varsa kökten çözülür.\n")
    (RAPOR / "b5_ozet.txt").write_text(ozet, encoding="utf-8")
    print(ozet)
    print(f"Tamam → {VERI / 'bolum_k5_bag_veri.csv'}, {CIKTI / 'k5_baglari_veri.sql'}, {RAPOR}")


if __name__ == "__main__":
    main()
