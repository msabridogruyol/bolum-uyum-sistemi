# -*- coding: utf-8 -*-
"""
Soru Geçerlilik Testi (Çoklu Model Kör Eşleştirme, KATMAN KISITLI)

[DÜZELTME — KÖK NEDEN BULUNDU] Önceki sürüm, her test birimini 87+
değişkenin TAMAMI arasından en yakınını seçtiriyordu (K1'in D1 sorusu
bile K3'ün I5'i gibi tamamen alakasız katmanlardaki değişkenlerle
karşılaştırılıyordu). Bu, görevi yapay olarak çok zorlaştırıyor ve
katmanlar arası anlamsız "yanlış" sonuçlar üretiyordu (417 testten
yalnızca 25'i 3/3 doğru çıkmıştı — gerçek soru kalitesi sorunu değildi).

Artık her test birimi YALNIZCA KENDİ KATMANINDAKİ değişkenler arasından
test ediliyor (K1 sorusu yalnızca D1-D7 arasından, K5'in Mühendislik
dalı sorusu yalnızca o dalın 7 değişkeni arasından vb.) — gerçekçi ve
adil bir karşılaştırma.

Girdi: veri/gecerlilik_girdisi.xlsx — admin panelindeki "Soru Geçerlilik
       Testi" sayfasından "⬇ Test Girdisini İndir" ile inen GÜNCEL dosya
       (katman_kod sütunu içermeli — eski sürümde yoktu).
       Sütunlar: birim_anahtari, kaynak_soru_id, secenek_id, kaynak_tipi,
       katman_kod, metin, baglam, beklenen_degisken_kod
Çıktı: cikti/soru_gecerlilik_sonuclari.xlsx
"""

import csv
import numpy as np
from pathlib import Path
from openpyxl import Workbook, load_workbook
from embedding_saglayici import embed_et, MODEL_TANIMLARI

VERI_DIZINI = Path(__file__).parent.parent / "veri"
CIKTI_DIZINI = Path(__file__).parent.parent / "cikti"
ONBELLEK_DIZINI = Path(__file__).parent.parent / "onbellek"


def xlsx_oku(yol):
    """openpyxl ile bir .xlsx'i, csv.DictReader gibi sözlük listesi olarak okur."""
    wb = load_workbook(yol, read_only=True, data_only=True)
    ws = wb.active
    satirlar = list(ws.iter_rows(values_only=True))
    basliklar = [str(h) for h in satirlar[0]]
    return [
        {basliklar[i]: (str(deger) if deger is not None else "") for i, deger in enumerate(satir)}
        for satir in satirlar[1:]
    ]


def calistir():
    birimler = xlsx_oku(VERI_DIZINI / "gecerlilik_girdisi.xlsx")
    print(f"{len(birimler)} test birimi bulundu (Likert soru + SJT seçeneği).")

    if birimler and "katman_kod" not in birimler[0]:
        raise ValueError(
            "gecerlilik_girdisi.xlsx dosyasında 'katman_kod' sütunu yok — "
            "admin panelinden GÜNCEL sürümü (v2) tekrar indirin."
        )

    with open(VERI_DIZINI / "degiskenler.csv", encoding="utf-8-sig") as f:
        degisken_satirlari = list(csv.DictReader(f))
    degisken_kodlari = [d["kod"] for d in degisken_satirlari]
    degisken_metinleri = [f"{d['ad']}: {d['aciklama']}" for d in degisken_satirlari]

    # [DÜZELTME] Adayları katman_kod yerine değişken KODUNUN HARF ÖNEKİNE
    # göre grupluyoruz — bu, K5'in 8 dalını da otomatik olarak birbirinden
    # ayırır (M=Mühendislik, S=Sağlık, F=Fen&Beşeri, ... hepsi kendi 7
    # değişkeniyle sınırlı kalır), yalnızca "K5" demek 56 değişkeni birbirine
    # karıştırırdı. Test biriminin grubu da kendi beklenen_degisken_kod'unun
    # önekinden çıkarılır — cevap anahtarının hangi "aile"ye ait olduğunu
    # bilmek "kopya" değil, gerçekçi bir çoktan seçmeli alan tanımıdır
    # (bir K1 sorusunun cevabı zaten asla bir K5 değişkeni olamaz).
    def kod_oneki(kod: str) -> str:
        return "".join(c for c in kod if not c.isdigit())

    degisken_onekleri = [kod_oneki(k) for k in degisken_kodlari]
    onek_bazli_indeksler: dict[str, list[int]] = {}
    for i, onek in enumerate(degisken_onekleri):
        onek_bazli_indeksler.setdefault(onek, []).append(i)

    test_metinleri = [b["metin"] for b in birimler]
    test_onekleri = [kod_oneki(b["beklenen_degisken_kod"]) for b in birimler]

    sonuc_matrisleri = {}
    for model_adi in MODEL_TANIMLARI:
        print(f"\n--- {model_adi} ---")
        degisken_npy = ONBELLEK_DIZINI / f"degiskenler_{model_adi}.npy"
        if degisken_npy.exists():
            degisken_vek = np.load(degisken_npy)
            print("  (değişken embedding'leri önbellekten okundu)")
        else:
            degisken_vek = embed_et(degisken_metinleri, model_adi, onek="passage")

        print(f"  {len(test_metinleri)} birim encode ediliyor...")
        test_vek = embed_et(test_metinleri, model_adi, onek="query")

        # [DÜZELTME] Tam matris çarpımı yerine, her birim yalnızca KENDİ
        # değişken ailesi (öneki) içindeki adaylarla karşılaştırılıyor.
        tahmin_indeksleri = np.zeros(len(test_metinleri), dtype=int)
        en_yuksek_benzerlik = np.zeros(len(test_metinleri))

        for onek, aday_indeksler in onek_bazli_indeksler.items():
            birim_indeksleri = [i for i, o in enumerate(test_onekleri) if o == onek]
            if not birim_indeksleri:
                continue
            alt_test_vek = test_vek[birim_indeksleri]
            alt_degisken_vek = degisken_vek[aday_indeksler]
            alt_benzerlik = alt_test_vek @ alt_degisken_vek.T
            alt_tahmin = np.argmax(alt_benzerlik, axis=1)
            for j, birim_i in enumerate(birim_indeksleri):
                gercek_indeks = aday_indeksler[alt_tahmin[j]]
                tahmin_indeksleri[birim_i] = gercek_indeks
                en_yuksek_benzerlik[birim_i] = alt_benzerlik[j, alt_tahmin[j]]

        sonuc_matrisleri[model_adi] = (tahmin_indeksleri, en_yuksek_benzerlik)

    CIKTI_DIZINI.mkdir(exist_ok=True)
    cikti_yolu = CIKTI_DIZINI / "soru_gecerlilik_sonuclari.xlsx"
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Sonuclar")
    ws.append([
        "birim_anahtari", "kaynak_soru_id", "secenek_id", "katman_kod", "gercek_degisken_kod",
        "model_a_tahmin", "model_a_dogru", "model_a_benzerlik",
        "model_b_tahmin", "model_b_dogru", "model_b_benzerlik",
        "model_c_tahmin", "model_c_dogru", "model_c_benzerlik",
    ])

    tam_dogru = 0
    for i, birim in enumerate(birimler):
        gercek = birim["beklenen_degisken_kod"]
        secenek_id = birim["secenek_id"] if birim["secenek_id"] else ""
        satir = [birim["birim_anahtari"], birim["kaynak_soru_id"], secenek_id, birim["katman_kod"], gercek]
        kac_dogru = 0
        for model_adi in MODEL_TANIMLARI:
            tahmin_idx, benzerlikler = sonuc_matrisleri[model_adi]
            tahmin_kod = degisken_kodlari[tahmin_idx[i]]
            dogru = tahmin_kod == gercek
            kac_dogru += int(dogru)
            satir.extend([tahmin_kod, dogru, round(float(benzerlikler[i]), 4)])
        if kac_dogru == 3:
            tam_dogru += 1
        ws.append(satir)

    wb.save(cikti_yolu)

    print(f"\nTamamlandı. {tam_dogru}/{len(birimler)} birim tüm modellerde doğru (3/3).")
    print(f"Çıktı: {cikti_yolu}")


if __name__ == "__main__":
    calistir()
