# -*- coding: utf-8 -*-
"""
B4 — Veriye dayalı bölüm profili (2026-10-18; son güncelleme 2026-11-01: bölüm → meslek kuralı bolum_meslek.py'de)

Yapay zekâ uzman / uzlaşı katmanlarının yerine gerçek veriye dayanan bölüm profilini üretir. Çıktı, eski a4 adımıyla AYNI
biçimdedir (cikti/bolum_agirliklari.xlsx) ve admin panelinden taslak olarak yüklenir; önizleme + onay olmadan
canlıya geçmez. Önceki çıktı her çalıştırmada yedekler/bolum_agirliklari_<tarih>.xlsx olarak saklanır.
[2026-11-01] Eski a4 adımı kaldırıldı; "bolum_agirliklari_a4.xlsx" kopyası artık üretilmez (dosya yoksa o anki bölüm profili çıktısını
yanlış adla "a4" diye kopyalıyordu).

ADIMLAR
  1) Meslek profili (K2–K4 ve D5–D7, 27 değişken; P4 dahil)
     • Hedef: O*NET 31.0 öğe puanlarından değişken bileşimi (çalışan anketi, analist, yapay zekâ+uzman karma; öğe
       listesi ESLEME'de). Mesleğin eğitim düzeyi (Job Zone) etkisi çıkarılır → bölümler arası ayrım, düzey değil.
     • ESCO mesleği → O*NET mesleği: O*NET'in resmî ESCO_to_ONET-SOC eşleştirmesi.
     • Tahmin: ridge regresyon; girdiler = 3 embedding modelinin meslek vektörleri (PCA 300) + m5 puanları +
       ESCO meslek–beceri bağları (AB uzmanlarının kurduğu, SVD 300). 5 katlı, O*NET meslek grubuna göre bölünmüş
       çapraz doğrulama; başarı cikti/rapor/b4_regresyon_basarisi.csv'ye yazılır.
     • Meslek puanı = %70 O*NET'in doğrudan değeri + %30 regresyon (O*NET karşılığı olmayan meslekte yalnız regresyon).
  2) D3 Statü (meslek): SIOPS-08 prestij ölçeği (ISCO-08; Ganzeboom & Treiman) %70 + regresyon %30.
     D4 Anlam (meslek): nesnel kaynak yok → m5 embedding puanı (düşük güvenli).
  3) Bölüm → meslek [2026-10-21]: bölüm metni (resmî ders planlarındaki ders adları; bulunamayan bölümde kısa açıklama,
     b2_bolum_ders_metni.py) ↔ ESCO meslek metni, 3 embedding modeli; en benzer UST_MESLEK meslek, ağırlık = benzerlik
     farkı. Kural (ISCO-08): öğretmenlik dışı bölümlerde öğretmen meslekleri (ISCO 23) sayılmaz. Bölüm → ABD programı
     (CIP) eşlemesi ve elle hazırlanmış hiçbir liste kullanılmaz.
  4) D1 Güvence, D2 Gelir (bölüm): TÜİK Yükseköğretim İstihdam Göstergeleri 2022–2025 (veri/tuik_lisans_bolum.csv;
     b3_tuik_esle.py üretir — ad kuralları + embedding, elle eş anlamlı listesi yok):
     D1 = istihdam oranı (+) ve ilk işe giriş süresi (−); D2 = kazanç grubu. TÜİK'te olmayan (küçük/yeni) bölümde
     bölüm açıklaması embedding'inden regresyonla tahmin edilir (rapor: kaynak = 'tahmin').
  4b) Öğretmenlik bölümlerinde eğilim katmanı (A1–A9) %50 branş bölümünden. Branş, YÖK program adından otomatik bulunur
      ("FİZİK ÖĞRETMENLİĞİ" → "FİZİK"); adı birebir bir bölüme karşılık gelmeyen öğretmenliklerde branş karışımı yapılmaz.
  5) Her değişken bölümler arasında 50 ± 10'a ölçeklenir (a4 SUTUN_OLCEK ile aynı). [2026-10-22] a4'ten sütun
     aktarılmaz: çıktı yalnızca bu adımın veriden hesapladığı 31 özelliktir (P4 O*NET'ten; ESLEME'ye bakın).
  Parametreler (O_NET_PAYI, SIOPS_PAYI, UST_MESLEK, BRANS_PAYI, UST_DUZEY_KURALI) 2026-10-22'de ÖSYM puan türü dış
  ölçütüyle sınandı; denenen değerler geçerliliği anlamlı değiştirmedi (çok sınıflı doğruluk 0,82–0,86), sıralamayı ise
  belirgin değiştirdi. Bu yüzden değerler korundu; gerekçe ve sonuçlar cikti/rapor/ altında değil, proje notunda.
Kullanım:  C:\\pipeline\\venv\\Scripts\\python.exe C:\\pipeline\\src\\b4_bolum_profili.py
"""
import collections
import csv
import io
import json
import shutil
import sys
import warnings
import zipfile
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook

warnings.filterwarnings("ignore")
from sklearn.decomposition import PCA, TruncatedSVD  # noqa: E402
from sklearn.linear_model import RidgeCV  # noqa: E402
from sklearn.model_selection import GroupKFold, KFold  # noqa: E402
from scipy import sparse  # noqa: E402
from scipy.stats import spearmanr  # noqa: E402

KOK = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1]
ONB, VERI, ONET, CIKTI = KOK / "onbellek", KOK / "veri", KOK / "onet", KOK / "cikti"
RAPOR = CIKTI / "rapor"
MODELLER = ("model_a", "model_b", "model_c")
O_NET_PAYI = 0.70
K1_K4 = "D1 D2 D3 D4 D5 D6 D7 P1 P2 P3 P4 P5 P6 P7 P8 I1 I2 I3 I4 I5 I6 I7 A1 A2 A3 A4 A5 A6 A7 A8 A9".split()

# Değişken → O*NET öğeleri (+ aynı yön, − ters yön). Öğe tutarlılığı (Cronbach α) O*NET verisinde sınandı.
ESLEME = {
    "D5": ["+Assisting and Caring for Others", "+Social Service", "+Health Care Service", "+Empathy"],
    "D6": ["+Freedom to Make Decisions", "+Determine Tasks, Priorities and Goals"],
    "D7": ["+Artistic", "+Visual Arts", "+Applied Arts and Design", "+Fine Arts", "+Thinking Creatively"],
    "P1": ["+Social Orientation", "+Contact With Others", "+Face-to-Face Discussions with Individuals and Within Teams",
           "+Deal With External Customers or the Public in General"],
    "P2": ["+Cooperation", "+Empathy", "+Work With or Contribute to a Work Group or Team"],
    "P3": ["+Dependability", "+Attention to Detail", "+Organizing, Planning, and Prioritizing Work"],
    # [2026-10-22] P4 artık O*NET'ten (önceden a4'ten — yapay zekâ uzman profili içeren eski adım — aktarılıyordu).
    # Yön, sistemin kendi kutup tanımıyla aynı (veri/degisken_kutuplari.csv, P4 yüksek = "duygusal olarak yıpratıcı,
    # kaygı ve stres yaratan durumlarla karşılaşır"): işin duygusal yükü. P4 eşleşmeye girmez; yalnızca koçlukta kullanılır.
    # "Time Pressure" adayı madde-toplam r = 0,09 olduğu için alınmadı (α 0,70 → 0,84).
    "P4": ["+Stress Tolerance", "+Self-Control", "+Dealing With Unpleasant, Angry, or Discourteous People"],
    "P5": ["+Innovation", "+Intellectual Curiosity", "+Adaptability", "+Thinking Creatively",
           "+Updating and Using Relevant Knowledge"],
    "P6": ["-Importance of Repeating Same Tasks", "+Adaptability", "-Pace Determined by Speed of Equipment",
           "+Making Decisions and Solving Problems"],
    "P7": ["+Tolerance for Ambiguity", "+Level of Competition", "-Cautiousness"],
    "P8": ["+Leadership Orientation", "+Coordinate or Lead Others in Accomplishing Work Activities",
           "+Guiding, Directing, and Motivating Subordinates", "+Management of Personnel Resources"],
    "I1": ["+Time Management", "+Organizing, Planning, and Prioritizing Work", "+Scheduling Work and Activities"],
    "I2": ["+Resolving Conflicts and Negotiating with Others", "+Conflict Situations", "+Negotiation",
           "+Developing and Building Teams"],
    "I3": ["+Stress Tolerance", "+Self-Control", "+Consequence of Error", "+Frequency of Decision Making"],
    "I4": ["+Integrity", "+Sincerity", "+Evaluating Information to Determine Compliance with Standards"],
    "I5": ["+Initiative", "+Achievement Orientation", "+Determine Tasks, Priorities and Goals"],
    "I6": ["+Self-Control", "+Humility", "+Active Learning"],
    "I7": ["+Administration and Management", "+Economics and Accounting", "+Developing Objectives and Strategies",
           "+Management of Financial Resources", "+Business Initiatives"],
    "A1": ["+Mathematics", "+Mathematical Reasoning", "+Number Facility", "+Analyzing Data or Information",
           "+Mathematics/Statistics"],
    "A2": ["+English Language", "+Writing", "+Reading Comprehension", "+Written Expression", "+Humanities",
           "+Creative Writing"],
    "A3": ["+Social", "+Psychology", "+Therapy and Counseling", "+Social Perceptiveness", "+Teaching/Education"],
    "A4": ["+Design", "+Visualization", "+Drafting, Laying Out, and Specifying Technical Devices, Parts, and Equipment",
           "+Applied Arts and Design", "+Visual Arts"],
    "A5": ["+Biology", "+Chemistry", "+Life Science", "+Nature/Outdoors", "+Science"],
    "A6": ["+Performing General Physical Activities", "+Stamina", "+Gross Body Coordination",
           "+Spend Time Walking or Running", "+Athletics"],
    "A7": ["+Conventional", "+Importance of Being Exact or Accurate", "+Attention to Detail", "+Information Ordering",
           "+Documenting/Recording Information"],
    "A8": ["+Originality", "+Fluency of Ideas", "+Inductive Reasoning", "+Thinking Creatively",
           "-Importance of Repeating Same Tasks"],
    "A9": ["+Enterprising", "+Selling or Influencing Others", "+Persuasion", "+Sales", "+Business Initiatives"],
}
BRANS = {}   # [2026-10-21] elle liste kaldırıldı; main() içinde YÖK program adından üretilir
UST_MESLEK = 40
BRANS_PAYI = 0.5
SIOPS_PAYI = 0.70

# O*NET dosyası → kullanılan ölçek. Aynı adlı öğede ilk dosya geçerlidir (ör. Mathematics = Knowledge).
KAYNAK = [("work_context", "CX"), ("work_activities", "IM"), ("knowledge", "IM"), ("essential_skills", "IM"),
          ("transferable_skills", "IM"), ("work_styles", "WI"), ("specific_interest_areas", "OI"),
          ("career_interest_types", "OI"), ("abilities", "IM")]


def oku_csv(yol, ayirici=","):
    with open(yol, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=ayirici))


def z(x):
    x = np.asarray(x, float)
    return (x - np.nanmean(x)) / (np.nanstd(x) + 1e-12)


def onet_oku():
    zp = zipfile.ZipFile(ONET / "db_31_0_csv.zip")
    ad = {Path(n).stem: n for n in zp.namelist() if n.endswith(".csv")}
    oge, sahip = collections.defaultdict(dict), {}
    for dosya, olcek in KAYNAK:
        for r in csv.DictReader(io.TextIOWrapper(zp.open(ad[dosya]), encoding="utf-8")):
            if r["Scale ID"] != olcek:
                continue
            e = r["Element Name"]
            sahip.setdefault(e, dosya)
            if sahip[e] == dosya:
                oge[e][r["O*NET-SOC Code"]] = float(r["Data Value"])
    jz = {r["O*NET-SOC Code"]: float(r["Job Zone"])
          for r in csv.DictReader(io.TextIOWrapper(zp.open(ad["job_zones"]), encoding="utf-8"))}
    return oge, jz


def xlsx_satirlar(yol, atla):
    ws = load_workbook(yol, read_only=True).worksheets[0]
    return list(ws.iter_rows(values_only=True))[atla:]


def main():
    RAPOR.mkdir(parents=True, exist_ok=True)
    print("[1/6] Girdiler okunuyor")
    meslek = oku_csv(VERI / "meslekler_esco.csv")
    kod_m = list(np.load(ONB / "degisken_kodlari.npy", allow_pickle=True))
    a3 = {m: np.load(ONB / f"m5_skor_matrisi_{m}.npy") for m in MODELLER}
    emb = [np.load(ONB / f"meslekler_{m}.npy") for m in MODELLER]
    X_emb = PCA(300, random_state=0).fit_transform(np.hstack([(e - e.mean(0)) / (e.std(0) + 1e-9) for e in emb]))
    # ESCO meslek–beceri bağları (temel = 1, isteğe bağlı = 0,5; IDF; satır normu)
    uri_i = {o["esco_uri"]: i for i, o in enumerate(meslek)}
    bag = oku_csv(VERI / "occupationSkillRelations_en.csv")
    say = collections.Counter(r["skillUri"] for r in bag)
    bj = {u: j for j, u in enumerate(u for u, n in say.items() if n >= 8)}
    R, C, D = [], [], []
    for r in bag:
        i, j = uri_i.get(r["occupationUri"]), bj.get(r["skillUri"])
        if i is not None and j is not None:
            R.append(i); C.append(j); D.append(1.0 if r["relationType"] == "essential" else 0.5)
    S = sparse.csr_matrix((D, (R, C)), shape=(len(meslek), len(bj)))
    idf = np.log(len(meslek) / np.maximum(1, np.asarray((S > 0).sum(0)).ravel()))
    S = S.multiply(idf).tocsr()
    S = sparse.diags(1 / (np.sqrt(np.asarray(S.multiply(S).sum(1)).ravel()) + 1e-9)) @ S
    X_bec = TruncatedSVD(300, random_state=0).fit_transform(S)

    oge, jz = onet_oku()
    cw = collections.defaultdict(set)
    for r in xlsx_satirlar(ONET / "ESCO_to_ONET-SOC.xlsx", 4):
        if r[0]:
            cw[str(r[0]).strip()].add(r[2])
    socs = sorted(jz)
    JZ = z([jz[s] for s in socs])

    print("[2/6] Meslek profilleri (O*NET + regresyon)")
    basari, puan, paylar = [], {}, {}
    for v, liste in ESLEME.items():
        M = []
        for o in liste:
            x = np.array([oge.get(o[1:], {}).get(s, np.nan) for s in socs])
            if np.isnan(x).all():
                raise SystemExit(f"O*NET öğesi bulunamadı: {o} ({v})")
            M.append(z(x) * (1 if o[0] == "+" else -1))
        t = np.nanmean(M, 0)
        ok = ~np.isnan(t)
        b = np.polyfit(JZ[ok], t[ok], 1)
        hedef = dict(zip(socs, t - (b[0] * JZ + b[1])))            # eğitim düzeyi etkisi çıkarıldı
        y, idx, grp, dogrudan = [], [], [], np.full(len(meslek), np.nan)
        for i, o in enumerate(meslek):
            ss = [s for s in cw.get(o["esco_kodu"], ()) if not np.isnan(hedef.get(s, np.nan))]
            if ss:
                dogrudan[i] = np.mean([hedef[s] for s in ss]); y.append(dogrudan[i]); idx.append(i)
                grp.append(sorted(ss)[0][:5])
        j = kod_m.index(v)
        taban = np.column_stack([a3[m][:, j] for m in MODELLER]) / 50
        Xe = np.hstack([X_emb, taban]); X = np.hstack([Xe / Xe.std(), X_bec / X_bec.std()])
        y, idx = np.array(y), np.array(idx)
        cv = np.zeros(len(y))
        for tr, te in GroupKFold(5).split(X[idx], y, grp):
            cv[te] = RidgeCV(alphas=np.logspace(-1, 4, 11)).fit(X[idx][tr], y[tr]).predict(X[idx][te])
        tahmin = RidgeCV(alphas=np.logspace(-1, 4, 11)).fit(X[idx], y).predict(X)
        tahmin[idx] = cv                                            # eğitim mesleklerinde sızıntısız tahmin
        pay = float(O_NET_PAYI)
        paylar[v] = round(pay, 3)
        son = np.where(np.isnan(dogrudan), tahmin, pay * np.nan_to_num(dogrudan) + (1 - pay) * tahmin)
        puan[v] = z(son)
        basari.append((v, len(y), round(float(spearmanr(cv, y)[0]), 3),
                       round(float(spearmanr(np.mean(taban, 1)[idx], y)[0]), 3)))
        print(f"   {v}: regresyon r={basari[-1][2]:.2f} (eski m5 r={basari[-1][3]:.2f}), O*NET'li meslek {len(y)}, O*NET payı {pay:.2f}")

    # D3 Statü: SIOPS-08 (ISCO 4 basamak)
    siops = json.load(open(VERI / "isco08_siops.json", encoding="utf-8"))["siops08"]
    isco = [o["isco_grubu"].zfill(4) for o in meslek]
    y3 = np.array([siops.get(c, np.nan) for c in isco], float)
    ok3 = ~np.isnan(y3); p3 = np.zeros(len(meslek))
    X3 = X_emb[ok3]; g3 = [c[:3] for c in np.array(isco)[ok3]]; cv3 = np.zeros(ok3.sum())
    for tr, te in GroupKFold(5).split(X3, y3[ok3], g3):
        cv3[te] = RidgeCV(alphas=np.logspace(0, 5, 11)).fit(X3[tr], y3[ok3][tr]).predict(X3[te])
    p3 = RidgeCV(alphas=np.logspace(0, 5, 11)).fit(X3, y3[ok3]).predict(X_emb); p3[ok3] = cv3
    puan["D3"] = z(np.where(ok3, SIOPS_PAYI * z(np.nan_to_num(y3)) + (1 - SIOPS_PAYI) * z(p3), z(p3)))
    basari.append(("D3", int(ok3.sum()), round(float(spearmanr(cv3, y3[ok3])[0]), 3), None))
    puan["D4"] = z(np.mean([a3[m][:, kod_m.index("D4")] for m in MODELLER], 0))   # nesnel kaynak yok
    basari.append(("D4", 0, None, None))

    print("[3/6] Bölüm → meslek bağlantıları (bölüm metni ↔ meslek metni)")
    bolumler = [r["ad"] for r in oku_csv(VERI / "bolumler.csv")]
    bad = list(np.load(ONB / "bolum_isimleri.npy", allow_pickle=True))
    from bolum_meslek import Kurallar, benzerlik       # [2026-11-01] b4 / b5 / b6 ortak bölüm → meslek kuralı
    benz = benzerlik(ONB, MODELLER)
    bad_i = {b: i for i, b in enumerate(bad)}

    kural = Kurallar(meslek)

    def agirliklar(b):
        # Bölüm metni ↔ ESCO meslek metni + ISCO-08 kuralları (bolum_meslek.py). CIP / elle eşleme yok.
        if b not in bad_i:
            return {}, "yok"
        return kural.agirliklar(b, benz[bad_i[b]], UST_MESLEK), "benzerlik"

    print("[4/6] Bölüm profilleri")
    P = np.full((len(bolumler), len(K1_K4)), np.nan); VAR = np.zeros_like(P); N = np.zeros(len(bolumler), int)
    kaynak = {}
    for bi, b in enumerate(bolumler):
        w, kay = agirliklar(b); kaynak[b] = [kay, len(w)]
        if not w:
            continue
        ii = np.array(list(w)); ww = np.array([w[i] for i in ii]); N[bi] = len(ii)
        for vj, v in enumerate(K1_K4):
            if v in puan:
                x = puan[v][ii]; P[bi, vj] = (x * ww).sum() / ww.sum(); VAR[bi, vj] = ((x - P[bi, vj]) ** 2 * ww).sum() / ww.sum()

    # Öğretmenlik bölümleri: O*NET öğretmen meslekleri branş ayırmıyor (ör. tüm lise öğretmenleri tek meslek). Eğilim
    # katmanı (A1–A9) %50 öğretmenlik + %50 branş bölümünün profili; değer / kişilik / yetkinlik öğretmenlik mesleğinden.
    bi_ = {b: i for i, b in enumerate(bolumler)}
    for b in bolumler:                                    # YÖK program adından: "X ÖĞRETMENLİĞİ" → "X" (varsa)
        if b.endswith(" ÖĞRETMENLİĞİ") and b[: -len(" ÖĞRETMENLİĞİ")] in bi_:
            BRANS[b] = [b[: -len(" ÖĞRETMENLİĞİ")]]
    a_j = [K1_K4.index(k) for k in K1_K4 if k.startswith("A")]
    for ogr, brans in BRANS.items():
        if ogr in bi_ and all(x in bi_ for x in brans):
            P[bi_[ogr], a_j] = (1 - BRANS_PAYI) * P[bi_[ogr], a_j] + BRANS_PAYI * np.nanmean([P[bi_[x], a_j] for x in brans], 0)
            kaynak[ogr][0] += "+branş"

    print("[5/6] D1 Güvence, D2 Gelir (TÜİK)")
    tuik = {r["bolum"]: r for r in oku_csv(VERI / "tuik_lisans_bolum.csv", ";")}
    sf = lambda b, k: float(tuik[b][k]) if b in tuik and str(tuik[b].get(k, "")).strip() else np.nan
    ist = np.array([sf(b, "istihdam_orani") for b in bolumler])
    sur = np.array([sf(b, "ilk_is_suresi_ay") for b in bolumler])
    kaz = np.array([sf(b, "kazanc_grubu_1_5") for b in bolumler])
    E = np.hstack([np.load(ONB / f"bolumler_{m}.npy") for m in MODELLER])
    Eb = np.array([E[bad_i[b]] if b in bad_i else np.zeros(E.shape[1]) for b in bolumler])
    for v, y in (("D1", np.nanmean([z(ist), -z(sur)], 0)), ("D2", z(kaz))):
        vj = K1_K4.index(v); ok = ~np.isnan(y)
        Xf = np.hstack([Eb, np.nan_to_num(P[:, [K1_K4.index("D3")]])])
        cvp = np.zeros(ok.sum())
        for tr, te in KFold(5, shuffle=True, random_state=0).split(Xf[ok]):
            cvp[te] = RidgeCV(alphas=np.logspace(-1, 4, 11)).fit(Xf[ok][tr], y[ok][tr]).predict(Xf[ok][te])
        tam = RidgeCV(alphas=np.logspace(-1, 4, 11)).fit(Xf[ok], y[ok]).predict(Xf)
        P[:, vj] = np.where(ok, y, tam)
        basari.append((v, int(ok.sum()), round(float(spearmanr(cvp, y[ok])[0]), 3), "bölüm düzeyi; TÜİK'te olmayan bölümde tahmin"))
    for b in bolumler:
        tam = b in tuik and all(str(tuik[b].get(k, "")).strip() for k in ("istihdam_orani", "ilk_is_suresi_ay", "kazanc_grubu_1_5"))
        kaynak[b].append("TÜİK" if tam else "TÜİK+tahmin (eksik gösterge)" if b in tuik else "tahmin")

    print("[6/6] Çıktı yazılıyor")
    eski = CIKTI / "bolum_agirliklari.xlsx"
    # [2026-10-22] Önceki çıktı zaman damgalı yedeklenir (eski sürüm her çalıştırmada a4 yedeğinin üzerine yazıyordu).
    # a4'ten sütun AKTARILMAZ: yalnızca bu adımın veriden hesapladığı 31 özellik yazılır (P4 dahil). Eski K5 dal
    # sütunları (C/E/F/M/R/S/T/U) puanlamada kullanılmıyor; canlıdaki eski sürümleri etkisizdir.
    if eski.exists():
        yedek = KOK / "yedekler"; yedek.mkdir(exist_ok=True)
        from datetime import datetime
        shutil.copy(eski, yedek / f"bolum_agirliklari_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    Pz = 50 + 10 * np.nan_to_num((P - np.nanmean(P, 0)) / (np.nanstd(P, 0) + 1e-12))
    wb = Workbook(); ws = wb.active; ws.title = "Bolum Agirliklari"
    ws.append(["bolum_adi", "degisken_kod", "agirlik_degeri", "yakinsama_skoru", "agirlikli_varyans", "etkin_meslek_sayisi"])
    for bi, b in enumerate(bolumler):
        for vj, v in enumerate(K1_K4):
            ws.append([b, v, round(float(np.clip(Pz[bi, vj], 0, 100)), 2), 0.0, round(float(VAR[bi, vj]), 4), int(N[bi])])
    wb.save(eski)
    with open(RAPOR / "b4_regresyon_basarisi.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";"); w.writerow(["degisken", "egitim_meslek_sayisi", "capraz_dogrulama_r", "eski_m5_r_veya_not"])
        w.writerows(basari)
    with open(RAPOR / "b4_bolum_kaynaklari.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";"); w.writerow(["bolum", "meslek_kaynagi", "meslek_sayisi", "d1_d2_kaynagi"])
        for b in bolumler:
            w.writerow([b, *kaynak[b]])
    print(f"Tamam: {len(bolumler)} bölüm → {eski}  (önceki çıktı: yedekler/; raporlar: {RAPOR})")


if __name__ == "__main__":
    main()
